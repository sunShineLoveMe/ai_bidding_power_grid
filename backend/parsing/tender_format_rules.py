"""从可编辑 DOCX 中保真提取投标文件格式表及其提交标记。"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _page_breaks(element: Any) -> int:
    return len(element.findall(".//" + qn("w:lastRenderedPageBreak"))) + len(
        element.findall(f'.//{qn("w:br")}[@{qn("w:type")}="page"]')
    )


def _table_pages(document: Document) -> dict[int, list[int]]:
    table_pages: dict[int, list[int]] = {}
    page = 1
    table_index = 0
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:tbl"):
            row_pages: list[int] = []
            for row in element.findall(qn("w:tr")):
                row_pages.append(page)
                page += _page_breaks(row)
            table_pages[table_index] = row_pages
            table_index += 1
        else:
            page += _page_breaks(element)
    return table_pages


def _document_text(document: Document) -> str:
    blocks = [_clean(paragraph.text) for paragraph in document.paragraphs if _clean(paragraph.text)]
    for table in document.tables:
        for row in table.rows:
            text = " | ".join(_clean(cell.text) for cell in row.cells)
            if text:
                blocks.append(text)
    return "\n".join(blocks)


def _volume_type(title: str) -> str | None:
    if "价格文件" in title:
        return "price"
    if "商务文件" in title:
        return "business"
    if "技术文件" in title:
        return "technical"
    return None


def extract_docx_format_rule_inputs(
    source_path: str | Path,
    *,
    source_display_name: str | None = None,
) -> dict[str, Any]:
    """提取最符合“价格/商务/技术 + 提交方式”的格式清单表。

    章节编号只保留为追溯信息，表格中的提交标记才参与是否纳入目录的判断。
    """
    path = Path(source_path)
    document = Document(path)
    table_pages = _table_pages(document)
    source_file = source_display_name or path.name
    clauses: list[str] = []
    best_table_index: int | None = None
    best_score = 0

    for table_index, table in enumerate(document.tables):
        table_text = "\n".join(" | ".join(_clean(cell.text) for cell in row.cells) for row in table.rows)
        score = sum(token in table_text for token in ("价格文件", "商务文件", "技术文件", "提交方式", "是否有格式要求"))
        if score > best_score:
            best_score = score
            best_table_index = table_index
        for row in table.rows:
            cells = [_clean(cell.text) for cell in row.cells]
            joined = " | ".join(cells)
            if (
                cells
                and re.fullmatch(r"\d+(?:\.\d+)*", cells[0] or "")
                and any(token in joined for token in (
                    "投标文件", "资格预审", "补充", "更新", "到期", "投标保证金",
                    "纸质", "签章", "按包", "按分标", "人员关系",
                ))
            ):
                clauses.append(joined)

    if best_table_index is None or best_score < 4:
        return {
            "source_file": source_file,
            "format_rows": [],
            "clauses": clauses,
            "document_text": _document_text(document),
            "warning": "未找到包含价格/商务/技术文件和提交标记的投标文件格式表",
        }

    rows: list[dict[str, Any]] = []
    table = document.tables[best_table_index]
    current_volume: str | None = None
    current_scope: str | None = None
    parent_key: str | None = None
    for row_index, row in enumerate(table.rows):
        cells = [_clean(cell.text) for cell in row.cells]
        if len(cells) < 3:
            continue
        sequence, title = cells[0], cells[1]
        detected_volume = _volume_type(title)
        if detected_volume:
            current_volume = detected_volume
            current_scope = "by_package" if "按包" in title else "by_lot" if "按分标" in title else None
            parent_key = None
            continue
        if not current_volume or row_index < 2 or not title:
            continue
        if sequence and "." not in sequence:
            parent_key = f"{current_volume}.{sequence}"
        pages = table_pages.get(best_table_index) or []
        rows.append({
            "sequence": sequence,
            "title": title,
            "volume_type": current_volume,
            "required_marker": cells[2],
            "format_marker": cells[3] if len(cells) > 3 else "",
            "port": cells[4] if len(cells) > 4 else "",
            "submission_scope": current_scope,
            "rule_scope": "package" if current_scope == "by_package" else "lot",
            "parent_key": parent_key if "." in sequence else None,
            "strict_format_table": True,
            "source_file": source_file,
            "source_section": "投标文件格式/投标文件组成清单",
            "source_page": pages[row_index] if row_index < len(pages) else None,
            "document_role": "main_tender_file",
            "detection_method": "native_docx_format_table_marker",
            "confidence": 0.99,
        })
    return {
        "source_file": source_file,
        "format_rows": rows,
        "clauses": clauses,
        "document_text": _document_text(document),
        "source_table_index": best_table_index,
    }
