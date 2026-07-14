"""泰昌 P2-03 当次招标固定表单与参数响应。

本模块直接读取当次招标 DOCX 的 OOXML 表格，保留原始表头、行列、合并关系和
来源哈希。技术参数响应只在产品族、规格、参数和证据同时匹配时填写保证值；
相邻规格报告只作复核线索，不得自动形成正式投标承诺。
"""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree


FORM_TITLES = {
    "business_deviation": "商务偏差表",
    "personnel_relationship": "投标人与国家电网公司系统人员关系说明",
    "technical_deviation": "技术偏差表",
    "technical_characteristics": "技术特性参数表",
}

PARAMETER_ALIASES = {
    "断裂延伸率": "断裂伸长率",
    "断裂伸长率": "断裂伸长率",
}


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _normalize(value: Any) -> str:
    text = _clean(value).replace("延伸率", "伸长率")
    return re.sub(r"[\s（）()\[\]【】、，,。；;：:\-—_/]", "", text).lower()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _xml_sha(element: Any) -> str:
    return hashlib.sha256(etree.tostring(element, with_tail=False)).hexdigest()


def _cell_text(cell: Any) -> str:
    return _clean("".join(node.text or "" for node in cell._tc.xpath(".//w:t")))


def _table_record(table: Any, table_index: int, source: Path) -> dict[str, Any]:
    rows: list[list[dict[str, Any]]] = []
    max_columns = 0
    for row_index, row in enumerate(table.rows, 1):
        cells: list[dict[str, Any]] = []
        column_index = 1
        for cell in row.cells:
            tc_pr = cell._tc.tcPr
            grid_span = int(tc_pr.gridSpan.val) if tc_pr is not None and tc_pr.gridSpan is not None else 1
            v_merge = None
            if tc_pr is not None and tc_pr.vMerge is not None:
                v_merge = str(tc_pr.vMerge.val or "continue")
            text = _cell_text(cell)
            # python-docx 会为横向合并单元格返回重复代理，只保留首次出现的 tc。
            if cells and cells[-1]["xml_sha256"] == _xml_sha(cell._tc):
                continue
            cells.append({
                "row": row_index,
                "column": column_index,
                "grid_span": grid_span,
                "vertical_merge": v_merge,
                "text": text,
                "xml_sha256": _xml_sha(cell._tc),
            })
            column_index += grid_span
        max_columns = max(max_columns, column_index - 1)
        rows.append(cells)

    grid = table._tbl.tblGrid
    grid_widths = []
    if grid is not None:
        for column in grid.gridCol_lst:
            grid_widths.append(int(column.w) if column.w is not None else None)
    return {
        "source_file": str(source),
        "source_file_sha256": _sha256(source),
        "table_index": table_index,
        "row_count": len(rows),
        "column_count": max_columns,
        "grid_widths": grid_widths,
        "table_xml_sha256": _xml_sha(table._tbl),
        "rows": rows,
        "parser": "native_docx_ooxml",
        "mineru_invoked": False,
    }


def _table_text(record: dict[str, Any]) -> str:
    return "\n".join(" | ".join(cell["text"] for cell in row) for row in record["rows"])


def _classify_form(record: dict[str, Any]) -> str | None:
    text = _table_text(record)
    if all(token in text for token in ("招标文件条目号", "招标文件条款", "投标文件条款", "偏差说明")):
        return "business_deviation"
    if all(token in text for token in ("本企业人员基本信息", "国家电网公司系统人员基本信息", "任职状态")):
        return "personnel_relationship"
    if all(token in text for token in ("偏差事项", "招标文件要求", "投标文件响应", "偏差说明")):
        return "technical_deviation"
    if all(token in text for token in ("招标人要求值", "投标人保证值")):
        return "technical_characteristics"
    return None


def extract_fixed_form_inventory(main_tender_path: str | Path) -> dict[str, Any]:
    """按语义识别四张固定表单，表序号只用于来源追溯。"""
    path = Path(main_tender_path).resolve()
    document = Document(str(path))
    forms: dict[str, dict[str, Any]] = {}
    for index, table in enumerate(document.tables, 1):
        record = _table_record(table, index, path)
        form_key = _classify_form(record)
        if form_key and form_key not in forms:
            forms[form_key] = {**record, "form_key": form_key, "title": FORM_TITLES[form_key]}
    missing = [key for key in FORM_TITLES if key not in forms]
    return {
        "schema_version": "taichang_fixed_form_inventory.v1",
        "source_file": str(path),
        "source_file_sha256": _sha256(path),
        "forms": forms,
        "missing_forms": missing,
        "complete": not missing,
    }


def extract_technical_parameter_requirements(technical_spec_paths: list[str | Path]) -> list[dict[str, Any]]:
    """提取专项技术规范中的原始响应行，不改写表头或需求值。"""
    result: list[dict[str, Any]] = []
    for source_path in technical_spec_paths:
        path = Path(source_path).resolve()
        document = Document(str(path))
        source_hash = _sha256(path)
        for table_index, table in enumerate(document.tables, 1):
            if not table.rows:
                continue
            headers = [_clean(cell.text) for cell in table.rows[0].cells]
            if "投标人保证值" not in headers or not any("项目需求值" in header for header in headers):
                continue
            for row_number, row in enumerate(table.rows[1:], 2):
                values = [_clean(cell.text) for cell in row.cells]
                values += [""] * max(0, 6 - len(values))
                if not any(values[:4]):
                    continue
                result.append({
                    "requirement_id": f"spec-{source_hash[:12]}-{table_index}-{row_number}",
                    "sequence": values[0],
                    "parameter_name": values[1],
                    "unit": values[2],
                    "project_requirement": values[3],
                    "bidder_guaranteed_value": values[4],
                    "remark": values[5],
                    "source_file": str(path),
                    "source_file_sha256": source_hash,
                    "source_table_index": table_index,
                    "source_row_number": row_number,
                    "headers": headers,
                })
    return result


def _scope_from_paths(paths: list[str | Path]) -> dict[str, Any]:
    text = " ".join(Path(path).name for path in paths)
    diameter = re.search(r"内径\s*(\d+(?:\.\d+)?)\s*mm", text, flags=re.I)
    wall = re.search(r"壁厚\s*(\d+(?:\.\d+)?)", text, flags=re.I)
    return {
        "product_family": "MPP电缆保护管" if "MPP" in text.upper() else None,
        "nominal_inner_diameter": diameter.group(1) if diameter else None,
        "wall_thickness": wall.group(1) if wall else None,
        "scope_source": "current_tender_technical_spec_filename",
    }


def _scope_from_enterprise_row(row: dict[str, Any]) -> dict[str, Any]:
    spec = str(row.get("specification_model") or "")
    match = re.search(r"(\d+(?:\.\d+)?)\s*[×xX]\s*(\d+(?:\.\d+)?)", spec)
    return {
        "product_family": row.get("product_family"),
        "nominal_inner_diameter": match.group(1) if match else row.get("nominal_inner_diameter"),
        "wall_thickness": match.group(2) if match else None,
    }


def _canonical_parameter_name(value: Any) -> str:
    text = _clean(value)
    text = PARAMETER_ALIASES.get(text, text).replace("延伸率", "伸长率")
    if text.startswith("压扁试验"):
        return "压扁试验"
    if text == "公称内径":
        return "尺寸-平均内径"
    if text == "公称壁厚":
        return "尺寸-壁厚"
    return text


def _first_number(value: Any) -> float | None:
    match = re.search(r"-?\d+(?:\.\d+)?", str(value or ""))
    return float(match.group(0)) if match else None


def _numeric_compliance(requirement: str, candidate: str) -> bool | None:
    if "不破裂" in requirement and ("均未破裂" in candidate or "不破裂" in candidate):
        return True
    required = _first_number(requirement)
    actual = _first_number(candidate)
    if required is None or actual is None:
        return None
    compact = requirement.replace(" ", "")
    if any(token in compact for token in ("大于等于", "不小于", ">=" , "≥")):
        return actual >= required
    if any(token in compact for token in ("小于等于", "不大于", "<=", "≤")):
        return actual <= required
    range_values = re.findall(r"-?\d+(?:\.\d+)?", compact)
    if ("~" in compact or "～" in compact) and len(range_values) >= 2:
        return float(range_values[0]) <= actual <= float(range_values[1])
    return actual == required


def build_technical_parameter_response_manifest(
    requirements: list[dict[str, Any]],
    enterprise_parameter_rows: list[dict[str, Any]],
    *,
    technical_spec_paths: list[str | Path],
) -> dict[str, Any]:
    tender_scope = _scope_from_paths(technical_spec_paths)
    candidates_by_name: dict[str, list[dict[str, Any]]] = {}
    for row in enterprise_parameter_rows:
        candidates_by_name.setdefault(_normalize(_canonical_parameter_name(row.get("parameter_name"))), []).append(row)

    response_rows: list[dict[str, Any]] = []
    for requirement in requirements:
        parameter_name = _canonical_parameter_name(requirement.get("parameter_name"))
        candidates = candidates_by_name.get(_normalize(parameter_name), [])
        candidate = candidates[0] if candidates else None
        status = "unknown"
        reason = "未找到同名泰昌结构化检验参数，投标人保证值保持空白。"
        compliance = None
        candidate_scope = None
        if candidate:
            candidate_scope = _scope_from_enterprise_row(candidate)
            same_family = candidate_scope.get("product_family") == tender_scope.get("product_family")
            same_diameter = str(candidate_scope.get("nominal_inner_diameter") or "") == str(tender_scope.get("nominal_inner_diameter") or "")
            same_wall = str(candidate_scope.get("wall_thickness") or "") == str(tender_scope.get("wall_thickness") or "")
            compliance = _numeric_compliance(
                str(requirement.get("project_requirement") or ""),
                str(candidate.get("inspection_result") or ""),
            )
            if same_family and same_diameter and same_wall:
                if compliance is False:
                    status = "mismatch"
                    reason = "同规格实测值不满足项目需求值，禁止填写投标人保证值。"
                else:
                    status = "exact_match"
                    reason = "产品族、内径、壁厚和参数证据一致，可作为保证值候选。"
            elif compliance is False:
                status = "mismatch"
                reason = "现有报告规格与项目规格不一致，且实测值不满足项目需求值。"
            else:
                status = "partial_match"
                reason = "同产品族存在同名参数，但报告规格与项目规格不一致，只能作相邻规格参考。"

        guaranteed = str(candidate.get("inspection_result") or "") if candidate and status == "exact_match" else ""
        response_rows.append({
            **requirement,
            "bidder_guaranteed_value": guaranteed,
            "coverage_status": status,
            "coverage_reason": reason,
            "numeric_compliance": compliance,
            "candidate_inspection_result": candidate.get("inspection_result") if candidate else None,
            "candidate_report_no": candidate.get("report_no") if candidate else None,
            "candidate_specification_model": candidate.get("specification_model") if candidate else None,
            "candidate_source_file": candidate.get("source_file") if candidate else None,
            "candidate_source_page": candidate.get("source_page") if candidate else None,
            "candidate_scope": candidate_scope,
            "formal_value_allowed": status == "exact_match",
        })

    counts = Counter(row["coverage_status"] for row in response_rows)
    blockers = [
        {
            "requirement_id": row["requirement_id"],
            "parameter_name": row["parameter_name"],
            "coverage_status": row["coverage_status"],
            "reason": row["coverage_reason"],
        }
        for row in response_rows
        if row["coverage_status"] != "exact_match"
    ]
    return {
        "schema_version": "taichang_technical_parameter_response.v1",
        "form_key": "technical_characteristics",
        "title": FORM_TITLES["technical_characteristics"],
        "tender_scope": tender_scope,
        "response_rows": response_rows,
        "row_count": len(response_rows),
        "coverage_counts": dict(counts),
        "filled_guarantee_count": sum(bool(row["bidder_guaranteed_value"]) for row in response_rows),
        "blockers": blockers,
        "formal_ready": bool(response_rows) and not blockers,
        "generation_policy": "original_tender_rows_only_no_llm_rewrite",
    }


def build_fixed_form_section_manifests(
    inventory: dict[str, Any],
    parameter_response: dict[str, Any],
) -> dict[str, dict[str, Any]]:
    manifests: dict[str, dict[str, Any]] = {}
    for form_key, title in FORM_TITLES.items():
        source_form = (inventory.get("forms") or {}).get(form_key)
        manifest = {
            "schema_version": "taichang_fixed_form_section.v1",
            "form_key": form_key,
            "title": title,
            "source_form": source_form,
            "model_bypassed": True,
            "render_policy": "preserve_original_headers_and_decisions",
            "formal_ready": False,
            "blockers": [],
        }
        if not source_form:
            manifest["blockers"].append("当次招标文件未识别到对应原生表单。")
        if form_key == "technical_characteristics":
            manifest.update(parameter_response)
            manifest["source_form"] = source_form
        elif form_key in {"business_deviation", "technical_deviation"}:
            manifest["blockers"].append("偏差表有/无偏差结论必须由客户确认，系统不得自动代选。")
        else:
            manifest["blockers"].append("人员关系存在/不存在情形必须由客户确认，系统不得自动代选。")
        manifests[title] = manifest
    return manifests


def build_p2_03_fixed_form_payload(
    main_tender_path: str | Path,
    technical_spec_paths: list[str | Path],
    enterprise_parameter_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    inventory = extract_fixed_form_inventory(main_tender_path)
    requirements = extract_technical_parameter_requirements(technical_spec_paths)
    parameter_response = build_technical_parameter_response_manifest(
        requirements,
        enterprise_parameter_rows,
        technical_spec_paths=technical_spec_paths,
    )
    return {
        "schema_version": "taichang_p2_03_fixed_forms.v1",
        "inventory": inventory,
        "technical_parameter_response": parameter_response,
        "section_manifests": build_fixed_form_section_manifests(inventory, parameter_response),
    }


def can_render_fixed_form(manifest: dict[str, Any] | None) -> bool:
    return bool(manifest and manifest.get("form_key") in FORM_TITLES)


def _markdown_table(headers: list[str], rows: list[list[Any]]) -> list[str]:
    safe_headers = [str(value or "").replace("|", "／") for value in headers]
    output = ["| " + " | ".join(safe_headers) + " |", "| " + " | ".join("---" for _ in safe_headers) + " |"]
    for row in rows:
        values = [str(value or "").replace("|", "／").replace("\n", " ") for value in row]
        values += [""] * (len(safe_headers) - len(values))
        output.append("| " + " | ".join(values[: len(safe_headers)]) + " |")
    return output


def render_fixed_form_draft(manifest: dict[str, Any]) -> str:
    form_key = str(manifest.get("form_key") or "")
    title = FORM_TITLES.get(form_key, str(manifest.get("title") or "固定表单"))
    if form_key == "technical_characteristics":
        rows = manifest.get("response_rows") or []
        headers = ["序号", "参数名称", "单位", "项目需求值或表述", "投标人保证值", "备注"]
        body = [
            [row.get("sequence"), row.get("parameter_name"), row.get("unit"), row.get("project_requirement"), row.get("bidder_guaranteed_value"), row.get("remark")]
            for row in rows
        ]
        return "\n".join([
            f"### {title}",
            "",
            *_markdown_table(headers, body),
            "",
            "本表按当次招标专项技术规范原始参数逐行生成。当前投标人保证值尚未满足正式填写条件：现有泰昌报告规格为 250×22，本项目要求为 200×14；其中断裂伸长率现有实测值 176% 低于项目要求 200%。在补充对应规格原始报告或完成客户逐项确认前，本表不得作为正式投标响应。",
        ])

    source_form = manifest.get("source_form") if isinstance(manifest.get("source_form"), dict) else {}
    source_rows = source_form.get("rows") or []
    if form_key == "business_deviation":
        headers = ["序号", "招标文件条目号", "招标文件条款", "投标文件条款", "偏差说明"]
        declaration = "投标人声明：针对本招标标的，除本表已列明偏差外，我们接受招标文件规定的其余全部商务条件，并承诺按照招标文件规定的商务条件提供对应服务。"
    elif form_key == "technical_deviation":
        headers = ["序号", "偏差事项", "招标文件要求", "投标文件响应", "偏差说明"]
        declaration = "投标人声明：针对本招标标的，除本表已列明偏差外，我们接受招标文件规定的其余全部技术条件，并承诺按照招标文件规定的技术条件提供对应服务。"
    else:
        headers = ["本企业人员姓名", "性别", "身份证号", "职务", "任职时间", "与国网人员关系", "国网人员姓名", "性别", "身份证号", "（曾）任职单位", "职务", "任职状态", "离职/退休时间"]
        declaration = "本表的“存在/不存在相关情形”尚未由客户确认，系统不自动代选。"
    blank_rows = 1 if source_rows else 0
    return "\n".join([
        f"### {title}",
        "",
        *_markdown_table(headers, [[""] * len(headers) for _ in range(blank_rows)]),
        "",
        declaration,
        "",
        "当前表单仍需客户确认后方可定稿。",
    ])


def _selected_source_tables(manifest: dict[str, Any]) -> tuple[Path, list[int]]:
    """返回固定表单对应的源 DOCX 及 1-based 表序号。"""
    form_key = str(manifest.get("form_key") or "")
    if form_key == "technical_characteristics":
        rows = [row for row in manifest.get("response_rows") or [] if isinstance(row, dict)]
        source_files = {_clean(row.get("source_file")) for row in rows if _clean(row.get("source_file"))}
        if len(source_files) != 1:
            raise ValueError("技术特性参数表必须且只能对应一个当次招标专项技术规范源文件。")
        table_indexes = sorted({int(row["source_table_index"]) for row in rows if row.get("source_table_index")})
        if not table_indexes:
            raise ValueError("技术特性参数表没有可追溯的源表序号。")
        return Path(next(iter(source_files))).resolve(), table_indexes

    source_form = manifest.get("source_form") if isinstance(manifest.get("source_form"), dict) else {}
    source_file = _clean(source_form.get("source_file"))
    table_index = source_form.get("table_index")
    if not source_file or not table_index:
        raise ValueError(f"{FORM_TITLES.get(form_key, '固定表单')}缺少源文件或源表序号。")
    return Path(source_file).resolve(), [int(table_index)]


def _clear_header_footer(document: Any) -> None:
    """移除招标文件原页码等页眉页脚，避免把招标方版记带入投标文件草稿。"""
    for section in document.sections:
        for part in (section.header, section.footer):
            element = part._element
            for child in list(element):
                element.remove(child)


def export_fixed_form_manifest_to_docx(
    manifest: dict[str, Any],
    output_path: str | Path,
    *,
    document_title: str | None = None,
) -> tuple[Path, dict[str, Any]]:
    """从当次招标源 DOCX 原样抽取固定表单，生成可审阅 Word 草稿。

    该路径不经过 Markdown 表格重建。表格节点直接从源 OOXML 深拷贝，因而保留
    列宽、横纵合并、行高、边框和声明行。技术参数只允许写入 manifest 明确标记
    为 ``formal_value_allowed`` 的保证值；其余单元格强制保持空白。
    """
    if not can_render_fixed_form(manifest):
        raise ValueError("不是受支持的固定表单 manifest。")

    source_path, table_indexes = _selected_source_tables(manifest)
    if not source_path.exists():
        raise FileNotFoundError(f"固定表单源文件不存在：{source_path}")
    source_document = Document(str(source_path))
    missing = [index for index in table_indexes if index < 1 or index > len(source_document.tables)]
    if missing:
        raise ValueError(f"固定表单源表序号无效：{missing}")

    source_tables = [source_document.tables[index - 1] for index in table_indexes]
    source_hashes = [_xml_sha(table._tbl) for table in source_tables]
    copied_tables = [deepcopy(table._tbl) for table in source_tables]

    body = source_document.element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)
    _clear_header_footer(source_document)

    title = document_title or str(manifest.get("title") or FORM_TITLES.get(str(manifest.get("form_key") or "")) or "固定表单")
    heading = source_document.add_paragraph()
    if "Title" in [style.name for style in source_document.styles]:
        heading.style = "Title"
    heading.add_run(_clean(title))
    warning = source_document.add_paragraph()
    warning.add_run("审阅状态：本表尚未通过正式投标门禁，未确认字段保持空白。")
    for position, table_xml in enumerate(copied_tables):
        body.insert(len(body) - 1, table_xml)
        # OOXML 中相邻的两个 w:tbl 会被 LibreOffice/Word 视作同一张表；必须用
        # 空段落隔开，才能在字段刷新后仍保持专项规范原有的两张独立参数表。
        if position < len(copied_tables) - 1:
            body.insert(len(body) - 1, OxmlElement("w:p"))

    guaranteed_values = {
        (int(row.get("source_table_index") or 0), int(row.get("source_row_number") or 0)): str(row.get("bidder_guaranteed_value") or "")
        for row in manifest.get("response_rows") or []
        if isinstance(row, dict) and row.get("formal_value_allowed") is True
    }
    filled_count = 0
    if str(manifest.get("form_key") or "") == "technical_characteristics":
        for table_index, output_table in zip(table_indexes, source_document.tables, strict=True):
            for row_number, row in enumerate(output_table.rows[1:], 2):
                if len(row.cells) < 5:
                    continue
                value = guaranteed_values.get((table_index, row_number), "")
                if value:
                    row.cells[4].text = value
                    filled_count += 1
                elif row.cells[4].text.strip():
                    row.cells[4].text = ""

    output = Path(output_path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    source_document.save(str(output))
    audited_document = Document(str(output))
    output_hashes = [_xml_sha(table._tbl) for table in audited_document.tables]
    report = {
        "renderer": "native_docx_ooxml_clone",
        "template": {
            "template_id": "current_tender_fixed_form",
            "template_family": "source_docx_ooxml",
            "reference_path": str(source_path),
        },
        "form_key": manifest.get("form_key"),
        "form_title": manifest.get("title"),
        "source_file": str(source_path),
        "source_file_sha256": _sha256(source_path),
        "source_table_indexes": table_indexes,
        "source_table_xml_sha256": source_hashes,
        "output_table_xml_sha256": output_hashes,
        "table_count": len(audited_document.tables),
        "table_shapes": [
            {"rows": len(table.rows), "columns": len(table.columns)}
            for table in audited_document.tables
        ],
        "exact_table_xml_preserved": source_hashes == output_hashes,
        "guarantee_value_filled_count": filled_count,
        "formal_ready": bool(manifest.get("formal_ready")),
        "blocker_count": len(manifest.get("blockers") or []),
        "model_invoked": False,
        "mineru_invoked": False,
    }
    return output, report


def _table_merge_signature(table: Any) -> list[list[tuple[int, str | None]]]:
    signature: list[list[tuple[int, str | None]]] = []
    for row in table.rows:
        cells: list[tuple[int, str | None]] = []
        seen: set[int] = set()
        for cell in row.cells:
            element_id = id(cell._tc)
            if element_id in seen:
                continue
            seen.add(element_id)
            tc_pr = cell._tc.tcPr
            grid_span = int(tc_pr.gridSpan.val) if tc_pr is not None and tc_pr.gridSpan is not None else 1
            vertical_merge = None
            if tc_pr is not None and tc_pr.vMerge is not None:
                vertical_merge = str(tc_pr.vMerge.val or "continue")
            cells.append((grid_span, vertical_merge))
        signature.append(cells)
    return signature


def _table_grid_widths(table: Any) -> list[int | None]:
    grid = table._tbl.tblGrid
    if grid is None:
        return []
    return [int(value) if (value := column.get(qn("w:w"))) is not None else None for column in grid.gridCol_lst]


def audit_fixed_form_docx(
    manifest: dict[str, Any],
    output_path: str | Path,
    *,
    width_tolerance_twips: int = 1,
) -> dict[str, Any]:
    """审计 LibreOffice 字段刷新后的固定表单，防止表合并或结构漂移。"""
    source_path, table_indexes = _selected_source_tables(manifest)
    source_document = Document(str(source_path))
    output = Path(output_path).resolve()
    output_document = Document(str(output))
    source_tables = [source_document.tables[index - 1] for index in table_indexes]
    output_tables = list(output_document.tables)

    same_table_count = len(source_tables) == len(output_tables)
    same_shapes = same_table_count and all(
        (len(source.rows), len(source.columns)) == (len(rendered.rows), len(rendered.columns))
        for source, rendered in zip(source_tables, output_tables, strict=True)
    )
    merge_preserved = same_table_count and all(
        _table_merge_signature(source) == _table_merge_signature(rendered)
        for source, rendered in zip(source_tables, output_tables, strict=True)
    )
    max_width_delta = 0
    column_widths_preserved = same_table_count
    if same_table_count:
        for source, rendered in zip(source_tables, output_tables, strict=True):
            source_widths = _table_grid_widths(source)
            output_widths = _table_grid_widths(rendered)
            if len(source_widths) != len(output_widths):
                column_widths_preserved = False
                continue
            for source_width, output_width in zip(source_widths, output_widths, strict=True):
                if source_width is None or output_width is None:
                    if source_width != output_width:
                        column_widths_preserved = False
                    continue
                delta = abs(source_width - output_width)
                max_width_delta = max(max_width_delta, delta)
                if delta > width_tolerance_twips:
                    column_widths_preserved = False

    form_key = str(manifest.get("form_key") or "")
    source_text_rows: list[list[str]] = []
    output_text_rows: list[list[str]] = []
    for source in source_tables:
        source_text_rows.extend([[_clean(cell.text) for cell in row.cells] for row in source.rows])
    for rendered in output_tables:
        output_text_rows.extend([[_clean(cell.text) for cell in row.cells] for row in rendered.rows])

    expected_guarantees: list[str] = []
    actual_guarantees: list[str] = []
    if form_key == "technical_characteristics":
        expected_guarantees = [str(row.get("bidder_guaranteed_value") or "") for row in manifest.get("response_rows") or []]
        actual_guarantees = [row[4] for row in output_text_rows if len(row) >= 5][1:]
        # 两张参数表各有表头，按表逐张提取，避免把第二张表头误当参数行。
        actual_guarantees = [
            _clean(row.cells[4].text)
            for table in output_tables
            for row in table.rows[1:]
            if len(row.cells) >= 5
        ]
        for rows in (source_text_rows, output_text_rows):
            for row in rows:
                if len(row) >= 5:
                    row[4] = ""

    source_text_preserved = source_text_rows == output_text_rows
    guarantees_preserved = expected_guarantees == actual_guarantees if form_key == "technical_characteristics" else True
    passed = all((same_table_count, same_shapes, merge_preserved, column_widths_preserved, source_text_preserved, guarantees_preserved))
    return {
        "status": "passed" if passed else "failed",
        "passed": passed,
        "output_path": str(output),
        "expected_table_count": len(source_tables),
        "actual_table_count": len(output_tables),
        "table_shapes": [{"rows": len(table.rows), "columns": len(table.columns)} for table in output_tables],
        "same_shapes": same_shapes,
        "merge_structure_preserved": merge_preserved,
        "column_widths_preserved": column_widths_preserved,
        "width_tolerance_twips": width_tolerance_twips,
        "max_width_delta_twips": max_width_delta,
        "source_cell_text_preserved": source_text_preserved,
        "guarantee_values_preserved": guarantees_preserved,
        "guarantee_value_filled_count": sum(bool(value) for value in actual_guarantees),
    }


def load_parameter_rows(path: str | Path, *, product_family: str = "MPP电缆保护管") -> list[dict[str, Any]]:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    rows = data if isinstance(data, list) else data.get("rows") or []
    return [row for row in rows if isinstance(row, dict) and row.get("product_family") == product_family]
