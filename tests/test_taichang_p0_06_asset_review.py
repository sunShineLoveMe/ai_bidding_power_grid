import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from scripts.rag.prepare_taichang_p0_06_asset_review import (  # noqa: E402
    apply_review_decisions,
    build_review_package,
    read_workbook_decisions,
    write_workbook,
)


DATA_DIR = ROOT_DIR / "docs/development/taichang-bid-v1-data"


class TaichangP006AssetReviewTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        matrix = json.loads((DATA_DIR / "asset_classification_matrix.json").read_text(encoding="utf-8"))
        queue = json.loads((DATA_DIR / "asset_review_queue.json").read_text(encoding="utf-8"))
        cls.package = build_review_package(matrix, queue)

    def test_review_package_covers_all_candidates_without_overlap(self):
        metadata = self.package["metadata"]
        self.assertEqual(metadata["candidate_count"], 896)
        self.assertEqual(metadata["source_authorized_count"], 896)
        self.assertEqual(metadata["policy_auto_accepted_count"], 169)
        self.assertEqual(metadata["review_candidate_count"], 290)
        self.assertEqual(metadata["review_group_count"], 60)
        self.assertEqual(metadata["auto_blocked_or_reference_count"], 437)
        self.assertEqual(metadata["auto_linked_existing_count"], 272)
        review_ids = {row["candidate_id"] for row in self.package["review_candidates"]}
        auto_ids = {row["candidate_id"] for row in self.package["policy_auto_accepted"]}
        blocked_ids = {row["candidate_id"] for row in self.package["auto_blocked_or_reference"]}
        self.assertFalse(review_ids & blocked_ids)
        self.assertFalse(review_ids & auto_ids)
        self.assertFalse(auto_ids & blocked_ids)
        self.assertEqual(len(review_ids | auto_ids | blocked_ids), 896)

    def test_default_package_auto_accepts_only_low_risk_knowledge(self):
        self.assertEqual(len(self.package["approved_decisions"]), 169)
        self.assertEqual(len(self.package["ready_for_ingestion"]), 169)
        self.assertTrue(all(row["review_status_after_review"] == "policy_auto_accepted"
                            for row in self.package["policy_auto_accepted"]))
        self.assertTrue(all(row["quality_tier_after_review"] == "knowledge_only"
                            for row in self.package["policy_auto_accepted"]))
        self.assertTrue(all(not row["allowed_for_bid"] and not row["formal_bid_ready"]
                            for row in self.package["policy_auto_accepted"]))
        self.assertTrue(all(row["source_domain"] == "enterprise_fact"
                            for row in self.package["policy_auto_accepted"]))
        self.assertTrue(all(not row["decision"] for row in self.package["review_candidates"]))

    def test_groups_reduce_page_level_work_without_losing_candidates(self):
        self.assertLess(len(self.package["groups"]), len(self.package["review_candidates"]))
        self.assertEqual(sum(group["candidate_count"] for group in self.package["groups"]), 290)
        group_ids = {group["review_group_id"] for group in self.package["groups"]}
        self.assertTrue(all(row["review_group_id"] in group_ids for row in self.package["review_candidates"]))

    def test_known_expired_ohs_certificate_is_forced_to_defer(self):
        rows = [
            row for row in self.package["review_candidates"]
            if "职业健康安全管理体系认证证书" in row["source_section"]
        ]
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(row["machine_recommendation"] == "延后补证" for row in rows))
        self.assertTrue(all("2026年6月18日到期" in row["required_check"] for row in rows))

    def test_incomplete_approval_is_rejected(self):
        candidate = next(
            row for row in self.package["review_candidates"]
            if row["dedup_status"] == "new" and row["sensitivity"] == "internal"
        )
        decisions = {row["candidate_id"]: {} for row in self.package["review_candidates"]}
        decisions[candidate["candidate_id"]] = {"decision": "批准为知识资料"}
        result = apply_review_decisions(self.package, decisions)
        self.assertEqual(len(result["approved_decisions"]), 169)
        self.assertEqual(len(result["ready_for_ingestion"]), 169)
        self.assertEqual(result["metadata"]["validation_error_count"], 1)

    def test_valid_simulated_knowledge_approval_can_pass_strict_gate(self):
        candidate = next(
            row for row in self.package["review_candidates"]
            if row["dedup_status"] == "new"
            and row["sensitivity"] == "internal"
            and row["evidence_type"] not in {
                "business_license", "certification", "finance", "personnel_certificate", "project_performance",
                "authorization", "inspection_report", "testing_capacity", "product_parameter_table",
                "production_capacity", "product_image",
            }
        )
        decision = {
            "decision": "批准为知识资料",
            "dedup_confirmation": "确认新资产",
            "reviewed_title": candidate["title"],
            "reviewed_target_library_label": candidate["target_library_label"],
            "original_evidence_verified": "是",
            "reviewer": "测试复核人",
            "reviewed_at": "2026-07-13",
        }
        decisions = {row["candidate_id"]: {} for row in self.package["review_candidates"]}
        decisions[candidate["candidate_id"]] = decision
        result = apply_review_decisions(self.package, decisions)
        self.assertEqual(result["metadata"]["validation_error_count"], 0)
        self.assertEqual(result["metadata"]["approved_decision_count"], 170)
        self.assertEqual(result["metadata"]["human_approved_decision_count"], 1)
        self.assertEqual(result["metadata"]["ready_for_ingestion_count"], 170)
        reviewed = next(row for row in result["ready_for_ingestion"] if row["candidate_id"] == candidate["candidate_id"])
        self.assertEqual(reviewed["quality_tier_after_review"], "knowledge_only")

    def test_parameter_candidate_requires_structured_extraction_not_direct_ingestion(self):
        candidate = next(
            row for row in self.package["review_candidates"]
            if row["dedup_status"] == "new"
            and row["evidence_type"] == "product_parameter_table"
            and row["product_families"]
        )
        decision = {
            "decision": "批准为知识资料",
            "dedup_confirmation": "确认新资产",
            "reviewed_title": candidate["title"],
            "reviewed_target_library_label": candidate["target_library_label"],
            "product_boundary_verified": "是",
            "original_evidence_verified": "是",
            "reviewer": "测试复核人",
            "reviewed_at": "2026-07-13",
        }
        decisions = {row["candidate_id"]: {} for row in self.package["review_candidates"]}
        decisions[candidate["candidate_id"]] = decision
        result = apply_review_decisions(self.package, decisions)
        self.assertEqual(result["metadata"]["validation_error_count"], 0)
        reviewed = next(row for row in result["approved_decisions"] if row["candidate_id"] == candidate["candidate_id"])
        self.assertEqual(reviewed["ingestion_action"], "structured_extraction_required")
        self.assertFalse(reviewed["ready_for_ingestion"])
        self.assertFalse(any(row["candidate_id"] == candidate["candidate_id"] for row in result["ready_for_ingestion"]))

    def test_sensitive_candidate_cannot_be_approved_by_spreadsheet_fields(self):
        candidate = next(row for row in self.package["review_candidates"] if row["sensitivity"] in {"sensitive", "restricted"})
        decision = {
            "decision": "批准为知识资料",
            "dedup_confirmation": "确认新资产",
            "reviewed_title": candidate["title"],
            "reviewed_target_library_label": candidate["target_library_label"],
            "original_evidence_verified": "是",
            "timeliness_verified": "是",
            "sensitive_authorized": "是",
            "reviewer": "测试复核人",
            "reviewed_at": "2026-07-13",
        }
        decisions = {row["candidate_id"]: {} for row in self.package["review_candidates"]}
        decisions[candidate["candidate_id"]] = decision
        result = apply_review_decisions(self.package, decisions)
        self.assertEqual(len(result["approved_decisions"]), 169)
        self.assertFalse(any(row["candidate_id"] == candidate["candidate_id"] for row in result["approved_decisions"]))
        self.assertTrue(any("受限或敏感资料" in error for error in result["validation_errors"][0]["errors"]))

    def test_tampered_candidate_set_blocks_all_approvals(self):
        candidate = next(
            row for row in self.package["review_candidates"]
            if row["dedup_status"] == "new"
            and row["sensitivity"] == "internal"
            and row["evidence_type"] not in {
                "business_license", "certification", "finance", "personnel_certificate", "project_performance",
                "authorization", "inspection_report", "testing_capacity", "product_parameter_table",
                "production_capacity", "product_image",
            }
        )
        decisions = {row["candidate_id"]: {} for row in self.package["review_candidates"]}
        decisions.pop(self.package["review_candidates"][-1]["candidate_id"])
        decisions[candidate["candidate_id"]] = {
            "decision": "批准为知识资料",
            "dedup_confirmation": "确认新资产",
            "reviewed_title": candidate["title"],
            "reviewed_target_library_label": candidate["target_library_label"],
            "original_evidence_verified": "是",
            "reviewer": "测试复核人",
            "reviewed_at": "2026-07-13",
        }
        result = apply_review_decisions(self.package, decisions)
        self.assertEqual(len(result["approved_decisions"]), 169)
        self.assertEqual(len(result["ready_for_ingestion"]), 169)
        self.assertEqual(result["metadata"]["human_approved_decision_count"], 0)
        self.assertTrue(any(error["candidate_id"] == "审批表完整性" for error in result["validation_errors"]))

    def test_workbook_has_required_sheets_validations_and_roundtrip(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "审批表.xlsx"
            write_workbook(path, self.package)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["审批说明", "证据分组", "系统自动接收", "异常候选复核", "自动阻断与仅参考"])
            self.assertEqual(wb["证据分组"].max_row, 61)
            self.assertEqual(wb["系统自动接收"].max_row, 170)
            self.assertEqual(wb["异常候选复核"].max_row, 291)
            self.assertEqual(wb["自动阻断与仅参考"].max_row, 438)
            self.assertGreaterEqual(len(wb["异常候选复核"].data_validations.dataValidation), 8)
            ws = wb["异常候选复核"]
            headers = {cell.value: cell.column for cell in ws[1]}
            ws.cell(2, headers["审批结论"], "延后补证")
            ws.cell(2, headers["复核人"], "资料管理员")
            wb.save(path)
            decisions = read_workbook_decisions(path)
            first_candidate = self.package["review_candidates"][0]["candidate_id"]
            self.assertEqual(decisions[first_candidate]["decision"], "延后补证")
            self.assertEqual(decisions[first_candidate]["reviewer"], "资料管理员")


if __name__ == "__main__":
    unittest.main()
