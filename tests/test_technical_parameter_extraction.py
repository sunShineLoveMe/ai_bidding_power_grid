import tempfile
import unittest
from pathlib import Path

from docx import Document


class TechnicalParameterExtractionTest(unittest.TestCase):
    def _docx_with_tables(self) -> Path:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "technical.docx"
        doc = Document()

        table = doc.add_table(rows=3, cols=4)
        for index, value in enumerate(["公称内径 （mm）", "内径允许偏差 （mm）", "最小壁厚 （mm）", "壁厚允许偏差 （mm）"]):
            table.rows[0].cells[index].text = value
        for index, value in enumerate(["200（内径）", "正公差 +1.0 负公差 0", "16", "正公差 +1.4 负公差 0"]):
            table.rows[1].cells[index].text = value
        for index, value in enumerate(["175（内径）", "正公差 +0.9 负公差 0", "14", "正公差 +1.4 负公差 0"]):
            table.rows[2].cells[index].text = value

        table = doc.add_table(rows=3, cols=3)
        for index, value in enumerate(["序号", "项目", "指标"]):
            table.rows[0].cells[index].text = value
        for index, value in enumerate(["1", "密度g/cm3", "0.90～0.94"]):
            table.rows[1].cells[index].text = value
        for index, value in enumerate(["2", "拉伸强度", "≥25MPa"]):
            table.rows[2].cells[index].text = value

        table = doc.add_table(rows=2, cols=5)
        for index, value in enumerate(["序号", "技术参数名称", "单位", "项目需求值或表述", "投标人响应值"]):
            table.rows[0].cells[index].text = value
        for index, value in enumerate(["1", "电缆保护管-材质:MPP,标称直径:φ200", "米", "参照通用部分5.技术要求", "满足"]):
            table.rows[1].cells[index].text = value

        doc.save(path)
        return path

    def test_extracts_dimension_performance_and_response_tables(self):
        from scripts.rag.extract_customer_technical_parameters import _extract_tables_from_docx

        tables = _extract_tables_from_docx(self._docx_with_tables())
        rows = [row for table in tables for row in table.rows]
        table_types = {table.table_type for table in tables}

        self.assertIn("dimension_parameter_table", table_types)
        self.assertIn("performance_parameter_table", table_types)
        self.assertIn("bid_response_parameter_table", table_types)
        self.assertTrue(any(row.get("parameter_name") == "公称内径 200（内径）" for row in rows))
        self.assertTrue(any(row.get("standard_value") == "≥25MPa" for row in rows))
        self.assertTrue(any(row.get("project_required_value") == "参照通用部分5.技术要求" for row in rows))
        self.assertTrue(any(row.get("bidder_response_value") == "满足" for row in rows))


if __name__ == "__main__":
    unittest.main()
