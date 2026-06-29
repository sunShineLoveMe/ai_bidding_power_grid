import subprocess
import tempfile
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from flask import Flask

from backend.api.routes import (
    build_project_bid_markdown,
    _asset_allowed_for_bid,
    _demote_body_markdown_headings,
    _normalize_body_outline_lines,
    _numbered_export_sections,
    _renumber_body_markdown_headings,
    _strip_export_guidance_blocks,
    _strip_duplicate_section_heading,
    _strip_untrusted_export_images,
)
from backend.export.md_to_word import (
    DOCX_BIDDER_FULL_NAME,
    DOCX_BODY_EAST_ASIA,
    DOCX_BODY_FIRST_LINE_INDENT_PT,
    DOCX_IMAGE_FRAME_HEIGHT_IN,
    DOCX_IMAGE_FRAME_WIDTH_IN,
    DOCX_BODY_LINE_SPACING,
    DOCX_BODY_LINE_SPACING_RULE,
    DOCX_TOC_ENTRY_FONT_SIZE,
    DOCX_TOC_ENTRY_LINE_SPACING,
    DOCX_TOC_PAGE_NUMBER_FONT_SIZE,
    DOCX_IMAGE_CAPTION_FONT_SIZE,
    DOCX_TABLE_EAST_ASIA,
    DOCX_LIST_HANGING_INDENT_PT,
    DOCX_LIST_LEFT_INDENT_PT,
    convert_md_to_word,
    extract_bid_cover_fields,
    refresh_docx_fields_with_soffice,
    should_start_heading_on_new_page,
    taichang_bid_document_title,
)
from backend.parsing.tender_metadata import extract_tender_project_metadata
from backend.services.formal_asset_naming import clean_formal_asset_title


class DocxExportRegressionTest(unittest.TestCase):
    def test_body_markdown_headings_are_not_exported_as_word_outline_headings(self):
        content = "\n".join(
            [
                "## 噪声与振动控制措施",
                "",
                "### 主要控制措施",
                "",
                "正文内容。",
                "",
                "监测与记录",
                "---",
                "",
                "| 项目 | 内容 |",
                "| --- | --- |",
                "| 监测 | 按要求执行 |",
            ]
        )

        normalized = _demote_body_markdown_headings(content)

        self.assertNotIn("## 噪声与振动控制措施", normalized)
        self.assertNotIn("### 主要控制措施", normalized)
        self.assertIn("<!-- BID_BODY_SUBHEADING: 噪声与振动控制措施 -->", normalized)
        self.assertIn("<!-- BID_BODY_SUBHEADING: 主要控制措施 -->", normalized)
        self.assertIn("<!-- BID_BODY_SUBHEADING: 监测与记录 -->", normalized)
        self.assertIn("| 项目 | 内容 |", normalized)

    def test_stale_bracket_body_headings_are_renumbered_for_formal_export(self):
        content = "\n".join(
            [
                "【5.1 概述】",
                "正文。",
                "【5.2 技术方案与产品性能响应】",
                "【5.2.1 产品执行标准与技术要求】",
                "5.2.2 产品关键技术参数",
                "1. 投标意愿与范围：这仍然是正文列表。",
            ]
        )

        normalized = _normalize_body_outline_lines(content, {"_export_order": "1"})

        self.assertIn("<!-- BID_BODY_SUBHEADING: 1.1 概述 -->", normalized)
        self.assertIn("<!-- BID_BODY_SUBHEADING: 1.2 技术方案与产品性能响应 -->", normalized)
        self.assertIn("<!-- BID_BODY_SUBHEADING: 1.2.1 产品执行标准与技术要求 -->", normalized)
        self.assertIn("<!-- BID_BODY_SUBHEADING: 1.2.2 产品关键技术参数 -->", normalized)
        self.assertIn("1. 投标意愿与范围：这仍然是正文列表。", normalized)
        self.assertNotIn("【", normalized)
        self.assertNotIn("】", normalized)

    def test_export_guidance_blocks_are_removed_from_formal_docx_content(self):
        content = "\n".join(
            [
                "正式正文。",
                "【编写要点】",
                "（需结合招标文件条款逐项响应，避免遗漏实质性要求。）",
                "【需准备资料】",
                "需人工补充企业资料、资信文件和证明材料",
                "【风险与复核】",
                "暂无明确风险，仍需结合招标文件复核。",
                "后续正文。",
            ]
        )

        cleaned = _strip_export_guidance_blocks(content)

        self.assertIn("正式正文。", cleaned)
        self.assertIn("后续正文。", cleaned)
        self.assertNotIn("编写要点", cleaned)
        self.assertNotIn("需准备资料", cleaned)
        self.assertNotIn("风险与复核", cleaned)

    def test_docx_navigation_headings_remain_official_section_headings_only(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "outline-lock.md"
            body = _demote_body_markdown_headings(
                "\n".join(
                    [
                        "## 噪声与振动控制措施",
                        "",
                        "### 主要控制措施",
                        "",
                        "正文内容。",
                    ]
                )
            )
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 测试投标文件",
                        "",
                        "## 1. 企业营业执照",
                        "",
                        body,
                        "",
                        "## 2. 安全生产许可证",
                        "",
                        "正文内容。",
                    ]
                ),
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            document = Document(str(output_path))
            headings = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.style and paragraph.style.name.startswith("Heading")
            ]

            self.assertIn("1. 企业营业执照", headings)
            self.assertIn("2. 安全生产许可证", headings)
            self.assertNotIn("噪声与振动控制措施", headings)
            self.assertNotIn("主要控制措施", headings)

    def test_duplicate_numbered_body_heading_is_removed_before_docx_export(self):
        content = "## 1. 企业营业执照\n\n正文内容。"
        cleaned = _strip_duplicate_section_heading(content, {"title": "企业营业执照", "level": 1})

        self.assertEqual(cleaned, "正文内容。")

    def test_body_heading_numbers_are_rewritten_under_export_section_order(self):
        content = "\n".join([
            "# 4. 投标保证金",
            "",
            "## 4.1 基本情况说明",
            "",
            "### 4.1.1 保证金提交形式",
        ])

        normalized = _renumber_body_markdown_headings(content, {"_export_order": "1.4"})

        self.assertIn("# 1.4.1 投标保证金", normalized)
        self.assertIn("## 1.4.1.1 基本情况说明", normalized)
        self.assertIn("### 1.4.1.1.1 保证金提交形式", normalized)
        self.assertNotIn("4. 投标保证金", normalized)
        self.assertNotIn("4.1 基本情况说明", normalized)

    def test_body_heading_number_rewrite_never_emits_zero_segments(self):
        content = "##### 4. 投标保证金"

        normalized = _renumber_body_markdown_headings(content, {"_export_order": "1.3"})

        self.assertIn("1.3.1", normalized)
        self.assertNotIn(".0", normalized)

    def test_body_heading_rewrite_removes_duplicate_section_title_headings(self):
        content = "\n".join([
            "## 产品购销合同",
            "",
            "# 产品购销合同",
            "",
            "## 1. 合同标的与供货范围",
            "",
            "### 1.1 供货产品清单",
        ])

        normalized = _renumber_body_markdown_headings(content, {
            "_export_order": "4.3.3.1",
            "title": "产品购销合同",
        })

        self.assertNotIn("4.3.3.1.1 产品购销合同", normalized)
        self.assertNotIn("产品购销合同\n", normalized)
        self.assertIn("## 4.3.3.1.1 合同标的与供货范围", normalized)
        self.assertIn("### 4.3.3.1.1.1 供货产品清单", normalized)

    def test_body_heading_rewrite_uses_relative_markdown_depth_for_unnumbered_headings(self):
        content = "\n".join([
            "## 资质业绩凭证 单",
            "",
            "### 资质业绩凭证单",
            "",
            "#### 一、本章说明",
            "",
            "#### 二、资格预审基本情况",
        ])

        normalized = _renumber_body_markdown_headings(content, {
            "_export_order": "2.3.2",
            "title": "资质业绩凭证 单",
        })

        self.assertNotIn("2.3.2.1 资质业绩凭证单", normalized)
        self.assertIn("#### 2.3.2.1 本章说明", normalized)
        self.assertIn("#### 2.3.2.2 资格预审基本情况", normalized)
        self.assertNotIn("2.3.2.1.1.1.1", normalized)

    def test_body_heading_rewrite_uses_relative_explicit_number_depth(self):
        content = "\n".join([
            "## 资格预审结果通知书",
            "",
            "### 1. 资格预审结果通知书",
            "",
            "#### 1.1 总则",
            "",
            "#### 1.2 资格预审结果确认",
        ])

        normalized = _renumber_body_markdown_headings(content, {
            "_export_order": "2.11.1.1",
            "title": "资格预审结果通知书",
        })

        self.assertIn("#### 2.11.1.1.1 总则", normalized)
        self.assertIn("#### 2.11.1.1.2 资格预审结果确认", normalized)
        self.assertNotIn("2.11.1.1.1.1", normalized)

    def test_body_heading_rewrite_skips_generic_volume_headings(self):
        content = "\n".join([
            "## 保险购买凭证-电缆保护管 CPVC",
            "",
            "### 第五章 商务文件",
            "",
            "#### 5.1 投标函",
            "",
            "#### 5.2 合同条款响应与商务偏离表",
        ])

        normalized = _renumber_body_markdown_headings(content, {
            "_export_order": "2.7.1.1",
            "title": "保险购买凭证-电缆保护管 CPVC",
        })

        self.assertNotIn("商务文件", normalized)
        self.assertIn("#### 2.7.1.1.1 投标函", normalized)
        self.assertIn("#### 2.7.1.1.2 合同条款响应与商务偏离表", normalized)
        self.assertNotIn("2.7.1.1.1.1", normalized)

    def test_export_sections_are_numbered_for_word_outline(self):
        sections = [
            {"id": "a", "level": 1, "title": "企业营业执照"},
            {"id": "b", "level": 1, "title": "2. 安全生产许可证"},
            {"id": "c", "level": 1, "title": "项目经理资格"},
            {"id": "d", "level": 2, "title": "项目经理简历表"},
        ]

        numbered = _numbered_export_sections(sections)

        self.assertEqual(numbered[0]["_export_title"], "1. 企业营业执照")
        self.assertEqual(numbered[1]["_export_title"], "2. 安全生产许可证")
        self.assertEqual(numbered[2]["_export_title"], "3. 项目经理资格")
        self.assertEqual(numbered[3]["_export_title"], "3.1 项目经理简历表")

    def test_export_section_numbering_collapses_skipped_levels(self):
        sections = [
            {"id": "a", "level": 1, "title": "投标人不得存在情形的声明"},
            {"id": "b", "level": 3, "title": "质量控制措施"},
            {"id": "c", "level": 3, "title": "环保与文明施工措施"},
        ]

        numbered = _numbered_export_sections(sections)

        self.assertEqual(numbered[0]["_export_title"], "1. 投标人不得存在情形的声明")
        self.assertEqual(numbered[1]["_export_title"], "1.1 质量控制措施")
        self.assertEqual(numbered[2]["_export_title"], "1.2 环保与文明施工措施")

    def test_export_section_numbering_collapses_leading_child_level_without_zero_prefix(self):
        sections = [
            {"id": "a", "level": 2, "title": "投标函及投标函附录"},
            {"id": "b", "level": 3, "title": "编制依据"},
            {"id": "c", "level": 1, "title": "资格审查资料"},
        ]

        numbered = _numbered_export_sections(sections)

        self.assertEqual(numbered[0]["_export_title"], "1. 投标函及投标函附录")
        self.assertEqual(numbered[1]["_export_title"], "1.1 编制依据")
        self.assertEqual(numbered[2]["_export_title"], "2. 资格审查资料")

    def test_export_section_titles_strip_repeated_parent_prefixes(self):
        sections = [
            {"id": "a", "level": 1, "title": "资格审查资料"},
            {"id": "b", "level": 2, "title": "企业基本资格资料"},
            {"id": "c", "level": 3, "title": "企业基本资格资料 - 响应要求"},
            {"id": "d", "level": 3, "title": "企业基本资格资料 - 资料清单"},
        ]

        numbered = _numbered_export_sections(sections)

        self.assertEqual(numbered[1]["_export_title"], "1.1 企业基本资格资料")
        self.assertEqual(numbered[2]["_export_title"], "1.1.1 响应要求")
        self.assertEqual(numbered[3]["_export_title"], "1.1.2 资料清单")

    def test_bid_markdown_export_prefers_editor_snapshot_over_stale_database_sections(self):
        project_id = "11111111-1111-1111-1111-111111111111"
        with tempfile.TemporaryDirectory() as tmpdir:
            app = Flask(__name__)
            app.config["GENERATED_FOLDER"] = tmpdir
            snapshot = [
                {
                    "id": "local-1",
                    "order_index": 1,
                    "level": 1,
                    "title": "投标函及格式文件",
                    "content": "在线工作台正文一。",
                },
                {
                    "id": "local-2",
                    "order_index": 2,
                    "level": 2,
                    "title": "投标函及投标函附录",
                    "content": "在线工作台正文二。",
                },
            ]
            stale_sections = [
                {
                    "id": "db-1",
                    "order_index": 1,
                    "level": 1,
                    "title": "资格审查资料封面及目录",
                    "content": "",
                }
            ]

            with (
                app.app_context(),
                patch("backend.api.routes.get_project_interpretation", return_value={
                    "project": {"id": project_id, "project_name": "测试项目"},
                    "analysis": {"project_meta": {"project_name": "测试投标文件"}},
                }),
                patch("backend.api.routes.list_bid_sections", return_value=stale_sections),
            ):
                markdown_path, _, _ = build_project_bid_markdown(project_id, sections_snapshot=snapshot)

            markdown = markdown_path.read_text(encoding="utf-8")
            self.assertIn("# 1. 投标函及格式文件", markdown)
            self.assertIn("## 1.1 投标函及投标函附录", markdown)
            self.assertIn("在线工作台正文一。", markdown)
            self.assertIn("在线工作台正文二。", markdown)
            self.assertNotIn("资格审查资料封面及目录", markdown)
            self.assertNotIn("待补充章节正文", markdown)

    def test_bid_markdown_export_uses_short_physical_path_for_long_project_name(self):
        project_id = "580b8c82-42c2-4a51-a0d2-17b60afa22b9"
        long_project_name = "国家电网有限公司2026年西北、西藏区域第一次联合采购10kV架空绝缘导线新疆技术标商务标超长项目名称用于回归路径长度问题"
        with tempfile.TemporaryDirectory() as tmpdir:
            app = Flask(__name__)
            app.config["GENERATED_FOLDER"] = tmpdir
            sections = [
                {
                    "id": "tech-1",
                    "order_index": 1,
                    "level": 1,
                    "title": "技术规范响应",
                    "content": "技术标正文。",
                    "metadata": {"volume_type": "technical"},
                },
                {
                    "id": "biz-1",
                    "order_index": 2,
                    "level": 1,
                    "title": "商务响应",
                    "content": "商务标正文。",
                    "metadata": {"volume_type": "business"},
                },
            ]
            cover_fields = {
                "项目名称": long_project_name,
                "招标编号": "SL265A",
                "包号": "包1",
                "包名称": "10kV架空绝缘导线-新疆",
            }

            with (
                app.app_context(),
                patch("backend.api.routes.get_project_interpretation", return_value={
                    "project": {"id": project_id, "project_name": long_project_name},
                    "analysis": {"project_meta": {"project_name": long_project_name, "cover_fields": cover_fields}},
                }),
                patch("backend.api.routes.list_bid_sections", return_value=sections),
            ):
                markdown_path, _, report = build_project_bid_markdown(project_id, volume_type="technical")

        self.assertEqual(markdown_path.parent.name, project_id[:8])
        self.assertLess(len(markdown_path.name), 80)
        self.assertTrue(markdown_path.name.startswith("泰昌_SL265A_包1_技术投标文件_"))
        self.assertNotIn(long_project_name, str(markdown_path))
        self.assertEqual(report["download_file_name"], markdown_path.with_suffix(".docx").name)
        self.assertEqual(report["output_naming"]["scheme"], "short_bidder_tender_package_volume_date.v1")

    def test_tender_metadata_extracts_cover_fields_from_uploaded_tender_text(self):
        markdown = "\n".join(
            [
                "# 国网辽宁电力2025年第三次物资协议库存招标采购招标文件",
                "",
                "| 字段 | 内容 |",
                "| --- | --- |",
                "| 招标编号 | 2225AC |",
                "| 分标编号 | 102-CPVC |",
                "| 分标名称 | 电缆保护管 |",
                "| 包号 | 包1 |",
                "| 包名称 | CPVC电缆保护管包1 |",
                "| 招标人 | 国网辽宁省电力有限公司 |",
                "| 招标代理机构 | 国网辽宁招标有限公司 |",
            ]
        )
        content_list = [{"text": "招标编号 2225AC 分标编号 102-CPVC", "page_idx": 0}]

        meta = extract_tender_project_metadata(markdown, content_list)

        self.assertEqual(meta["project_name"], "国网辽宁电力2025年第三次物资协议库存招标采购")
        self.assertEqual(meta["tender_no"], "2225AC")
        self.assertEqual(meta["project_no"], "2225AC")
        self.assertEqual(meta["tender_unit"], "国网辽宁省电力有限公司")
        self.assertEqual(meta["agency"], "国网辽宁招标有限公司")
        self.assertEqual(meta["cover_fields"]["招标编号"], "2225AC")
        self.assertEqual(meta["cover_fields"]["分标编号"], "102-CPVC")
        self.assertEqual(meta["cover_fields"]["包名称"], "CPVC电缆保护管包1")
        self.assertIn("招标编号", meta["cover_field_sources"])

    def test_tender_metadata_does_not_treat_next_label_as_empty_project_name_value(self):
        markdown = "\n".join(
            [
                "项目名称：",
                "招标编号：2225AC",
                "招标人：",
                "招标代理机构：国网辽宁招标有限公司",
            ]
        )

        meta = extract_tender_project_metadata(markdown, [])

        self.assertNotEqual(meta.get("project_name"), "招标编号：")
        self.assertNotIn("项目名称", meta["cover_fields"])
        self.assertNotEqual(meta.get("tender_unit"), "招标代理机构：")
        self.assertEqual(meta["cover_fields"]["招标编号"], "2225AC")

    def test_bid_markdown_report_carries_structured_cover_fields_from_project_meta(self):
        project_id = "11111111-1111-1111-1111-111111111111"
        with tempfile.TemporaryDirectory() as tmpdir:
            app = Flask(__name__)
            app.config["GENERATED_FOLDER"] = tmpdir
            sections = [
                {
                    "id": "section-1",
                    "order_index": 1,
                    "level": 1,
                    "title": "投标函",
                    "content": "正文内容。",
                }
            ]
            cover_fields = {
                "项目名称": "国网辽宁电力2025年第三次物资协议库存招标采购",
                "文件类型": "技术投标文件",
                "招标编号": "2225AC",
                "分标编号": "102-CPVC",
                "包号": "包1",
            }

            with (
                app.app_context(),
                patch("backend.api.routes.get_project_interpretation", return_value={
                    "project": {"id": project_id, "project_name": "上传文件名"},
                    "analysis": {"project_meta": {"project_name": "上传文件名", "cover_fields": cover_fields}},
                }),
                patch("backend.api.routes.list_bid_sections", return_value=sections),
            ):
                _, document_title, report = build_project_bid_markdown(project_id)

        self.assertEqual(document_title, "上传文件名投标文件")
        self.assertEqual(report["cover_field_source"], "uploaded_tender_structured_extract")
        self.assertEqual(report["cover_fields"]["招标编号"], "2225AC")
        self.assertEqual(report["cover_fields"]["分标编号"], "102-CPVC")

    def test_bid_markdown_volume_export_sets_delivery_file_type_and_scope(self):
        project_id = "11111111-1111-1111-1111-111111111111"
        with tempfile.TemporaryDirectory() as tmpdir:
            app = Flask(__name__)
            app.config["GENERATED_FOLDER"] = tmpdir
            sections = [
                {
                    "id": "technical-1",
                    "order_index": 1,
                    "level": 1,
                    "title": "技术响应文件",
                    "content": "技术标正文。",
                    "metadata": {"volume_type": "technical"},
                },
                {
                    "id": "business-1",
                    "order_index": 2,
                    "level": 1,
                    "title": "商务偏差表",
                    "content": "商务标正文。",
                    "metadata": {"volume_type": "business"},
                },
            ]
            cover_fields = {
                "项目名称": "国网辽宁电力2025年第三次物资协议库存招标采购",
                "招标编号": "2225AC",
            }

            with (
                app.app_context(),
                patch("backend.api.routes.get_project_interpretation", return_value={
                    "project": {"id": project_id, "project_name": "上传文件名"},
                    "analysis": {"project_meta": {"project_name": "上传文件名", "cover_fields": cover_fields}},
                }),
                patch("backend.api.routes.list_bid_sections", return_value=sections),
            ):
                markdown_path, document_title, report = build_project_bid_markdown(project_id, volume_type="technical")
                markdown = markdown_path.read_text(encoding="utf-8")

        self.assertEqual(document_title, "上传文件名投标文件-技术标")
        self.assertEqual(report["scope"], "volume")
        self.assertEqual(report["volume_type"], "technical")
        self.assertEqual(report["volume_name"], "技术标")
        self.assertEqual(report["delivery_file_type"], "技术投标文件")
        self.assertEqual(report["cover_fields"]["文件类型"], "技术投标文件")
        self.assertEqual(report["formal_readiness"]["template_id"], "technical_bid_standard")
        self.assertEqual(report["formal_readiness"]["template_family"], "formal_bid_xinjiang_sgcc_reference")
        self.assertEqual(report["section_count"], 1)
        self.assertIn("技术标正文", markdown)
        self.assertNotIn("商务标正文", markdown)

    def test_bid_markdown_formal_readiness_rejects_simulated_customer_values(self):
        project_id = "11111111-1111-1111-1111-111111111111"
        with tempfile.TemporaryDirectory() as tmpdir:
            app = Flask(__name__)
            app.config["GENERATED_FOLDER"] = tmpdir
            sections = [
                {
                    "id": "section-1",
                    "order_index": 1,
                    "level": 1,
                    "title": "投标函",
                    "content": "投标人（盖章）：河北泰昌电力器材科技有限公司",
                }
            ]
            confirmed_values = {
                "project_name": "测试项目",
                "tender_no": "TEST-001",
                "tender_unit": "测试招标人",
                "package_no": "包1",
                "package_name": "测试包",
                "material_category": "电缆保护管",
                "goods_list_summary": "测试货物清单",
                "bidder_name": DOCX_BIDDER_FULL_NAME,
                "company_address": "河北省保定市",
                "legal_representative": "晁坤琳",
                "unified_social_credit_code": "91130607056539515C",
                "total_bid_price": "8888888元（内部测试模拟值，非正式报价）",
                "bid_bond_amount": "100000元（内部测试模拟值，非正式保证金金额）",
                "delivery_period": "内部测试模拟为合同签订后30日",
                "warranty_period": "内部测试模拟为验收后12个月",
                "bid_validity_days": "90",
                "authorized_representative": "张三（内部测试模拟授权代表）",
                "authorized_representative_id": "110101199001011234（内部测试模拟身份证号）",
                "signature_date": "2026年06月25日（内部测试模拟日期）",
                "technical_parameter_summary": "测试参数",
                "technical_deviation_candidates": "测试偏差候选",
            }

            with (
                app.app_context(),
                patch("backend.api.routes.get_project_interpretation", return_value={
                    "project": {"id": project_id, "project_name": "测试项目"},
                    "analysis": {"project_meta": {"bid_prefill": {"confirmed_values": confirmed_values}}},
                }),
                patch("backend.api.routes.list_bid_sections", return_value=sections),
            ):
                _, _, report = build_project_bid_markdown(project_id)

        readiness = report["formal_readiness"]
        self.assertEqual("formal_bid_standard", readiness["template_id"])
        self.assertFalse(readiness["ready"])
        missing_keys = {item["key"] for item in readiness["missing_formal_required_fields"]}
        self.assertIn("total_bid_price", missing_keys)
        self.assertIn("authorized_representative", missing_keys)

    def test_bid_markdown_with_images_loads_taichang_assets(self):
        project_id = "11111111-1111-1111-1111-111111111111"
        with tempfile.TemporaryDirectory() as tmpdir:
            app = Flask(__name__)
            app.config["GENERATED_FOLDER"] = tmpdir
            sections = [
                {
                    "id": "section-1",
                    "order_index": 1,
                    "level": 1,
                    "title": "试验检测能力",
                    "content": "泰昌具备电子天平和万能试验机等试验检测能力。",
                    "metadata": {"volume_type": "technical"},
                }
            ]
            assets = [
                {
                    "id": "asset-1",
                    "title": "泰昌电子天平",
                    "category": "试验检测",
                    "asset_type": "image",
                    "metadata": {
                        "enterprise": "泰昌",
                        "doc_owner": DOCX_BIDDER_FULL_NAME,
                        "source_domain": "enterprise_fact",
                        "evidence_type": "testing_capacity",
                        "target_library": "product_library",
                        "library_type": "product",
                        "reference_only": False,
                    },
                }
            ]

            with (
                app.app_context(),
                patch("backend.api.routes.get_project_interpretation", return_value={
                    "project": {"id": project_id, "project_name": "国网辽宁电力2025年第三次物资协议库存招标采购招标文件"},
                    "analysis": {"project_meta": {"project_name": "国网辽宁电力2025年第三次物资协议库存招标采购招标文件"}},
                }),
                patch("backend.api.routes.list_bid_sections", return_value=sections),
                patch("backend.api.routes.list_knowledge_assets", return_value=assets) as list_assets_mock,
            ):
                markdown_path, document_title, report = build_project_bid_markdown(project_id, with_images=True)

            markdown = markdown_path.read_text(encoding="utf-8")
            self.assertEqual("国网辽宁电力2025年第三次物资协议库存招标采购投标文件", document_title)
            list_assets_mock.assert_called_once()
            self.assertEqual(1, report["asset_candidates"])
            self.assertEqual(1, report["selected"])
            self.assertIn("/api/bidding/knowledge/assets/asset-1/file?variant=original", markdown)

    def test_manual_asset_image_survives_formal_export_image_cleanup(self):
        content = (
            "泰昌产品资料如下。\n\n"
            "![泰昌MPP生产线资料](/api/knowledge/assets/11111111-1111-1111-1111-111111111111/file?variant=original)\n\n"
            "![临时截图](https://example.com/temp.png)\n"
        )

        cleaned = _strip_untrusted_export_images(content, remove_all=True)

        self.assertIn("![泰昌MPP生产线资料](/api/knowledge/assets/11111111-1111-1111-1111-111111111111/file?variant=original)", cleaned)
        self.assertNotIn("https://example.com/temp.png", cleaned)

    def test_formal_asset_title_removes_empty_brackets(self):
        self.assertEqual("社保证明", clean_formal_asset_title("社保证明（）"))
        self.assertEqual("社保证明", clean_formal_asset_title("社保证明（ ）"))
        self.assertEqual("社保证明", clean_formal_asset_title("社保证明（-）"))

    def test_bid_markdown_with_images_does_not_repeat_same_asset(self):
        project_id = "11111111-1111-1111-1111-111111111111"
        with tempfile.TemporaryDirectory() as tmpdir:
            app = Flask(__name__)
            app.config["GENERATED_FOLDER"] = tmpdir
            sections = [
                {
                    "id": "section-1",
                    "order_index": 1,
                    "level": 1,
                    "title": "试验检测能力",
                    "content": "泰昌具备电子天平试验检测能力。",
                    "metadata": {"volume_type": "technical"},
                },
                {
                    "id": "section-2",
                    "order_index": 2,
                    "level": 1,
                    "title": "试验检测设备",
                    "content": "泰昌万能试验机和电子天平配置完善。",
                    "metadata": {"volume_type": "technical"},
                },
            ]
            asset = {
                "id": "asset-1",
                "title": "泰昌电子天平",
                "category": "试验检测",
                "asset_type": "image",
                "metadata": {
                    "enterprise": "泰昌",
                    "doc_owner": DOCX_BIDDER_FULL_NAME,
                    "source_domain": "enterprise_fact",
                    "evidence_type": "testing_capacity",
                    "target_library": "product_library",
                    "library_type": "product",
                    "reference_only": False,
                },
            }

            with (
                app.app_context(),
                patch("backend.api.routes.get_project_interpretation", return_value={
                    "project": {"id": project_id, "project_name": "测试招标文件"},
                    "analysis": {"project_meta": {"project_name": "测试招标文件"}},
                }),
                patch("backend.api.routes.list_bid_sections", return_value=sections),
                patch("backend.api.routes.list_knowledge_assets", return_value=[asset]),
            ):
                markdown_path, _, report = build_project_bid_markdown(project_id, with_images=True)

            markdown = markdown_path.read_text(encoding="utf-8")
            self.assertEqual(1, report["selected"])
            self.assertEqual(1, markdown.count("/api/bidding/knowledge/assets/asset-1/file?variant=original"))

    def test_bid_markdown_with_images_prefers_project_performance_assets(self):
        project_id = "11111111-1111-1111-1111-111111111111"
        with tempfile.TemporaryDirectory() as tmpdir:
            app = Flask(__name__)
            app.config["GENERATED_FOLDER"] = tmpdir
            sections = [
                {
                    "id": "section-1",
                    "order_index": 1,
                    "level": 1,
                    "title": "类似项目业绩",
                    "content": "提供同类产品供货合同、中标通知书和业绩证明材料。",
                    "metadata": {"volume_type": "qualification"},
                }
            ]
            assets = [
                {
                    "id": "cert-1",
                    "title": "泰昌体系认证证书",
                    "category": "资质证书",
                    "asset_type": "qualification_image",
                    "metadata": {
                        "enterprise": "泰昌",
                        "doc_owner": DOCX_BIDDER_FULL_NAME,
                        "source_domain": "enterprise_fact",
                        "evidence_type": "certification",
                        "target_library": "qualification_library",
                        "library_type": "qualification",
                        "reference_only": False,
                    },
                },
                {
                    "id": "award-1",
                    "title": "泰昌电缆保护管中标通知书第1页",
                    "category": "项目业绩",
                    "asset_type": "qualification_image",
                    "metadata": {
                        "enterprise": "泰昌",
                        "doc_owner": DOCX_BIDDER_FULL_NAME,
                        "source_domain": "enterprise_fact",
                        "evidence_type": "project_performance",
                        "target_library": "qualification_library",
                        "library_type": "qualification",
                        "reference_only": False,
                    },
                    "searchable_text": "中标通知书 招标编号 0322AB 包号 157-保护管",
                },
            ]

            with (
                app.app_context(),
                patch("backend.api.routes.get_project_interpretation", return_value={
                    "project": {"id": project_id, "project_name": "测试招标文件"},
                    "analysis": {"project_meta": {"project_name": "测试招标文件"}},
                }),
                patch("backend.api.routes.list_bid_sections", return_value=sections),
                patch("backend.api.routes.list_knowledge_assets", return_value=assets),
            ):
                markdown_path, _, report = build_project_bid_markdown(project_id, with_images=True)

            markdown = markdown_path.read_text(encoding="utf-8")
            self.assertEqual(1, report["selected"])
            self.assertIn("/api/bidding/knowledge/assets/award-1/file?variant=original", markdown)
            self.assertNotIn("/api/bidding/knowledge/assets/cert-1/file?variant=original", markdown)
            self.assertEqual("project_performance", report["manifest"][0]["evidence_type"])

    def test_docx_first_page_is_formal_toc_and_title_is_not_outline_heading(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "toc.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 招标文件",
                        "",
                        "# 1. 企业营业执照",
                        "",
                        "正文内容。",
                        "",
                        "# 2. 安全生产许可证",
                        "",
                        "正文内容。",
                        "",
                        "## 2.1 安全生产许可范围",
                        "",
                        "正文内容。",
                    ]
                ),
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            document = Document(str(output_path))
            non_empty_paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            headings = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.style and paragraph.style.name.startswith("Heading")
            ]
            hyperlinks = document._element.xpath(".//w:hyperlink")
            bookmarks = document._element.xpath(".//w:bookmarkStart")
            field_codes = [
                node.text or ""
                for node in document._element.xpath(".//w:instrText")
            ]
            document_xml = document._element.xml

            self.assertEqual(non_empty_paragraphs[0], "投标文件")
            self.assertIn(f"投标人：{DOCX_BIDDER_FULL_NAME}", non_empty_paragraphs[:8])
            self.assertIn("法定代表人或其委托代理人：        （签名）", non_empty_paragraphs[:8])
            self.assertIn("目  录", non_empty_paragraphs[:10])
            self.assertTrue(any(item.startswith("1. 企业营业执照") for item in non_empty_paragraphs[:10]))
            self.assertTrue(any(item.startswith("2. 安全生产许可证") for item in non_empty_paragraphs[:10]))
            self.assertNotIn("招标文件", headings)
            self.assertIn("1. 企业营业执照", headings)
            self.assertIn("2. 安全生产许可证", headings)
            self.assertEqual(len(hyperlinks), 0)
            self.assertGreaterEqual(len(bookmarks), 3)
            self.assertTrue(any("PAGEREF bid_heading_1" in code for code in field_codes))
            self.assertIn('w:leader="dot"', document_xml)
            self.assertIn('w:dirty="true"', document_xml)

    def test_docx_requests_field_update_on_open_for_toc_page_numbers(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "toc-fields.md"
            markdown_path.write_text(
                "# 招标文件\n\n# 1. 企业营业执照\n\n正文内容。",
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            document = Document(str(output_path))

            self.assertTrue(document.settings.element.xpath(".//w:updateFields[@w:val='true']"))
            self.assertTrue(document._element.xpath(".//w:fldChar[@w:dirty='true']"))

    def test_formal_toc_stability_level_limit_indents_fields_and_markers(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "toc-stability.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 目录稳定性测试投标文件",
                        "",
                        "# 1. 一级章节",
                        "",
                        "正文内容。",
                        "",
                        "## 1.1 二级章节",
                        "",
                        "正文内容。",
                        "",
                        "### 1.1.1 三级章节",
                        "",
                        "正文内容。",
                        "",
                        "#### 1.1.1.1 四级章节",
                        "",
                        "正文内容。",
                        "",
                        "##### 1.1.1.1.1 五级章节",
                        "",
                        "正文内容。",
                    ]
                ),
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            document = Document(str(output_path))
            non_empty = [p.text for p in document.paragraphs if p.text.strip()]
            toc_entries = [p for p in document.paragraphs if "\t" in p.text and p.text.strip()[0].isdigit()]
            field_codes = [node.text or "" for node in document._element.xpath(".//w:instrText")]

            self.assertIn("目  录", non_empty[:8])
            self.assertEqual(2, len(toc_entries))
            self.assertEqual([
                "1. 一级章节",
                "1.1 二级章节",
            ], [p.text.split("\t")[0] for p in toc_entries])
            self.assertNotIn("1.1.1 三级章节", "\n".join(p.text for p in toc_entries))
            self.assertNotIn("1.1.1.1.1 五级章节", "\n".join(p.text for p in toc_entries))

            expected_indents = [0, 18]
            for paragraph, expected_indent in zip(toc_entries, expected_indents):
                self.assertEqual(expected_indent, paragraph.paragraph_format.left_indent.pt)
                self.assertEqual(DOCX_TOC_ENTRY_LINE_SPACING, paragraph.paragraph_format.line_spacing.pt)
                self.assertEqual(0, paragraph.paragraph_format.first_line_indent.pt)
                for run in paragraph.runs:
                    if run.text.strip():
                        self.assertFalse(run.bold)
                        if run.font.size:
                            expected_size = DOCX_TOC_PAGE_NUMBER_FONT_SIZE if run.text.strip().isdigit() else DOCX_TOC_ENTRY_FONT_SIZE
                            self.assertEqual(expected_size, run.font.size.pt)
                self.assertFalse(paragraph._p.xpath(".//w:keepLines"))
                self.assertFalse(paragraph._p.xpath(".//w:keepNext"))
                self.assertTrue(paragraph._p.xpath(".//w:tab[@w:val='right'][@w:leader='dot']"))

            for index in range(1, 3):
                self.assertTrue(any(f"PAGEREF bid_heading_{index}" in code for code in field_codes))
            self.assertFalse(any("PAGEREF bid_heading_3" in code for code in field_codes))

            with ZipFile(output_path) as docx_zip:
                document_xml = docx_zip.read("word/document.xml").decode("utf-8")
                styles_xml = docx_zip.read("word/styles.xml").decode("utf-8")
            self.assertNotIn("w:keepLines", document_xml)
            self.assertNotIn("w:keepNext", document_xml)
            self.assertNotIn("w:pageBreakBefore", document_xml)
            self.assertNotIn("w:keepLines", styles_xml)
            self.assertNotIn("w:keepNext", styles_xml)
            self.assertNotIn("w:pageBreakBefore", styles_xml)

    def test_formal_image_caption_is_sanitized_centered_and_small(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "caption.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 图片题注测试投标文件",
                        "",
                        "# 1. 检验报告",
                        "",
                        "图示：泰昌CPVC电缆保护管检验报告内径250第1页",
                        "",
                        "资料：3.职业健康安全管理体系认证证书第1页",
                        "",
                        "图示：晁坤琳2原图（脱敏示意图）",
                        "",
                        "图示：泰昌试验设备台账原图",
                    ]
                ),
                encoding="utf-8",
            )

            output_path, report = convert_md_to_word(markdown_path, return_report=True)
            document = Document(str(output_path))
            captions = [p for p in document.paragraphs if p.text.startswith("资料：")]

            self.assertEqual([
                "资料：CPVC电缆保护管检验报告",
                "资料：职业健康安全管理体系认证证书",
                "资料：身份证明文件",
                "资料：试验设备台账",
            ], [p.text for p in captions])
            for caption in captions:
                self.assertEqual(WD_ALIGN_PARAGRAPH.CENTER, caption.alignment)
                self.assertEqual(0, caption.paragraph_format.first_line_indent.pt)
                self.assertEqual(DOCX_IMAGE_CAPTION_FONT_SIZE, caption.runs[0].font.size.pt)
            self.assertNotIn("图示", "\n".join(p.text for p in document.paragraphs))
            self.assertNotIn("内径250", "\n".join(p.text for p in document.paragraphs))
            self.assertNotIn("原图", "\n".join(p.text for p in document.paragraphs))
            self.assertEqual(4, report["captions"]["formalized"])

    def test_plain_paragraph_after_image_prefix_is_not_caption(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "plain_image_word.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 图片后续正文测试投标文件",
                        "",
                        "# 1. 手工插图章节",
                        "",
                        "图片之后继续编辑文字，导出 Word 应保持文字和图片顺序。",
                    ]
                ),
                encoding="utf-8",
            )

            output_path, report = convert_md_to_word(markdown_path, return_report=True)
            document = Document(str(output_path))
            paragraph_text = "\n".join(p.text for p in document.paragraphs)

            self.assertIn("图片之后继续编辑文字，导出 Word 应保持文字和图片顺序。", paragraph_text)
            self.assertFalse(any(p.text.startswith("资料：之后继续编辑文字") for p in document.paragraphs))
            self.assertEqual(0, report["captions"]["formalized"])

    def test_formal_bid_text_is_black_and_major_headings_start_new_page(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "black-pagination.md"
            markdown_path.write_text(
                "# 颜色分页测试投标文件\n\n# 1. 商务响应\n\n正文内容。\n\n## 1.1 投标函\n\n正文内容。",
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            self.assertFalse(should_start_heading_on_new_page(2, "1.1 投标函", 1))
            self.assertTrue(should_start_heading_on_new_page(1, "2. 商务响应", 1))
            with ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml").decode("utf-8")
                styles_xml = archive.read("word/styles.xml").decode("utf-8")
            self.assertNotIn("1F4E79", document_xml)
            self.assertNotIn("1F4E79", styles_xml)
            self.assertIn('w:val="000000"', document_xml)
            self.assertIn('w:type="page"', document_xml)

    def test_mermaid_fence_source_is_not_exported_when_conversion_fails(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "mermaid.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 流程图测试",
                        "",
                        "# 1. 现场施工流程",
                        "",
                        "```mermaid",
                        "graph TD",
                        "A[接收订单] --> B[生产备货]",
                        "```",
                    ]
                ),
                encoding="utf-8",
            )

            with patch("backend.export.md_to_word.process_mermaid", return_value=False):
                output_path, report = convert_md_to_word(markdown_path, return_report=True)

            document = Document(str(output_path))
            text = "\n".join(p.text for p in document.paragraphs)
            self.assertNotIn("```mermaid", text)
            self.assertNotIn("graph TD", text)
            self.assertEqual(1, report["mermaid"]["found"])
            self.assertEqual(1, report["mermaid"]["skipped"])

    def test_soffice_refresh_replaces_docx_and_reports_success(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            source = Path(tmpdir) / "test.docx"
            source.write_bytes(b"original")

            def fake_run(cmd, capture_output, text, timeout, check):
                outdir = Path(cmd[cmd.index("--outdir") + 1])
                (outdir / source.name).write_bytes(b"refreshed")

                class Result:
                    returncode = 0
                    stdout = "converted"
                    stderr = ""

                return Result()

            with (
                patch.dict("os.environ", {"DOCX_REFRESH_FIELDS": "true", "SOFFICE_BIN": "soffice-test"}),
                patch("backend.export.md_to_word.subprocess.run", side_effect=fake_run),
            ):
                refreshed_path, report = refresh_docx_fields_with_soffice(source)

            self.assertEqual(source, refreshed_path)
            self.assertEqual(b"refreshed", source.read_bytes())
            self.assertEqual("refreshed", report["status"])
            self.assertFalse(report["manual_refresh_required"])
            self.assertIn("自动刷新", report["user_message"])

    def test_soffice_refresh_can_be_disabled(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            source = Path(tmpdir) / "test.docx"
            source.write_bytes(b"original")

            with patch.dict("os.environ", {"DOCX_REFRESH_FIELDS": "false"}):
                refreshed_path, report = refresh_docx_fields_with_soffice(source)

            self.assertEqual(source, refreshed_path)
            self.assertEqual(b"original", source.read_bytes())
            self.assertEqual("skipped", report["status"])
            self.assertTrue(report["manual_refresh_required"])
            self.assertIn("关闭", report["user_message"])

    def test_soffice_refresh_reports_missing_configured_binary_without_blocking_docx(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            source = Path(tmpdir) / "test.docx"
            source.write_bytes(b"original")

            with patch.dict("os.environ", {"DOCX_REFRESH_FIELDS": "true", "SOFFICE_BIN": "/not/exist/soffice"}):
                refreshed_path, report = refresh_docx_fields_with_soffice(source)

            self.assertEqual(source, refreshed_path)
            self.assertEqual(b"original", source.read_bytes())
            self.assertEqual("failed", report["status"])
            self.assertTrue(report["manual_refresh_required"])
            self.assertIn("does not exist", report["reason"])
            self.assertIn("Word/WPS", report["user_message"])

    def test_soffice_refresh_timeout_reports_manual_refresh_required(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            source = Path(tmpdir) / "test.docx"
            source.write_bytes(b"original")

            def fake_run(*args, **kwargs):
                raise subprocess.TimeoutExpired(cmd=kwargs.get("args") or "soffice", timeout=3)

            with (
                patch.dict("os.environ", {
                    "DOCX_REFRESH_FIELDS": "true",
                    "SOFFICE_BIN": "soffice-test",
                    "DOCX_REFRESH_TIMEOUT_SECONDS": "3",
                }),
                patch("backend.export.md_to_word.subprocess.run", side_effect=fake_run),
            ):
                refreshed_path, report = refresh_docx_fields_with_soffice(source)

            self.assertEqual(source, refreshed_path)
            self.assertEqual(b"original", source.read_bytes())
            self.assertEqual("failed", report["status"])
            self.assertTrue(report["manual_refresh_required"])
            self.assertIn("timed out", report["reason"])

    def test_formal_docx_cleans_generation_notes_emoji_and_preserves_table(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "formal.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 某水库除险加固工程投标文件",
                        "",
                        "## 商务响应文件",
                        "",
                        "（本章节正文共计约2980字，符合目标字数要求，结构完整，可用于直接插入标书资格文件分册）",
                        "",
                        "关键提醒：📌 本投标单位承诺严格响应招标文件要求。",
                        "",
                        "| 项目 | 响应 |",
                        "| --- | --- |",
                        "| 工期 | 满足招标文件要求 |",
                    ]
                ),
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            document = Document(str(output_path))
            full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

            self.assertIn("商务响应文件", full_text)
            self.assertIn("本投标单位承诺严格响应招标文件要求", full_text)
            self.assertNotIn("本章节正文共计", full_text)
            self.assertNotIn("📌", full_text)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(document.tables[0].cell(1, 0).text, "工期")

    def test_formal_bid_template_uses_sgcc_default_fonts_layout_and_metadata(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "sgcc-template.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 国网山西电力2026年第二次物资协议库存公开招标采购投标文件",
                        "",
                        "# 1. 投标函及格式文件",
                        "",
                        "本企业承诺严格响应招标文件第六章投标文件格式要求。",
                        "",
                        "| 序号 | 文件名称 | 响应情况 |",
                        "| --- | --- | --- |",
                        "| 1 | 商务投标文件 | 已按要求编制 |",
                    ]
                ),
                encoding="utf-8",
            )

            output_path, report = convert_md_to_word(markdown_path, return_report=True)
            document = Document(str(output_path))
            normal = document.styles["Normal"]
            section = document.sections[0]
            body_paragraph = next(p for p in document.paragraphs if p.text.startswith("本企业承诺"))
            body_run = body_paragraph.runs[0]
            body_rfonts = body_run._element.rPr.rFonts
            table_run = document.tables[0].cell(1, 1).paragraphs[0].runs[0]
            table = document.tables[0]
            table_pr = table._tbl.tblPr
            table_width = table_pr.find(qn("w:tblW"))
            table_layout = table_pr.find(qn("w:tblLayout"))

            self.assertEqual("formal_bid_standard", report["template"]["template_id"])
            self.assertEqual(DOCX_BIDDER_FULL_NAME, report["template"]["bidder_full_name"])
            self.assertEqual(DOCX_BODY_EAST_ASIA, report["template"]["body_font"])
            self.assertEqual(DOCX_BODY_FIRST_LINE_INDENT_PT, report["template"]["body_first_line_indent_pt"])
            self.assertEqual(DOCX_BODY_LINE_SPACING_RULE, report["template"]["body_line_spacing_rule"])
            self.assertEqual(DOCX_BODY_LINE_SPACING, report["template"]["body_line_spacing"])
            self.assertEqual(DOCX_LIST_LEFT_INDENT_PT, report["template"]["list_left_indent_pt"])
            self.assertEqual(DOCX_LIST_HANGING_INDENT_PT, report["template"]["list_hanging_indent_pt"])
            self.assertFalse(report["template"]["heading_keep_with_next"])
            self.assertFalse(report["template"]["toc_entry_bold_all"])
            self.assertEqual("assets/template_words/5d2a2c833dad4bb3b3ccc0856f755b54.docx", report["template"]["reference_templates"][0]["path"])
            self.assertTrue(report["template"]["image_layout"]["uniform_frame_enabled"])
            self.assertEqual(DOCX_IMAGE_FRAME_WIDTH_IN, report["template"]["image_layout"]["frame_width_in"])
            self.assertEqual(DOCX_IMAGE_FRAME_HEIGHT_IN, report["template"]["image_layout"]["frame_height_in"])
            self.assertEqual(18, report["template"]["table_line_spacing_pt"])
            self.assertEqual(DOCX_BODY_EAST_ASIA, normal._element.rPr.rFonts.get(qn("w:eastAsia")))
            self.assertEqual(DOCX_BODY_EAST_ASIA, body_rfonts.get(qn("w:eastAsia")))
            self.assertEqual(12, body_run.font.size.pt)
            self.assertEqual(WD_LINE_SPACING.ONE_POINT_FIVE, body_paragraph.paragraph_format.line_spacing_rule)
            self.assertEqual(DOCX_BODY_LINE_SPACING, body_paragraph.paragraph_format.line_spacing)
            self.assertEqual(DOCX_BODY_FIRST_LINE_INDENT_PT, body_paragraph.paragraph_format.first_line_indent.pt)
            self.assertEqual(DOCX_TABLE_EAST_ASIA, table_run._element.rPr.rFonts.get(qn("w:eastAsia")))
            self.assertEqual(12, table_run.font.size.pt)
            self.assertEqual("5000", table_width.get(qn("w:w")))
            self.assertEqual("pct", table_width.get(qn("w:type")))
            self.assertEqual("fixed", table_layout.get(qn("w:type")))
            self.assertEqual(21, round(section.page_width.cm))
            self.assertEqual(29.7, round(section.page_height.cm, 1))
            self.assertEqual(2.5, round(section.top_margin.cm, 1))
            self.assertEqual(2.5, round(section.bottom_margin.cm, 1))
            self.assertEqual(2.8, round(section.left_margin.cm, 1))
            self.assertEqual(2.5, round(section.right_margin.cm, 1))
            self.assertIn("国网山西电力2026年第二次物资协议库存公开招标采购投标文件", section.header.paragraphs[0].text)
            self.assertIn("投标文件", section.header.paragraphs[0].text)
            self.assertFalse(report["template"]["header_footer"]["header_logo"])

    def test_formal_bid_body_headings_and_lists_use_stable_paragraph_format(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "body-format.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 正文格式测试投标文件",
                        "",
                        "",
                        "# 1. 施工组织设计",
                        "",
                        "本节正文用于检查正式投标文件的正文段落格式。",
                        "",
                        "",
                        "## 1.1 组织措施",
                        "",
                        "- 配置项目经理和技术负责人。",
                        "- 建立质量、安全、进度协调机制。",
                        "",
                        "1. 明确资料提交节点。",
                        "2. 明确现场配合责任。",
                        "",
                        "后续正文不得因为 Markdown 空行出现异常空白段落。",
                    ]
                ),
                encoding="utf-8",
            )

            output_path, report = convert_md_to_word(markdown_path, return_report=True)
            document = Document(str(output_path))
            body_paragraph = next(p for p in document.paragraphs if p.text.startswith("本节正文"))
            trailing_body = next(p for p in document.paragraphs if p.text.startswith("后续正文"))
            heading = next(p for p in document.paragraphs if p.text == "1. 施工组织设计")
            subheading = next(p for p in document.paragraphs if p.text == "1.1 组织措施")
            bullet = next(p for p in document.paragraphs if p.text.startswith("配置项目经理"))
            numbered = next(p for p in document.paragraphs if p.text.startswith("1. 明确资料提交"))

            self.assertEqual(DOCX_BODY_FIRST_LINE_INDENT_PT, report["template"]["body_first_line_indent_pt"])
            self.assertEqual(WD_LINE_SPACING.ONE_POINT_FIVE, body_paragraph.paragraph_format.line_spacing_rule)
            self.assertEqual(DOCX_BODY_LINE_SPACING, body_paragraph.paragraph_format.line_spacing)
            self.assertEqual(DOCX_BODY_FIRST_LINE_INDENT_PT, body_paragraph.paragraph_format.first_line_indent.pt)
            self.assertEqual(0, body_paragraph.paragraph_format.space_before.pt)
            self.assertEqual(0, body_paragraph.paragraph_format.space_after.pt)
            self.assertIsNone(body_paragraph.paragraph_format.keep_together)
            self.assertIsNone(body_paragraph.paragraph_format.widow_control)
            self.assertIsNone(trailing_body.paragraph_format.keep_together)

            self.assertIsNone(heading.paragraph_format.keep_with_next)
            self.assertIsNone(heading.paragraph_format.keep_together)
            self.assertEqual(0, heading.paragraph_format.first_line_indent.pt)
            self.assertEqual(WD_ALIGN_PARAGRAPH.CENTER, heading.alignment)
            self.assertEqual(22, heading.runs[0].font.size.pt)
            self.assertEqual(15, subheading.runs[0].font.size.pt)
            self.assertTrue(subheading.runs[0].font.bold)

            for paragraph in (bullet, numbered):
                self.assertEqual(WD_LINE_SPACING.ONE_POINT_FIVE, paragraph.paragraph_format.line_spacing_rule)
                self.assertEqual(DOCX_BODY_LINE_SPACING, paragraph.paragraph_format.line_spacing)
                self.assertEqual(DOCX_LIST_LEFT_INDENT_PT, paragraph.paragraph_format.left_indent.pt)
                self.assertEqual(-DOCX_LIST_HANGING_INDENT_PT, paragraph.paragraph_format.first_line_indent.pt)

    def test_numbered_lists_preserve_source_markers_and_restart_between_sections(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "numbering-restart.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 编号回归测试投标文件",
                        "",
                        "# 1. 商务响应",
                        "",
                        "1. 第一项商务要求。",
                        "2. 第二项商务要求。",
                        "",
                        "## 1.1 技术响应",
                        "",
                        "1. 第一项技术要求。",
                        "2. 第二项技术要求。",
                        "",
                        "5.4.1 生产工艺文件控制",
                        "5.4.2 过程检验",
                    ]
                ),
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            document = Document(str(output_path))
            texts = [p.text for p in document.paragraphs]

            self.assertEqual(2, texts.count("1. 第一项商务要求。") + texts.count("1. 第一项技术要求。"))
            self.assertIn("1. 第一项商务要求。", texts)
            self.assertIn("2. 第二项商务要求。", texts)
            self.assertIn("1. 第一项技术要求。", texts)
            self.assertIn("2. 第二项技术要求。", texts)
            self.assertIn("5.4.1 生产工艺文件控制", texts)
            self.assertIn("5.4.2 过程检验", texts)
            with ZipFile(output_path) as docx_zip:
                document_xml = docx_zip.read("word/document.xml").decode("utf-8")
            self.assertNotIn("<w:numPr>", document_xml)

    def test_formal_docx_does_not_emit_black_square_paragraph_markers(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "paragraph-marker.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 段落格式标记测试投标文件",
                        "",
                        "# 1. 商务响应",
                        "",
                        "正文段落一。",
                        "",
                        "## 1.1 响应要求",
                        "",
                        "正文段落二。",
                    ]
                ),
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            document = Document(str(output_path))
            normal = document.styles["Normal"]
            toc_entry = next(p for p in document.paragraphs if p.text.startswith("1. 商务响应") and "\t" in p.text)
            body_paragraph = next(p for p in document.paragraphs if p.text == "正文段落一。")
            heading = next(p for p in document.paragraphs if p.text == "1. 商务响应" and p.style.name.startswith("Heading"))

            self.assertIsNone(normal.paragraph_format.keep_together)
            self.assertIsNone(normal.paragraph_format.keep_with_next)
            self.assertIsNone(toc_entry.paragraph_format.keep_together)
            self.assertIsNone(toc_entry.paragraph_format.keep_with_next)
            self.assertIsNone(body_paragraph.paragraph_format.keep_together)
            self.assertIsNone(body_paragraph.paragraph_format.keep_with_next)
            self.assertIsNone(heading.paragraph_format.keep_together)
            self.assertIsNone(heading.paragraph_format.keep_with_next)
            with ZipFile(output_path) as docx_zip:
                document_xml = docx_zip.read("word/document.xml").decode("utf-8")
                styles_xml = docx_zip.read("word/styles.xml").decode("utf-8")
            for xml in (document_xml, styles_xml):
                self.assertNotIn("w:keepLines", xml)
                self.assertNotIn("w:keepNext", xml)
                self.assertNotIn("w:pageBreakBefore", xml)

            _, report = convert_md_to_word(markdown_path, return_report=True)
            self.assertEqual(0, report["marker_cleanup"]["counts_after"]["keepLines"])
            self.assertEqual(0, report["marker_cleanup"]["counts_after"]["keepNext"])

    def test_formal_bid_tables_use_repeat_header_width_and_cell_spacing(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "formal-table.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 表格正式化测试投标文件",
                        "",
                        "# 1. 资格审查资料",
                        "",
                        "| 序号 | 文件名称 | 响应情况 | 说明 |",
                        "| --- | --- | --- | --- |",
                        "| 1 | 营业执照 | 已提供 | 原件扫描件清晰完整，满足招标文件要求 |",
                        "| 2 | 质量管理体系认证证书 | 已提供 | 证书在有效期内 |",
                    ]
                ),
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            document = Document(str(output_path))
            table = document.tables[0]
            table_pr = table._tbl.tblPr
            table_width = table_pr.find(qn("w:tblW"))
            table_layout = table_pr.find(qn("w:tblLayout"))
            cell_margins = table_pr.find(qn("w:tblCellMar"))
            first_cell_margins = table.cell(1, 1)._tc.get_or_add_tcPr().find(qn("w:tcMar"))
            header_pr = table.rows[0]._tr.get_or_add_trPr()
            header_repeat = header_pr.find(qn("w:tblHeader"))
            header_shading = table.cell(0, 0)._tc.get_or_add_tcPr().find(qn("w:shd"))
            header_run = table.cell(0, 1).paragraphs[0].runs[0]
            body_paragraph = table.cell(1, 1).paragraphs[0]
            centered_paragraph = table.cell(1, 0).paragraphs[0]

            self.assertEqual("5000", table_width.get(qn("w:w")))
            self.assertEqual("pct", table_width.get(qn("w:type")))
            self.assertEqual("fixed", table_layout.get(qn("w:type")))
            self.assertEqual("true", header_repeat.get(qn("w:val")))
            self.assertEqual("EDEDED", header_shading.get(qn("w:fill")))
            self.assertEqual("100", cell_margins.find(qn("w:left")).get(qn("w:w")))
            self.assertEqual("100", first_cell_margins.find(qn("w:left")).get(qn("w:w")))
            self.assertTrue(header_run.bold)
            self.assertEqual(WD_ALIGN_PARAGRAPH.LEFT, body_paragraph.alignment)
            self.assertEqual(WD_ALIGN_PARAGRAPH.CENTER, centered_paragraph.alignment)
            self.assertEqual(WD_LINE_SPACING.EXACTLY, body_paragraph.paragraph_format.line_spacing_rule)
            self.assertEqual(18, body_paragraph.paragraph_format.line_spacing.pt)
            self.assertEqual(0, body_paragraph.paragraph_format.first_line_indent.pt)

    def test_sixth_chapter_forms_preserve_signature_layout_and_deviation_table_rows(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "sixth-chapter-forms.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 第六章格式表单测试投标文件",
                        "",
                        "# 1. 投标函及投标函附录",
                        "",
                        "致：国网辽宁省电力有限公司",
                        "",
                        "投标人（盖单位章）：河北泰昌电力器材科技有限公司",
                        "法定代表人（签字或盖章）：晁坤琳",
                        "授权代表（签字）：内部测试模拟授权代表",
                        "日期：2026年06月25日（内部测试模拟日期）",
                        "",
                        "【投标保证金承诺函】",
                        "",
                        "投标人（盖单位章）：河北泰昌电力器材科技有限公司",
                        "",
                        "## 1.1 技术偏差表",
                        "",
                        "| 序号 | 技术规范条目 | 标准值（技术要求） | 保证值（我公司承诺） | 偏差说明 |",
                        "| --- | --- | --- | --- | --- |",
                        "| 1 | CPVC平均内径 | 符合技术规范要求 | 250.2～250.4mm | 无偏差 |",
                        "| 2 | MPP环刚度 | ≥40kN/m² | 66.40kN/m² | 正偏差 |",
                    ]
                ),
                encoding="utf-8",
            )

            output_path, report = convert_md_to_word(markdown_path, return_report=True)
            document = Document(str(output_path))
            signature_paragraphs = [
                paragraph for paragraph in document.paragraphs
                if paragraph.text.startswith(("投标人（", "法定代表人（", "授权代表（", "日期："))
            ]
            table = document.tables[0]
            widths = [
                int(cell._tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w")))
                for cell in table.rows[0].cells
            ]
            grid_widths = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]

            self.assertEqual(5, len(signature_paragraphs))
            self.assertTrue(all(paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT for paragraph in signature_paragraphs))
            self.assertEqual(["bid_letter", "commitment", "technical_deviation"], report["formal_forms"]["detected_types"])
            self.assertEqual(1, report["formal_forms"]["tables"])
            self.assertEqual(2, report["formal_forms"]["non_split_rows"])
            self.assertEqual(1, report["formal_forms"]["subheadings"])
            self.assertEqual(5, report["formal_forms"]["signature_lines"])
            self.assertLess(widths[0], widths[2])
            self.assertLess(widths[1], widths[3])
            self.assertEqual(widths, grid_widths)
            for row in table.rows:
                cant_split = row._tr.get_or_add_trPr().find(qn("w:cantSplit"))
                self.assertIsNotNone(cant_split)
                self.assertEqual("true", cant_split.get(qn("w:val")))

    def test_taichang_bid_document_title_rewrites_tender_file_title(self):
        self.assertEqual(
            "国网辽宁电力2025年第三次物资协议库存招标采购投标文件",
            taichang_bid_document_title("国网辽宁电力2025年第三次物资协议库存招标采购招标文件"),
        )

    def test_docx_cover_includes_extracted_formal_bid_fields(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "cover-fields.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 国网辽宁电力2025年第三次物资协议库存招标采购商务投标文件",
                        "",
                        "**招标编号：** 2225AC",
                        "**分标编号**：2225AC-1408006-3401",
                        "**分标名称**：电缆保护管CPVC",
                        "**包号**：包1-包2",
                        "",
                        "# 1. 商务偏差表",
                        "",
                        "全部响应，无偏差。",
                    ]
                ),
                encoding="utf-8",
            )

            output_path, report = convert_md_to_word(markdown_path, return_report=True)
            document = Document(str(output_path))
            non_empty_paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

            self.assertIn("投标文件", non_empty_paragraphs[:10])
            self.assertIn("文件类别：商务", non_empty_paragraphs[:10])
            self.assertNotIn("文件类型：商务投标文件", non_empty_paragraphs[:10])
            self.assertIn("招标编号：2225AC", non_empty_paragraphs[:10])
            self.assertIn("分标编号：2225AC-1408006-3401", non_empty_paragraphs[:10])
            self.assertIn("分标名称：电缆保护管CPVC", non_empty_paragraphs[:10])
            self.assertIn("包    号：包1-包2", non_empty_paragraphs[:10])
            self.assertEqual("2225AC", report["template"]["cover_fields"]["招标编号"])

    def test_extract_bid_cover_fields_from_markdown_table(self):
        fields = extract_bid_cover_fields(
            "\n".join(
                [
                    "# 投标文件",
                    "",
                    "| 字段 | 内容 |",
                    "| --- | --- |",
                    "| 招标编号 | 2225AC |",
                    "| 分标名称 | 电缆保护管MPP |",
                ]
            )
        )

        self.assertEqual("2225AC", fields["招标编号"])
        self.assertEqual("电缆保护管MPP", fields["分标名称"])

    def test_extract_bid_cover_fields_rejects_adjacent_field_name_as_value(self):
        fields = extract_bid_cover_fields(
            "\n".join(
                [
                    "# 投标文件",
                    "",
                    "招标编号：2225AC",
                    "分标编号：分标名称",
                ]
            )
        )

        self.assertEqual("2225AC", fields["招标编号"])
        self.assertNotIn("分标编号", fields)

    def test_docx_cover_prefers_structured_tender_fields_over_markdown_fallback(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "cover.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 旧标题招标文件",
                        "",
                        "招标编号：OLD-NO",
                        "",
                        "# 1. 投标函",
                        "",
                        "正文内容。",
                    ]
                ),
                encoding="utf-8",
            )

            output_path, report = convert_md_to_word(
                markdown_path,
                return_report=True,
                cover_fields={
                    "项目名称": "国网辽宁电力2025年第三次物资协议库存招标采购",
                    "文件类型": "技术投标文件",
                    "招标编号": "2225AC",
                    "分标编号": "102-CPVC",
                    "包号": "包1",
                    "招标人": "国网辽宁省电力有限公司",
                },
            )
            document = Document(str(output_path))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            cover_text = "\n".join(paragraph.text for paragraph in document.paragraphs[:14])
            header_text = "\n".join(
                paragraph.text
                for section in document.sections
                for paragraph in section.header.paragraphs
            )
            with ZipFile(output_path) as archive:
                footer_xml = "\n".join(
                    archive.read(name).decode("utf-8", errors="ignore")
                    for name in archive.namelist()
                    if name.startswith("word/footer")
                )

        self.assertEqual(report["template"]["cover_fields"]["招标编号"], "2225AC")
        self.assertIn("国网辽宁电力2025年第三次物资协议", text)
        self.assertIn("库存招标采购", text)
        self.assertIn("投标文件", text)
        self.assertIn("文件类别：技术", text)
        self.assertNotIn("文件类型：技术投标文件", text)
        self.assertIn("招标编号：2225AC", text)
        self.assertIn("分标编号：102-CPVC", text)
        self.assertIn("包    号：包1", text)
        self.assertIn("招标人：国网辽宁省电力有限公司", text)
        self.assertNotIn("招标编号：OLD-NO", cover_text)
        self.assertEqual("", header_text.strip())
        self.assertEqual(report["template"]["header_footer"]["header_text"], "")
        self.assertEqual(report["template"]["header_footer"]["header_text_policy"], "blank")
        self.assertEqual(report["template"]["header_footer"]["page_number_field"], "PAGE")
        self.assertIn("PAGE", footer_xml)
        self.assertNotIn("NUMPAGES", footer_xml)

    def test_technical_bid_reference_profile_uses_volume_template_layout_and_toc_depth(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "technical-template.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 技术投标文件",
                        "",
                        "# （一）技术偏差表",
                        "",
                        "## 1. 技术偏差表",
                        "",
                        "### 1.1 技术特性参数表",
                        "",
                        "#### 附:技术规范点对点应答",
                        "",
                        "正文内容。",
                        "",
                        "| 参数 | 响应 |",
                        "| --- | --- |",
                        "| 额定电压 | 满足 |",
                    ]
                ),
                encoding="utf-8",
            )

            output_path, report = convert_md_to_word(
                markdown_path,
                return_report=True,
                cover_fields={
                    "项目名称": "国网新疆10kV架空绝缘导线采购",
                    "文件类型": "技术投标文件",
                    "招标编号": "SL265A",
                    "分标编号": "TC-TECH-001",
                    "分标名称": "10kV架空绝缘导线测试分标",
                    "包号": "包1-13",
                },
            )
            document = Document(str(output_path))
            section = document.sections[0]
            toc_entries = [p.text.split("\t")[0] for p in document.paragraphs if "\t" in p.text and p.text.strip()]
            full_text = "\n".join(p.text for p in document.paragraphs)
            cover_paragraphs = [p.text for p in document.paragraphs[:12] if p.text.strip()]
            cover_run_specs = {
                p.text: (
                    p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else None,
                    p.runs[0].bold if p.runs else None,
                    p.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) if p.runs and p.runs[0]._element.rPr is not None else None,
                )
                for p in document.paragraphs[:12]
                if p.text.strip()
            }
            first_heading = next(p for p in document.paragraphs if p.text == "（一）技术偏差表")
            body_paragraph = next(p for p in document.paragraphs if p.text == "正文内容。")
            table_run = document.tables[0].cell(1, 1).paragraphs[0].runs[0]
            header_text = "\n".join(
                paragraph.text
                for doc_section in document.sections
                for paragraph in doc_section.header.paragraphs
            )
            with ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
                footer_xml = "\n".join(
                    archive.read(name).decode("utf-8", errors="ignore")
                    for name in archive.namelist()
                    if name.startswith("word/footer")
                )

        self.assertEqual("technical_bid_standard", report["template"]["template_id"])
        self.assertEqual("formal_bid_xinjiang_sgcc_reference", report["template"]["template_family"])
        self.assertEqual("assets/template_words/技术文件 - 10kV架空绝缘导线-新疆.docx", report["template"]["reference_path"])
        self.assertIn("技术规范点对点应答", report["template"]["reference_outline"])
        self.assertEqual(4, report["template"]["toc_max_level"])
        self.assertEqual("宋体", report["template"]["toc_font"])
        self.assertEqual("宋体", report["template"]["body_font"])
        self.assertEqual("宋体", report["template"]["table_font"])
        self.assertEqual("sgcc_reference_volume_cover", report["template"]["cover_layout"]["style"])
        self.assertEqual("after_project_title", report["template"]["cover_layout"]["tender_no_position"])
        self.assertEqual("宋体", report["template"]["header_footer"]["header_font"])
        self.assertEqual("", report["template"]["header_footer"]["header_text"])
        self.assertEqual("blank", report["template"]["header_footer"]["header_text_policy"])
        self.assertEqual("纯数字页码", report["template"]["header_footer"]["page_number_format"])
        self.assertEqual("PAGE", report["template"]["header_footer"]["page_number_field"])
        self.assertFalse(report["template"]["header_footer"]["different_first_page_header_footer"])
        self.assertEqual(3.17, report["template"]["margins_cm"]["left"])
        self.assertEqual(3.17, round(section.left_margin.cm, 2))
        self.assertEqual(3.17, round(section.right_margin.cm, 2))
        self.assertEqual(1.5, round(section.header_distance.cm, 1))
        self.assertFalse(section.different_first_page_header_footer)
        self.assertEqual("", header_text.strip())
        self.assertNotIn("w:titlePg", document_xml)
        self.assertIn("PAGE", footer_xml)
        self.assertNotIn("NUMPAGES", footer_xml)
        self.assertNotIn(">第<", footer_xml)
        self.assertNotIn(">共<", footer_xml)
        self.assertEqual("宋体", body_paragraph.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")))
        self.assertEqual("宋体", table_run._element.rPr.rFonts.get(qn("w:eastAsia")))
        self.assertLess(cover_paragraphs.index("招标编号：SL265A"), cover_paragraphs.index("投标文件"))
        self.assertIn("投标文件", cover_paragraphs)
        self.assertIn("分标编号：TC-TECH-001", cover_paragraphs)
        self.assertIn("分标名称：10kV架空绝缘导线测试分标", cover_paragraphs)
        self.assertIn("包    号：包1-13", cover_paragraphs)
        self.assertIn("文件类别：技术", cover_paragraphs)
        self.assertEqual((18.0, True, "宋体"), cover_run_specs["招标编号：SL265A"])
        self.assertEqual((36.0, True, "宋体"), cover_run_specs["投标文件"])
        self.assertEqual((14.0, True, "宋体"), cover_run_specs["分标编号：TC-TECH-001"])
        self.assertEqual((16.0, True, "宋体"), cover_run_specs[f"投标人：{DOCX_BIDDER_FULL_NAME}（盖单位章）"])
        self.assertEqual((15.0, True, "宋体"), cover_run_specs["法定代表人（单位负责人）或其授权代表人：       （签字）"])
        self.assertIn(f"投标人：{DOCX_BIDDER_FULL_NAME}（盖单位章）", full_text)
        self.assertIn("法定代表人（单位负责人）或其授权代表人：       （签字）", full_text)
        self.assertEqual(WD_ALIGN_PARAGRAPH.LEFT, first_heading.alignment)
        self.assertEqual("宋体", first_heading.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")))
        self.assertEqual(12, first_heading.runs[0].font.size.pt)
        self.assertIn("附:技术规范点对点应答", toc_entries)
        self.assertNotIn("保定铠蒂电力器材有限公司", full_text)
        self.assertNotIn("SL265A-1402005-0001", full_text)

    def test_business_bid_reference_profile_uses_business_template_metadata(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "business-template.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 商务投标文件",
                        "",
                        "# （一）商务偏差表",
                        "",
                        "## 1. 商务偏差表",
                        "",
                        "### 1.1 查询报告及截图",
                        "",
                        "正文内容。",
                    ]
                ),
                encoding="utf-8",
            )

            _, report = convert_md_to_word(
                markdown_path,
                return_report=True,
                cover_fields={
                    "项目名称": "国网新疆10kV架空绝缘导线采购",
                    "文件类型": "商务投标文件",
                    "招标编号": "SL265A",
                },
            )

        self.assertEqual("business_bid_standard", report["template"]["template_id"])
        self.assertEqual("formal_bid_xinjiang_sgcc_reference", report["template"]["template_family"])
        self.assertEqual("assets/template_words/商务文件 - 10kV架空绝缘导线-新疆(1).docx", report["template"]["reference_path"])
        self.assertIn("商务偏差表", report["template"]["reference_outline"])
        self.assertIn("查询报告及截图", report["template"]["reference_outline"])
        self.assertEqual("宋体", report["template"]["toc_font"])
        self.assertEqual("宋体", report["template"]["body_font"])
        self.assertEqual("宋体", report["template"]["table_font"])
        self.assertEqual("sgcc_reference_volume_cover", report["template"]["cover_layout"]["style"])
        self.assertEqual("", report["template"]["header_footer"]["header_text"])
        self.assertEqual("PAGE", report["template"]["header_footer"]["page_number_field"])
        self.assertIn("禁止复用参考稿企业事实", report["template"]["runtime_policy"])

    def test_formal_form_subheading_switch_starts_on_new_page(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "formal-forms.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 商务投标文件",
                        "",
                        "# 1. 投标函及法定格式文件",
                        "",
                        "【一、投标函】",
                        "",
                        "投标人名称：河北泰昌电力器材科技有限公司",
                        "",
                        "【二、法定代表人授权书】",
                        "",
                        "授权代表：李明",
                    ]
                ),
                encoding="utf-8",
            )

            output_path, report = convert_md_to_word(
                markdown_path,
                return_report=True,
                cover_fields={
                    "项目名称": "国网新疆10kV架空绝缘导线采购",
                    "文件类型": "商务投标文件",
                },
            )

            with ZipFile(output_path) as docx:
                root = ET.fromstring(docx.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paragraphs = root.findall(".//w:p", ns)
            auth_index = next(
                index
                for index, paragraph in enumerate(paragraphs)
                if "二、法定代表人授权书" in "".join(t.text or "" for t in paragraph.findall(".//w:t", ns))
            )
            previous = paragraphs[auth_index - 1]
            previous_has_page_break = any(
                br.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type") == "page"
                for br in previous.findall(".//w:br", ns)
            )

        self.assertEqual(1, report["formal_forms"]["subheading_page_breaks"])
        self.assertTrue(previous_has_page_break)

    def test_bid_export_assets_are_limited_to_taichang_enterprise_facts(self):
        taichang_asset = {
            "metadata": {
                "enterprise": "泰昌",
                "doc_owner": DOCX_BIDDER_FULL_NAME,
                "source_domain": "enterprise_fact",
                "reference_only": False,
            }
        }
        haoqian_reference = {
            "metadata": {
                "doc_owner": "河北豪乾电气设备科技有限公司",
                "source_domain": "reference_template",
                "reference_only": True,
            }
        }
        liaoning_tender = {
            "metadata": {
                "doc_owner": "国网辽宁省电力有限公司",
                "source_domain": "tender_requirement",
                "reference_only": False,
            }
        }

        self.assertTrue(_asset_allowed_for_bid(taichang_asset))
        self.assertFalse(_asset_allowed_for_bid(haoqian_reference))
        self.assertFalse(_asset_allowed_for_bid(liaoning_tender))

    def test_markdown_image_limit_records_skipped_images(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            image_path = tmp / "tiny.png"
            image_path.write_bytes(
                bytes.fromhex(
                    "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753de"
                    "0000000c49444154789c63606060000000040001f61738550000000049454e44ae426082"
                )
            )
            markdown_path = tmp / "images.md"
            markdown_path.write_text(
                f"# 图文测试\n\n![图1]({image_path})\n\n![图2]({image_path})\n",
                encoding="utf-8",
            )

            import backend.export.md_to_word as md_to_word

            original_limit = md_to_word.MARKDOWN_IMAGE_MAX_COUNT
            md_to_word.MARKDOWN_IMAGE_MAX_COUNT = 1
            try:
                _, report = convert_md_to_word(markdown_path, return_report=True)
            finally:
                md_to_word.MARKDOWN_IMAGE_MAX_COUNT = original_limit

            self.assertEqual(report["found"], 2)
            self.assertEqual(report["inserted"], 1)
            self.assertEqual(report["skipped"], 1)

    def test_docx_reference_cover_and_header_omit_logo(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "logo.md"
            markdown_path.write_text("# 测试投标文件\n\n# 1. 企业简介\n\n正文内容。\n", encoding="utf-8")

            output_path, report = convert_md_to_word(markdown_path, return_report=True)

            self.assertFalse(report["logo"]["cover"]["inserted"])
            self.assertEqual("reference_template_cover_has_no_logo", report["logo"]["cover"]["reason"])
            self.assertFalse(report["logo"]["header"]["inserted"])
            self.assertEqual("formal_header_has_no_logo", report["logo"]["header"]["reason"])
            with ZipFile(output_path) as archive:
                media_names = [name for name in archive.namelist() if name.startswith("word/media/")]
                header_xml = "\n".join(
                    archive.read(name).decode("utf-8", errors="ignore")
                    for name in archive.namelist()
                    if name.startswith("word/header")
                )
            self.assertEqual([], media_names)
            self.assertNotIn("<w:drawing>", header_xml)

    def test_markdown_images_use_uniform_frame_without_cropping_source_content(self):
        from PIL import Image as PILImage

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            portrait_path = tmp / "full-page-scan.png"
            landscape_path = tmp / "factory-photo.png"
            PILImage.new("RGB", (800, 1600), "white").save(portrait_path)
            PILImage.new("RGB", (1600, 800), "white").save(landscape_path)
            markdown_path = tmp / "image-fit.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 图文测试",
                        "",
                        "# 1. 合同整页扫描件",
                        "",
                        f"![合同整页扫描件]({portrait_path})",
                        "",
                        f"![厂房实景照片]({landscape_path})",
                    ]
                ),
                encoding="utf-8",
            )

            output_path, report = convert_md_to_word(markdown_path, return_report=True)
            document = Document(str(output_path))
            body_shape_sizes = [
                (round(shape.width / 914400, 3), round(shape.height / 914400, 3))
                for shape in document.inline_shapes
                if shape.height
            ]

            self.assertEqual(report["inserted"], 2)
            self.assertEqual(1, len(set(body_shape_sizes)))
            self.assertEqual((DOCX_IMAGE_FRAME_WIDTH_IN, DOCX_IMAGE_FRAME_HEIGHT_IN), body_shape_sizes[0])
            self.assertTrue(all(event["fit"]["uniform_frame_enabled"] for event in report["events"]))
            self.assertTrue(all(event["fit"]["content_aspect_ratio_preserved"] for event in report["events"]))
            self.assertEqual(0.5, round(report["events"][0]["fit"]["source_aspect_ratio"], 3))
            self.assertEqual(2.0, round(report["events"][1]["fit"]["source_aspect_ratio"], 3))


if __name__ == "__main__":
    unittest.main()
