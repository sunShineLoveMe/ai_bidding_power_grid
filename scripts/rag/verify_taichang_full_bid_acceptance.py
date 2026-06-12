#!/usr/bin/env python3
"""Run customer-demo acceptance checks for the real Taichang full bid DOCX."""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from dotenv import load_dotenv
from PIL import Image

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RUNS_DIR = PROJECT_ROOT / "docs" / "development" / "runs"
PROJECT_ID = "4bc3ee73-9ec5-4184-aafd-eaede9f90798"
SUPPLEMENT_BATCH_ID = "customer_taichang_supplement_20260611"
FORBIDDEN_TEXT_TOKENS = [
    "```mermaid",
    "graph TD",
    "source_batch_id",
    "source_domain",
    "metadata",
    "匹配依据",
    "相似度",
    "rerank",
    "vector_score",
    "河北豪乾电气设备科技有限公司投标人",
]

sys.path.insert(0, str(PROJECT_ROOT))
load_dotenv(PROJECT_ROOT / ".env")


def _rel(path: Path) -> str:
    path = path if path.is_absolute() else PROJECT_ROOT / path
    return str(path.relative_to(PROJECT_ROOT))


def _read_zip_text(docx_path: Path, names: list[str]) -> str:
    with ZipFile(docx_path) as archive:
        parts = []
        for name in names:
            if name in archive.namelist():
                parts.append(archive.read(name).decode("utf-8", errors="ignore"))
        return "\n".join(parts)


def _docx_xml_parts(docx_path: Path) -> dict[str, str]:
    with ZipFile(docx_path) as archive:
        parts = {}
        for name in archive.namelist():
            if name == "word/document.xml" or name == "word/styles.xml" or name.startswith("word/header") or name.startswith("word/footer"):
                parts[name] = archive.read(name).decode("utf-8", errors="ignore")
        return parts


def _manifest_counts(manifest: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "by_batch": dict(Counter(item.get("source_batch_id") or "unknown" for item in manifest)),
        "by_evidence": dict(Counter(item.get("evidence_type") or "unknown" for item in manifest)),
        "by_batch_evidence": dict(Counter(
            f"{item.get('source_batch_id') or 'unknown'}|{item.get('evidence_type') or 'unknown'}"
            for item in manifest
        )),
    }


def _audit_docx_images(docx_path: Path) -> dict[str, Any]:
    ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    }
    with ZipFile(docx_path) as archive:
        media_sizes: dict[str, tuple[int, int]] = {}
        for name in archive.namelist():
            if not name.startswith("word/media/"):
                continue
            with archive.open(name) as handle:
                image = Image.open(handle)
                media_sizes[name] = image.size

        rel_maps: dict[str, dict[str, str]] = {}
        for name in archive.namelist():
            if not name.startswith("word/_rels/"):
                continue
            try:
                root = ET.fromstring(archive.read(name))
            except Exception:
                continue
            rel_maps[name] = {
                rel.attrib.get("Id"): "word/" + str(rel.attrib.get("Target") or "").lstrip("/")
                for rel in root
                if rel.attrib.get("Id") and rel.attrib.get("Target")
            }

        checked = 0
        crop_markers = 0
        max_ratio_delta = 0.0
        bad_ratio: list[dict[str, Any]] = []
        for part in archive.namelist():
            if not (part == "word/document.xml" or part.startswith("word/header")):
                continue
            xml = archive.read(part)
            crop_markers += xml.count(b"a:srcRect")
            try:
                root = ET.fromstring(xml)
            except Exception:
                continue
            rel_name = "word/_rels/document.xml.rels"
            if part.startswith("word/header"):
                rel_name = f"word/_rels/{Path(part).name}.rels"
            rels = rel_maps.get(rel_name, {})
            for inline in root.findall(".//wp:inline", ns):
                extent = inline.find("wp:extent", ns)
                blip = inline.find(".//a:blip", ns)
                if extent is None or blip is None:
                    continue
                rel_id = blip.attrib.get(f"{{{ns['r']}}}embed")
                media_name = rels.get(rel_id or "")
                if not media_name or media_name not in media_sizes:
                    continue
                cx = float(extent.attrib.get("cx") or 0)
                cy = float(extent.attrib.get("cy") or 0)
                if cx <= 0 or cy <= 0:
                    continue
                width, height = media_sizes[media_name]
                source_ratio = width / height
                display_ratio = cx / cy
                delta = abs(display_ratio - source_ratio) / source_ratio
                checked += 1
                max_ratio_delta = max(max_ratio_delta, delta)
                if delta > 0.01:
                    bad_ratio.append({
                        "part": part,
                        "media": media_name,
                        "pixel_width": width,
                        "pixel_height": height,
                        "ratio_delta": round(delta, 6),
                    })
        return {
            "media_count": len(media_sizes),
            "checked_inline_images": checked,
            "crop_marker_count": crop_markers,
            "max_ratio_delta": round(max_ratio_delta, 6),
            "bad_ratio": bad_ratio[:10],
        }


def _paragraph_summary(document: Document) -> dict[str, Any]:
    paragraphs = [p.text.strip() for p in document.paragraphs]
    non_empty = [text for text in paragraphs if text]
    headings = [
        p.text.strip()
        for p in document.paragraphs
        if p.text.strip() and p.style and p.style.name.startswith("Heading")
    ]
    toc_entries = [
        p.text.strip()
        for p in document.paragraphs
        if "\t" in p.text and p.text.strip() and p.text.strip()[0].isdigit()
    ]
    return {
        "paragraph_count": len(paragraphs),
        "non_empty_paragraph_count": len(non_empty),
        "heading_count": len(headings),
        "toc_entry_count": len(toc_entries),
        "first_non_empty": non_empty[:16],
        "heading_samples": headings[:20],
        "toc_samples": toc_entries[:20],
    }


def _table_audit(document: Document) -> dict[str, Any]:
    samples: list[dict[str, Any]] = []
    full_width = 0
    fixed_layout = 0
    repeat_header = 0
    with_margins = 0
    for table in document.tables:
        table_pr = table._tbl.tblPr
        table_width = table_pr.find(qn("w:tblW"))
        table_layout = table_pr.find(qn("w:tblLayout"))
        cell_margins = table_pr.find(qn("w:tblCellMar"))
        header_repeat = None
        if table.rows:
            header_repeat = table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader"))
        is_full = bool(table_width is not None and table_width.get(qn("w:w")) == "5000" and table_width.get(qn("w:type")) == "pct")
        is_fixed = bool(table_layout is not None and table_layout.get(qn("w:type")) == "fixed")
        is_repeat = bool(header_repeat is not None and header_repeat.get(qn("w:val")) == "true")
        has_margins = cell_margins is not None
        full_width += int(is_full)
        fixed_layout += int(is_fixed)
        repeat_header += int(is_repeat)
        with_margins += int(has_margins)
        if len(samples) < 8:
            samples.append({
                "rows": len(table.rows),
                "cols": len(table.columns),
                "full_width": is_full,
                "fixed_layout": is_fixed,
                "repeat_header": is_repeat,
                "cell_margins": has_margins,
            })
    count = len(document.tables)
    return {
        "table_count": count,
        "full_width_count": full_width,
        "fixed_layout_count": fixed_layout,
        "repeat_header_count": repeat_header,
        "cell_margin_count": with_margins,
        "all_tables_full_width": count > 0 and full_width == count,
        "all_tables_fixed_layout": count > 0 and fixed_layout == count,
        "all_tables_repeat_header": count > 0 and repeat_header == count,
        "all_tables_have_cell_margins": count > 0 and with_margins == count,
        "samples": samples,
    }


def _export_pdf(docx_path: Path) -> dict[str, Any]:
    from backend.export.md_to_word import _soffice_bin  # noqa: WPS433

    soffice_bin = _soffice_bin()
    report = {"enabled": bool(soffice_bin), "soffice_bin": soffice_bin, "status": "skipped", "pdf_path": None}
    if not soffice_bin:
        return report
    with tempfile.TemporaryDirectory(prefix="docx-pdf-preview-") as tmpdir:
        output_dir = Path(tmpdir)
        cmd = [
            soffice_bin,
            "--headless",
            "--invisible",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--nofirststartwizard",
            "--norestore",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(docx_path),
        ]
        completed = subprocess.run(cmd, capture_output=True, text=True, timeout=180, check=False)
        report.update({
            "returncode": completed.returncode,
            "stdout": (completed.stdout or "")[-1000:],
            "stderr": (completed.stderr or "")[-1000:],
        })
        generated = output_dir / f"{docx_path.stem}.pdf"
        if completed.returncode == 0 and generated.exists() and generated.stat().st_size > 0:
            pdf_path = docx_path.with_suffix(".pdf")
            if pdf_path.exists():
                pdf_path.unlink()
            generated.replace(pdf_path)
            report.update({"status": "generated", "pdf_path": _rel(pdf_path), "size": pdf_path.stat().st_size})
        else:
            report.update({"status": "failed", "reason": "LibreOffice did not produce PDF preview"})
    return report


def _validate(report: dict[str, Any]) -> tuple[list[str], list[str]]:
    failures: list[str] = []
    warnings: list[str] = []
    sections = report["source_sections"]
    docx = report["docx_audit"]
    image_selection = report["image_selection"]
    image_conversion = report["image_conversion"]
    field_refresh = report["field_refresh"]
    tables = docx["tables"]
    images = docx["images"]
    checks = docx["checks"]

    if sections["total"] < report["expected_min_sections"]:
        failures.append(f"章节数量不足：{sections['total']} < {report['expected_min_sections']}")
    if sections["non_empty"] < report["expected_min_non_empty_sections"]:
        failures.append(f"有正文的章节数量不足：{sections['non_empty']} < {report['expected_min_non_empty_sections']}")
    if image_selection.get("selected", 0) < 20:
        failures.append(f"正式图片选中数量不足：{image_selection.get('selected')}")
    if image_conversion.get("inserted", 0) < image_selection.get("selected", 0):
        failures.append("DOCX 图片插入数量少于选中数量")
    if image_conversion.get("failed", 0) != 0:
        failures.append(f"DOCX 图片插入失败：{image_conversion.get('failed')}")
    if field_refresh.get("status") != "refreshed":
        failures.append(f"LibreOffice 字段刷新失败：{field_refresh.get('status')}")
    if tables["table_count"] < 20:
        failures.append(f"表格数量异常偏少：{tables['table_count']}")
    for key in ["all_tables_full_width", "all_tables_fixed_layout", "all_tables_repeat_header", "all_tables_have_cell_margins"]:
        if not tables.get(key):
            failures.append(f"表格格式检查失败：{key}")
    if images["crop_marker_count"] != 0:
        failures.append(f"图片存在裁剪标记：{images['crop_marker_count']}")
    if images["bad_ratio"]:
        failures.append(f"图片比例变形：{images['bad_ratio']}")
    if report["pdf_preview"].get("enabled") and report["pdf_preview"].get("status") != "generated":
        warnings.append(f"PDF 预览未生成：{report['pdf_preview'].get('reason') or report['pdf_preview'].get('status')}")
    for key, passed in checks.items():
        if not passed:
            failures.append(f"DOCX 成品检查失败：{key}")
    return failures, warnings


def _write_report(report: dict[str, Any]) -> Path:
    RUNS_DIR.mkdir(parents=True, exist_ok=True)
    run_id = report["run_id"]
    json_path = RUNS_DIR / f"{run_id}.json"
    md_path = RUNS_DIR / f"{run_id}.md"
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    status = "PASS" if not report["failures"] else "FAIL"
    lines = [
        f"# {run_id} — 泰昌完整标书客户演示验收",
        "",
        f"- 生成时间：{report['generated_at']}",
        f"- 项目 ID：`{report['project_id']}`",
        f"- 状态：{status}",
        "",
        "## 输出文件",
        "",
        f"- Markdown：`{report['markdown_path']}`",
        f"- DOCX：`{report['docx_path']}`",
    ]
    if report["pdf_preview"].get("pdf_path"):
        lines.append(f"- PDF 预览：`{report['pdf_preview']['pdf_path']}`")
    lines.extend([
        "",
        "## 源项目与正文",
        "",
        "| 指标 | 数量 |",
        "| --- | ---: |",
        f"| 章节节点 | {report['source_sections']['total']} |",
        f"| 有正文章节 | {report['source_sections']['non_empty']} |",
        f"| Markdown 字符数 | {report['markdown_audit']['chars']} |",
        f"| DOCX 段落 | {report['docx_audit']['paragraphs']['non_empty_paragraph_count']} |",
        f"| DOCX 标题 | {report['docx_audit']['paragraphs']['heading_count']} |",
        f"| 目录条目 | {report['docx_audit']['paragraphs']['toc_entry_count']} |",
        "",
        "## 图片与表格",
        "",
        "| 指标 | 数量/结果 |",
        "| --- | ---: |",
        f"| 图片候选 | {report['image_selection'].get('asset_candidates')} |",
        f"| 图片选中 | {report['image_selection'].get('selected')} |",
        f"| 图片插入 | {report['image_conversion'].get('inserted')} |",
        f"| 图片失败 | {report['image_conversion'].get('failed')} |",
        f"| DOCX 媒体文件 | {report['docx_audit']['images']['media_count']} |",
        f"| 图片裁剪标记 | {report['docx_audit']['images']['crop_marker_count']} |",
        f"| 图片最大比例偏差 | {report['docx_audit']['images']['max_ratio_delta']} |",
        f"| 表格数量 | {report['docx_audit']['tables']['table_count']} |",
        f"| 表格全宽 | {report['docx_audit']['tables']['all_tables_full_width']} |",
        f"| 表格固定布局 | {report['docx_audit']['tables']['all_tables_fixed_layout']} |",
        f"| 表头跨页重复 | {report['docx_audit']['tables']['all_tables_repeat_header']} |",
        "",
        "## 封面/目录/页眉页脚",
        "",
        "| 检查项 | 结果 |",
        "| --- | --- |",
    ])
    for key, value in report["docx_audit"]["checks"].items():
        lines.append(f"| `{key}` | {'通过' if value else '失败'} |")
    lines.extend([
        "",
        "## 字段刷新",
        "",
        f"- LibreOffice 字段刷新：`{report['field_refresh'].get('status')}`",
        f"- 页码/总页数字段：`{report['docx_audit']['field_counts']}`",
        "",
        "## 结论",
        "",
    ])
    if report["failures"]:
        lines.extend(f"- FAIL：{failure}" for failure in report["failures"])
    else:
        lines.append("- 完整标书真实导出与成品结构验收通过，可作为客户演示版本。")
    if report["warnings"]:
        lines.extend(f"- WARN：{warning}" for warning in report["warnings"])
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return md_path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-id", default="run_20260612_taichang_full_bid_customer_acceptance")
    parser.add_argument("--project-id", default=PROJECT_ID)
    parser.add_argument("--expected-min-sections", type=int, default=30)
    parser.add_argument("--expected-min-non-empty-sections", type=int, default=30)
    parser.add_argument("--pdf-preview", action="store_true")
    args = parser.parse_args()

    os.environ["APP_AUTH_ENABLED"] = "false"
    os.environ["APP_LOGIN_ENABLED"] = "false"
    os.environ.setdefault("APP_ENV", "testing")

    import main as flask_main  # noqa: WPS433
    from backend.api.routes import build_project_bid_markdown  # noqa: WPS433
    from backend.db.supabase_repo import get_project_interpretation, list_bid_sections  # noqa: WPS433
    from backend.export.md_to_word import DOCX_BODY_EAST_ASIA, DOCX_HEADING_EAST_ASIA, convert_md_to_word, refresh_docx_fields_with_soffice  # noqa: WPS433

    with flask_main.app.app_context():
        sections = list_bid_sections(args.project_id)
        project_payload = get_project_interpretation(args.project_id)
        markdown_path, title, image_selection = build_project_bid_markdown(args.project_id, with_images=True)
        docx_path, image_conversion = convert_md_to_word(
            markdown_path,
            return_report=True,
            cover_fields=(image_selection or {}).get("cover_fields") or None,
        )
        docx_path, field_refresh = refresh_docx_fields_with_soffice(Path(docx_path))

    markdown_path = Path(markdown_path)
    docx_path = Path(docx_path)
    document = Document(str(docx_path))
    xml_parts = _docx_xml_parts(docx_path)
    document_xml = xml_parts.get("word/document.xml", "")
    styles_xml = xml_parts.get("word/styles.xml", "")
    all_xml = "\n".join(xml_parts.values())
    first_page_header_is_empty = 'w:type="first"' in document_xml and not any(
        text.strip()
        for section in document.sections
        for text in [p.text for p in section.first_page_header.paragraphs]
    )
    all_text = "\n".join(p.text for p in document.paragraphs)
    headers = "\n".join(p.text for section in document.sections for p in section.header.paragraphs)
    footers = "\n".join(p.text for section in document.sections for p in section.footer.paragraphs)
    field_codes = re.findall(r"<w:instrText[^>]*>\s*([^<]+?)\s*</w:instrText>", all_xml)
    manifest = image_selection.get("manifest") or []
    markdown_text = markdown_path.read_text(encoding="utf-8")
    forbidden_hits = [token for token in FORBIDDEN_TEXT_TOKENS if token in all_xml or token in all_text]
    repeated_title_pattern = re.findall(r"\d+(?:\.\d+)+\s+[^。\n\t]{2,30}\s+-\s+[^。\n\t]{2,30}", all_text)
    section_text_lengths = [len(str(section.get("content") or "")) for section in sections]
    report = {
        "run_id": args.run_id,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "project_id": args.project_id,
        "project_name": ((project_payload.get("project") or {}).get("project_name")),
        "document_title": title,
        "expected_min_sections": args.expected_min_sections,
        "expected_min_non_empty_sections": args.expected_min_non_empty_sections,
        "markdown_path": _rel(markdown_path),
        "docx_path": _rel(docx_path),
        "docx_size": docx_path.stat().st_size,
        "source_sections": {
            "total": len(sections),
            "non_empty": sum(1 for length in section_text_lengths if length > 0),
            "total_content_chars": sum(section_text_lengths),
            "empty_titles": [section.get("title") for section in sections if not str(section.get("content") or "").strip()][:30],
        },
        "markdown_audit": {
            "chars": len(markdown_text),
            "heading_count": len(re.findall(r"^#", markdown_text, flags=re.MULTILINE)),
            "image_refs": len(re.findall(r"^!\[", markdown_text, flags=re.MULTILINE)),
            "table_blocks": len(re.findall(r"^\|.+\|$", markdown_text, flags=re.MULTILINE)),
        },
        "image_selection": image_selection,
        "image_conversion": image_conversion,
        "field_refresh": field_refresh,
        "manifest_counts": _manifest_counts(manifest),
        "pdf_preview": _export_pdf(docx_path) if args.pdf_preview else {"enabled": False, "status": "skipped"},
        "docx_audit": {
            "configured_fonts": {
                "body_east_asia": DOCX_BODY_EAST_ASIA,
                "heading_east_asia": DOCX_HEADING_EAST_ASIA,
            },
            "paragraphs": _paragraph_summary(document),
            "tables": _table_audit(document),
            "images": _audit_docx_images(docx_path),
            "field_counts": {
                "instrText": len(field_codes),
                "PAGEREF": sum(1 for code in field_codes if "PAGEREF" in code),
                "PAGE": sum(1 for code in field_codes if code.strip() == "PAGE"),
                "NUMPAGES": sum(1 for code in field_codes if code.strip() == "NUMPAGES"),
            },
            "checks": {
                "docx_opens_with_python_docx": True,
                "cover_has_bid_title": "投标文件" in "\n".join(_paragraph_summary(document)["first_non_empty"][:6]),
                "cover_has_bidder": "投标人：河北泰昌电力器材科技有限公司" in "\n".join(_paragraph_summary(document)["first_non_empty"][:8]),
                "cover_has_project_name": "国网辽宁电力2025年第三次物资协议库存招标采购" in all_text,
                "cover_has_tender_no": "2225AC" in all_text,
                "cover_first_page_header_empty": first_page_header_is_empty,
                "toc_title_exists": "目  录" in all_text,
                "toc_has_entries": _paragraph_summary(document)["toc_entry_count"] > 0,
                "toc_has_dot_leader": 'w:leader="dot"' in document_xml,
                "toc_has_pageref": any("PAGEREF" in code for code in field_codes),
                "header_has_taichang_bid": "河北泰昌电力器材科技有限公司投标文件" in headers,
                "footer_has_page_text": "第 " in footers and " 页，共 " in footers,
                "footer_has_page_fields": any(code.strip() == "PAGE" for code in field_codes) and any(code.strip() == "NUMPAGES" for code in field_codes),
                "uses_configured_cjk_fonts": DOCX_BODY_EAST_ASIA in all_xml and DOCX_HEADING_EAST_ASIA in all_xml,
                "no_forbidden_internal_tokens": not forbidden_hits,
                "no_repeated_parent_title_pattern": not repeated_title_pattern,
                "no_black_square_markers": all(token not in document_xml and token not in styles_xml for token in ("w:keepLines", "w:keepNext", "w:pageBreakBefore")),
                "no_image_crop_or_distortion": _audit_docx_images(docx_path)["crop_marker_count"] == 0 and not _audit_docx_images(docx_path)["bad_ratio"],
                "has_supplement_project_performance_assets": (Counter(f"{item.get('source_batch_id')}|{item.get('evidence_type')}" for item in manifest).get(f"{SUPPLEMENT_BATCH_ID}|project_performance", 0) >= 2),
                "has_supplement_testing_assets": (Counter(f"{item.get('source_batch_id')}|{item.get('evidence_type')}" for item in manifest).get(f"{SUPPLEMENT_BATCH_ID}|testing_capacity", 0) >= 1),
                "has_supplement_inspection_report_asset": (Counter(f"{item.get('source_batch_id')}|{item.get('evidence_type')}" for item in manifest).get(f"{SUPPLEMENT_BATCH_ID}|inspection_report", 0) >= 1),
            },
            "forbidden_hits": forbidden_hits,
            "repeated_title_pattern_samples": repeated_title_pattern[:20],
        },
    }
    failures, warnings = _validate(report)
    report["failures"] = failures
    report["warnings"] = warnings
    md_path = _write_report(report)
    print(json.dumps({
        "report": _rel(md_path),
        "docx": report["docx_path"],
        "pdf": report["pdf_preview"].get("pdf_path"),
        "failures": failures,
        "warnings": warnings,
    }, ensure_ascii=False, indent=2))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
