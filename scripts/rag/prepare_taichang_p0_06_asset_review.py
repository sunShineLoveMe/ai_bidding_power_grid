#!/usr/bin/env python3
"""生成并校验泰昌历史标书 P0-06 分级接收与异常复核包。

本脚本只生成离线审批材料，不连接数据库、不修改 RAG、不改变 DOCX 选图。
客户主动提供两份历史标书即视为授权系统处理来源文件。低风险、来源明确且不含
精确时效事实的资料可按策略自动接收为 ``knowledge_only``；只有异常项需要人工
确认。自动接收不等于 ``formal_bid_ready``，也不允许直接进入正式标书。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = PROJECT_ROOT / "docs/development/taichang-bid-v1-data"
DEFAULT_MATRIX = DATA_DIR / "asset_classification_matrix.json"
DEFAULT_REVIEW_QUEUE = DATA_DIR / "asset_review_queue.json"
DEFAULT_OUTPUT_DIR = DATA_DIR / "p0_06_review"
DEFAULT_WORKBOOK = DEFAULT_OUTPUT_DIR / "泰昌历史标书资产增量入库审批表.xlsx"

ENTERPRISE = "河北泰昌电力器材科技有限公司"
DECISIONS = ("批准为知识资料", "批准为正式资产", "延后补证", "禁止入库", "保留历史参考")
YES_NO_NA = ("是", "否", "不适用")
DEDUP_CONFIRMATIONS = ("确认新资产", "确认重复", "同证据不同载体", "待进一步核验")
TARGET_LIBRARY_LABELS = ("知识库资料", "产品库资料", "资信库资料")
SENSITIVE_LEVELS = {"sensitive", "restricted"}
AUTO_KNOWLEDGE_EVIDENCE_TYPES = {"enterprise_evidence", "production_capacity", "green_low_carbon"}
AUTO_KNOWLEDGE_SECTION_KEYWORDS = (
    "工艺流程",
    "各工序控制点工艺文件",
    "生产制造环境",
    "试组装环境",
    "绿色低碳生产及绿色回收",
    "ESG",
    "废水、废气、废固",
    "现状环境影响评估报告",
    "绿色发展规划报告",
    "数字领航企业评价报告",
    "智能制造优秀场景报告",
    "创新激励机制评价报告",
)
AUTO_KNOWLEDGE_BLOCKED_SECTION_KEYWORDS = (
    "证书",
    "凭证",
    "查询",
    "绩效评价",
    "股权",
    "合同",
    "中标",
    "授权",
    "人员",
    "审计",
    "保证金",
)
TIMELINESS_EVIDENCE_TYPES = {
    "business_license",
    "certification",
    "finance",
    "personnel_certificate",
    "project_performance",
    "authorization",
    "inspection_report",
    "testing_capacity",
}
PRODUCT_EVIDENCE_TYPES = {
    "inspection_report",
    "product_parameter_table",
    "production_capacity",
    "testing_capacity",
    "product_image",
}
FORMAL_ORIGINAL_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg", ".webp", ".doc", ".docx", ".xls", ".xlsx", ".csv"}
FORBIDDEN_DISPLAY_PATTERNS = (
    "taichang_", "power_grid_", "production_capacity", "testing_capacity", "green_low_carbon",
    "business_license", "certification", "product_image", "technical", "页面_", "原图",
    "parsed_outputs", "rag_seed", "/api/",
)
UUID_RE = re.compile(r"\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b", re.I)
HASH_RE = re.compile(r"\b[0-9a-f]{24,64}\b", re.I)

GROUP_FIELDS = [
    "证据组编号",
    "来源标书",
    "历史章节",
    "证据类型",
    "建议目标库",
    "产品族",
    "候选数量",
    "去重状态",
    "敏感级别",
    "机器复核建议",
    "必须核验事项",
]

CANDIDATE_FIELDS = [
    "候选编号",
    "证据组编号",
    "来源标书",
    "历史页序",
    "历史章节",
    "资料名称",
    "证据类型",
    "建议目标库",
    "产品族",
    "去重状态",
    "敏感级别",
    "机器复核建议",
    "必须核验事项",
    "审批结论",
    "去重结论确认",
    "复核后资料名称",
    "复核后目标库",
    "产品边界已核验",
    "原始证据已核验",
    "原始证据路径",
    "有效性时效已核验",
    "敏感资料授权",
    "正式展示质量已确认",
    "复核人",
    "复核日期",
    "复核意见",
]

BLOCKED_FIELDS = [
    "候选编号",
    "来源标书",
    "历史页序",
    "历史章节",
    "资料名称",
    "证据类型",
    "去重状态",
    "处理结论",
    "处理原因",
]

AUTO_ACCEPTED_FIELDS = [
    "候选编号",
    "来源标书",
    "历史章节",
    "资料名称",
    "证据类型",
    "目标库",
    "质量等级",
    "处理结论",
    "使用限制",
]


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _join(values: Any) -> str:
    if isinstance(values, list):
        return "、".join(_text(value) for value in values if _text(value))
    return _text(values)


def _is_expired_ohs_certificate(row: dict[str, Any]) -> bool:
    section = _text(row.get("source_section"))
    return "职业健康安全管理体系认证证书" in section


def _review_guidance(row: dict[str, Any]) -> tuple[str, str]:
    dedup_status = _text(row.get("dedup_status"))
    evidence_type = _text(row.get("evidence_type"))
    sensitivity = _text(row.get("sensitivity"))
    if _is_expired_ohs_certificate(row):
        return "延后补证", "历史证书已于2026年6月18日到期，必须提供有效续证，禁止正式引用旧证"
    if dedup_status.startswith("possible_"):
        return "优先核重", "查看原件并确认是新资产、重复资产或同一证据的不同载体"
    if dedup_status == "same_evidence_new_rendition":
        return "核验证据包", "不得新建重复业务主记录，仅可在确认后关联为同一证据的不同载体"
    if sensitivity in SENSITIVE_LEVELS:
        return "延后授权", "身份、银行、财务、社保、股权或授权资料必须确认权限和使用范围"
    if evidence_type == "product_parameter_table":
        return "优先结构化核验", "核对产品族、规格型号、参数值和原始表格，批准后仍须走结构化抽取"
    if evidence_type == "inspection_report":
        return "优先核验原始报告", "核对报告编号、产品、规格、检验值、日期和完整原始报告"
    if evidence_type == "certification":
        return "优先核验有效期", "核对证书编号、认证范围、签发日期、失效日期和续证情况"
    if evidence_type == "testing_capacity":
        return "核验最新校准证书", "以可追溯的较新原始校准证书为准，不采用历史台账旧日期"
    if evidence_type in {"production_capacity", "green_low_carbon"}:
        return "核验整份原始资料", "连续内嵌页应按完整报告或证据包审核，不逐页新建业务资产"
    return "核验原始证据", "确认资料归属泰昌、内容真实、时效有效、中文展示合规且目标库正确"


def _auto_accept_knowledge_candidate(row: dict[str, Any]) -> bool:
    """只自动接收明确的低风险知识资料，精确事实和正式证据继续人工复核。"""
    section = _text(row.get("source_section"))
    return all((
        _text(row.get("record_type")) == "media",
        _text(row.get("dedup_status")) == "new",
        _text(row.get("sensitivity")) not in SENSITIVE_LEVELS,
        _text(row.get("evidence_type")) in AUTO_KNOWLEDGE_EVIDENCE_TYPES,
        any(keyword in section for keyword in AUTO_KNOWLEDGE_SECTION_KEYWORDS),
        not any(keyword in section for keyword in AUTO_KNOWLEDGE_BLOCKED_SECTION_KEYWORDS),
        not _is_expired_ohs_certificate(row),
    ))


def _knowledge_title(row: dict[str, Any], sequence: int) -> str:
    leaf = _text(row.get("source_section")).split("/")[-1].strip()
    leaf = re.sub(r"^[（(]?\d+(?:\.\d+)*[）).、]?\s*", "", leaf).strip()
    label = leaf or _text(row.get("evidence_type_label")) or "企业资料"
    return f"泰昌{label}知识资料第{sequence:03d}项"


def _page_sort_value(value: Any) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _auto_accepted_knowledge_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    accepted: list[dict[str, Any]] = []
    section_sequences: Counter[str] = Counter()
    for row in sorted(
        rows,
        key=lambda item: (
            _text(item.get("source_document_name")),
            _page_sort_value(item.get("source_page")),
            item["candidate_id"],
        ),
    ):
        section = _text(row.get("source_section"))
        section_sequences[section] += 1
        accepted.append({
            **row,
            "origin_source_domain": row.get("source_domain"),
            "source_domain": "enterprise_fact",
            "source_authorized": True,
            "source_authorization_basis": "客户主动提供历史技术标/商务标用于系统整理与复用",
            "fact_source_allowed_for_enterprise": True,
            "quality_tier_after_review": "knowledge_only",
            "review_status_after_review": "policy_auto_accepted",
            "decision": "系统自动接收为知识资料",
            "reviewed_title": _knowledge_title(row, section_sequences[section]),
            "reviewed_target_library_label": row.get("target_library_label", "知识库资料"),
            "allowed_for_bid": False,
            "formal_bid_ready": False,
            "ingestion_action": "extract_word_media_as_knowledge_only",
            "ingestion_readiness": "ready_for_extraction_and_validation",
            "ready_for_ingestion": True,
            "requires_fact_cross_check_for_precise_values": True,
            "usage_restriction": "仅作泰昌内部知识资料；不得作为精确参数、有效证书或正式投标附件直接引用",
            "reviewer": "系统分级策略",
            "reviewed_at": date.today().isoformat(),
        })
    return accepted


def _auto_blocked_row(row: dict[str, Any], in_review_queue: bool) -> dict[str, Any] | None:
    status = _text(row.get("dedup_status"))
    review_status = _text(row.get("review_status"))
    if status == "duplicate_exact":
        conclusion, reason = "关联已有资产", "与现有数字资产精确重复，无需再次入库"
    elif status == "same_evidence_new_rendition":
        conclusion, reason = "关联同一证据包", "已确认属于既有报告的不同载体，不建立第二个业务主记录"
    elif status == "fact_conflict":
        conclusion, reason = "禁止复用", "历史项目字段或事实值与当前基线冲突"
    elif review_status == "not_asset_reference":
        conclusion, reason = "保留历史参考", "章节或表单只用于历史结构对照，不属于泰昌企业事实资产"
    elif not in_review_queue:
        conclusion, reason = "禁止独立入库", "通用响应话术或缺少独立证据意义，不作为正式资产"
    else:
        return None
    return {
        "candidate_id": row["candidate_id"],
        "source_document_name": row.get("source_document_name", ""),
        "source_page": row.get("source_page", ""),
        "source_section": row.get("source_section", ""),
        "title": row.get("title", ""),
        "evidence_type_label": row.get("evidence_type_label", ""),
        "dedup_status_label": row.get("dedup_status_label", ""),
        "disposition": conclusion,
        "reason": reason,
        "ready_for_ingestion": False,
    }


def build_review_package(matrix: dict[str, Any], review_queue: dict[str, Any]) -> dict[str, Any]:
    original_review_records = review_queue.get("records") or []
    review_ids = {_text(row.get("candidate_id")) for row in original_review_records}
    auto_accepted_source_rows = [row for row in original_review_records if _auto_accept_knowledge_candidate(row)]
    auto_accepted_ids = {row["candidate_id"] for row in auto_accepted_source_rows}
    auto_linked_ids = {
        row["candidate_id"] for row in original_review_records
        if _text(row.get("dedup_status")) == "same_evidence_new_rendition"
    }
    review_records = [
        row for row in original_review_records
        if row["candidate_id"] not in auto_accepted_ids | auto_linked_ids
    ]
    auto_accepted = _auto_accepted_knowledge_rows(auto_accepted_source_rows)
    blocked = [
        blocked_row
        for row in (matrix.get("records") or [])
        if (blocked_row := _auto_blocked_row(row, _text(row.get("candidate_id")) in review_ids)) is not None
    ]

    group_keys = sorted({
        (
            _text(row.get("source_document_name")),
            _text(row.get("source_section")),
            _text(row.get("evidence_type")),
            _join(row.get("product_families")),
            _text(row.get("target_library")),
        )
        for row in review_records
    })
    group_ids = {key: f"P006-G{index:03d}" for index, key in enumerate(group_keys, start=1)}
    grouped: dict[tuple[str, str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    candidate_rows: list[dict[str, Any]] = []
    for row in review_records:
        key = (
            _text(row.get("source_document_name")),
            _text(row.get("source_section")),
            _text(row.get("evidence_type")),
            _join(row.get("product_families")),
            _text(row.get("target_library")),
        )
        grouped[key].append(row)
        guidance, required_check = _review_guidance(row)
        candidate_rows.append({
            **row,
            "review_group_id": group_ids[key],
            "machine_recommendation": guidance,
            "required_check": required_check,
            "decision": "",
            "dedup_confirmation": "",
            "reviewed_title": row.get("title", ""),
            "reviewed_target_library_label": row.get("target_library_label", ""),
            "product_boundary_verified": "否" if _text(row.get("evidence_type")) in PRODUCT_EVIDENCE_TYPES else "不适用",
            "original_evidence_verified": "否",
            "original_evidence_path": "",
            "timeliness_verified": "否" if _text(row.get("evidence_type")) in TIMELINESS_EVIDENCE_TYPES else "不适用",
            "sensitive_authorized": "否" if _text(row.get("sensitivity")) in SENSITIVE_LEVELS else "不适用",
            "formal_display_quality_verified": "否",
            "reviewer": "",
            "reviewed_at": "",
            "review_notes": "",
        })

    groups: list[dict[str, Any]] = []
    for key in group_keys:
        rows = grouped[key]
        recommendations = sorted({_review_guidance(row)[0] for row in rows})
        checks = sorted({_review_guidance(row)[1] for row in rows})
        groups.append({
            "review_group_id": group_ids[key],
            "source_document_name": key[0],
            "source_section": key[1],
            "evidence_type": key[2],
            "evidence_type_label": rows[0].get("evidence_type_label", ""),
            "target_library": key[4],
            "target_library_label": rows[0].get("target_library_label", ""),
            "product_families": rows[0].get("product_families") or [],
            "candidate_count": len(rows),
            "dedup_statuses": sorted({_text(row.get("dedup_status_label")) for row in rows}),
            "sensitivity_labels": sorted({_text(row.get("sensitivity_label")) for row in rows}),
            "machine_recommendations": recommendations,
            "required_checks": checks,
        })

    auto_linked_existing_count = sum(
        row["disposition"] in {"关联已有资产", "关联同一证据包"} for row in blocked
    )
    counts = {
        "candidate_count": len(matrix.get("records") or []),
        "source_authorized_count": len(matrix.get("records") or []),
        "original_review_queue_count": len(original_review_records),
        "policy_auto_accepted_count": len(auto_accepted),
        "review_candidate_count": len(candidate_rows),
        "review_group_count": len(groups),
        "auto_blocked_or_reference_count": len(blocked),
        "auto_linked_existing_count": auto_linked_existing_count,
        "human_approved_decision_count": 0,
        "approved_decision_count": len(auto_accepted),
        "ready_for_ingestion_count": len(auto_accepted),
        "validation_error_count": 0,
        "decision_counts": {"系统自动接收为知识资料": len(auto_accepted)} if auto_accepted else {},
    }
    return {
        "metadata": {
            "schema_version": "taichang_p0_06_review_v2",
            "generated_at": datetime.now(timezone.utc).astimezone().isoformat(),
            "enterprise": ENTERPRISE,
            "pilot_only": True,
            "read_only": True,
            "database_written": False,
            "rag_updated": False,
            "docx_selection_updated": False,
            **counts,
        },
        "groups": groups,
        "review_candidates": sorted(candidate_rows, key=lambda row: (row["review_group_id"], row["candidate_id"])),
        "auto_blocked_or_reference": sorted(blocked, key=lambda row: row["candidate_id"]),
        "policy_auto_accepted": auto_accepted,
        "approved_decisions": list(auto_accepted),
        "ready_for_ingestion": list(auto_accepted),
        "validation_errors": [],
    }


def _approval_errors(row: dict[str, Any]) -> list[str]:
    decision = _text(row.get("decision"))
    if not decision:
        return []
    errors: list[str] = []
    if decision not in DECISIONS:
        return ["审批结论不在允许范围内"]
    if decision not in {"批准为知识资料", "批准为正式资产"}:
        return errors
    required = {
        "dedup_confirmation": "必须确认去重结论",
        "reviewed_title": "必须填写复核后资料名称",
        "reviewed_target_library_label": "必须确认目标库",
        "original_evidence_verified": "必须核验原始证据",
        "reviewer": "必须填写复核人",
        "reviewed_at": "必须填写复核日期",
    }
    for field, message in required.items():
        if not _text(row.get(field)):
            errors.append(message)
    dedup_confirmation = _text(row.get("dedup_confirmation"))
    if dedup_confirmation not in {"确认新资产", "同证据不同载体"}:
        errors.append("批准项的去重结论必须为确认新资产或同证据不同载体")
    if _text(row.get("original_evidence_verified")) != "是":
        errors.append("批准项必须确认原始证据已核验")
    if _text(row.get("reviewed_target_library_label")) not in TARGET_LIBRARY_LABELS:
        errors.append("复核后目标库不合法")
    reviewed_title = _text(row.get("reviewed_title"))
    lowered_title = reviewed_title.lower()
    if not any("\u4e00" <= char <= "\u9fff" for char in reviewed_title):
        errors.append("复核后资料名称必须使用中文业务名称")
    if any(pattern.lower() in lowered_title for pattern in FORBIDDEN_DISPLAY_PATTERNS):
        errors.append("复核后资料名称包含内部枚举、解析路径或禁用表达")
    if UUID_RE.search(reviewed_title) or HASH_RE.search(reviewed_title):
        errors.append("复核后资料名称不得包含UUID或长哈希")
    reviewed_at = _text(row.get("reviewed_at"))
    if reviewed_at:
        try:
            reviewed_date = date.fromisoformat(reviewed_at[:10])
            if reviewed_date > date.today():
                errors.append("复核日期不得晚于当前日期")
        except ValueError:
            errors.append("复核日期必须为有效日期")
    if _text(row.get("sensitivity")) in SENSITIVE_LEVELS or _text(row.get("quality_tier")) == "restricted":
        errors.append("受限或敏感资料不得通过本批自动生成入库候选")
    if _text(row.get("evidence_type")) in PRODUCT_EVIDENCE_TYPES:
        if _text(row.get("product_boundary_verified")) != "是":
            errors.append("产品类资料必须确认产品边界")
        if not _join(row.get("product_families")):
            errors.append("产品类资料缺少明确产品族")
    if _text(row.get("evidence_type")) in TIMELINESS_EVIDENCE_TYPES:
        if _text(row.get("timeliness_verified")) != "是":
            errors.append("证书、报告或时效资料必须确认有效性")
    if _is_expired_ohs_certificate(row):
        errors.append("已知过期职业健康安全管理体系证书禁止批准")
    if decision == "批准为正式资产":
        if _text(row.get("formal_display_quality_verified")) != "是":
            errors.append("正式资产必须确认展示质量")
        evidence_path = _text(row.get("original_evidence_path"))
        if not evidence_path:
            errors.append("正式资产必须填写独立原始证据路径")
        else:
            path = Path(evidence_path).expanduser()
            if not path.is_absolute():
                path = PROJECT_ROOT / path
            if not path.is_file():
                errors.append("正式资产的原始证据路径不存在")
            elif path.suffix.lower() not in FORMAL_ORIGINAL_EXTENSIONS:
                errors.append("正式资产的原始证据文件类型不受支持")
            source_path = Path(_text(row.get("source_file"))).expanduser()
            if not source_path.is_absolute():
                source_path = PROJECT_ROOT / source_path
            if path.resolve() == source_path.resolve():
                errors.append("历史标书本身不能替代独立原始证据文件")
        if _text(row.get("record_type")) != "media":
            errors.append("文本/参数候选不得直接提升为正式展示资产")
    return list(dict.fromkeys(errors))


def apply_review_decisions(package: dict[str, Any], decisions: dict[str, dict[str, Any]]) -> dict[str, Any]:
    approved: list[dict[str, Any]] = list(package.get("policy_auto_accepted") or [])
    ready: list[dict[str, Any]] = list(package.get("policy_auto_accepted") or [])
    errors: list[dict[str, Any]] = []
    decision_counts: Counter[str] = Counter({"系统自动接收为知识资料": len(approved)})
    human_approved_count = 0
    expected_ids = {row["candidate_id"] for row in package["review_candidates"]}
    actual_ids = set(decisions)
    missing_ids = sorted(expected_ids - actual_ids)
    unknown_ids = sorted(actual_ids - expected_ids)
    if missing_ids:
        errors.append({"candidate_id": "审批表完整性", "errors": [f"缺少{len(missing_ids)}个候选编号"]})
    if unknown_ids:
        errors.append({"candidate_id": "审批表完整性", "errors": [f"存在{len(unknown_ids)}个未知候选编号"]})
    integrity_ok = not missing_ids and not unknown_ids
    for base_row in package["review_candidates"]:
        decision = decisions.get(base_row["candidate_id"], {})
        row = {**base_row, **decision}
        decision_name = _text(row.get("decision"))
        if decision_name:
            decision_counts[decision_name] += 1
        row_errors = _approval_errors(row)
        if row_errors:
            errors.append({"candidate_id": row["candidate_id"], "errors": row_errors})
            continue
        if decision_name not in {"批准为知识资料", "批准为正式资产"}:
            continue
        quality_tier = "formal_bid_ready" if decision_name == "批准为正式资产" else "knowledge_only"
        action = "structured_extraction_required" if row["evidence_type"] == "product_parameter_table" else (
            "attach_rendition" if row["dedup_confirmation"] == "同证据不同载体" else "create_new_asset"
        )
        reviewed = {
            **row,
            "quality_tier_after_review": quality_tier,
            "review_status_after_review": "approved",
            "fact_source_allowed_for_enterprise": True,
            "source_domain": "enterprise_fact",
            "ingestion_action": action,
            "ready_for_ingestion": action in {"create_new_asset", "attach_rendition"},
        }
        if integrity_ok:
            approved.append(reviewed)
            human_approved_count += 1
            if reviewed["ready_for_ingestion"]:
                ready.append(reviewed)

    result = dict(package)
    result["approved_decisions"] = approved
    result["ready_for_ingestion"] = ready
    result["validation_errors"] = errors
    result["metadata"] = {
        **package["metadata"],
        "approved_decision_count": len(approved),
        "human_approved_decision_count": human_approved_count,
        "ready_for_ingestion_count": len(ready),
        "validation_error_count": len(errors),
        "decision_counts": dict(sorted(decision_counts.items())),
    }
    return result


def _style_sheet(ws, widths: dict[int, int] | None = None) -> None:
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = dark_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if widths:
        for index, width in widths.items():
            ws.column_dimensions[get_column_letter(index)].width = width


def _add_list_validation(ws, column: int, values: tuple[str, ...], start: int, end: int) -> None:
    validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    validation.error = "请选择下拉列表中的值"
    validation.errorTitle = "审批值不合法"
    validation.prompt = "请从下拉列表选择，不要自行输入其他值"
    validation.promptTitle = "审批填写提示"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    ws.add_data_validation(validation)
    validation.add(f"{get_column_letter(column)}{start}:{get_column_letter(column)}{end}")


def write_workbook(path: Path, package: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "审批说明"
    ws.sheet_view.showGridLines = False
    title_fill = PatternFill("solid", fgColor="1F4E78")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    ws.merge_cells("A1:F1")
    ws["A1"] = "泰昌历史标书资产增量入库审批表"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center")
    instructions = [
        ("适用企业", ENTERPRISE),
        ("来源授权候选", package["metadata"]["source_authorized_count"]),
        ("系统自动接收", package["metadata"]["policy_auto_accepted_count"]),
        ("异常复核候选", package["metadata"]["review_candidate_count"]),
        ("证据分组", package["metadata"]["review_group_count"]),
        ("自动阻断/仅参考", package["metadata"]["auto_blocked_or_reference_count"]),
        ("当前可进入处理流程", package["metadata"]["ready_for_ingestion_count"]),
        ("填写顺序", "“系统自动接收”无需逐项审批；只需在“异常候选复核”处理证照、财务、人员、合同、参数值等风险项"),
        ("人工批准前提", "必须核验原始证据、去重、产品边界、有效性、中文展示和复核责任人"),
        ("严格限制", "受限/敏感资料、已过期证书、未确认疑似重复、产品不明项不得批准"),
        ("数据边界", "本表不会直接写数据库；自动接收项仅可进入媒体提取和质量校验，不等于正式投标可用资产"),
    ]
    for index, (label, value) in enumerate(instructions, start=3):
        ws.cell(index, 1, label).font = Font(bold=True)
        ws.cell(index, 2, value)
        ws.merge_cells(start_row=index, start_column=2, end_row=index, end_column=6)
        ws.cell(index, 2).alignment = Alignment(wrap_text=True, vertical="top")
        if label in {"严格限制", "数据边界"}:
            ws.cell(index, 1).fill = warning_fill
            ws.cell(index, 2).fill = warning_fill
    ws.column_dimensions["A"].width = 18
    for column in "BCDEF":
        ws.column_dimensions[column].width = 22

    group_ws = wb.create_sheet("证据分组")
    group_ws.append(GROUP_FIELDS)
    for group in package["groups"]:
        group_ws.append([
            group["review_group_id"], group["source_document_name"], group["source_section"],
            group["evidence_type_label"], group["target_library_label"], _join(group["product_families"]),
            group["candidate_count"], _join(group["dedup_statuses"]), _join(group["sensitivity_labels"]),
            _join(group["machine_recommendations"]), _join(group["required_checks"]),
        ])
    _style_sheet(group_ws, {1: 14, 2: 18, 3: 55, 4: 18, 5: 16, 6: 22, 7: 12, 8: 26, 9: 16, 10: 22, 11: 55})

    auto_ws = wb.create_sheet("系统自动接收")
    auto_ws.append(AUTO_ACCEPTED_FIELDS)
    for row in package["policy_auto_accepted"]:
        auto_ws.append([
            row["candidate_id"], row["source_document_name"], row["source_section"], row["reviewed_title"],
            row["evidence_type_label"], row["reviewed_target_library_label"],
            row["quality_tier_after_review"], row["decision"], row["usage_restriction"],
        ])
    _style_sheet(auto_ws, {1: 24, 2: 18, 3: 55, 4: 38, 5: 18, 6: 16, 7: 18, 8: 24, 9: 60})
    for cell in auto_ws["A"][1:]:
        cell.number_format = "@"

    candidate_ws = wb.create_sheet("异常候选复核")
    candidate_ws.append(CANDIDATE_FIELDS)
    for row in package["review_candidates"]:
        candidate_ws.append([
            row["candidate_id"], row["review_group_id"], row["source_document_name"], row["source_page"],
            row["source_section"], row["title"], row["evidence_type_label"], row["target_library_label"],
            _join(row["product_families"]), row["dedup_status_label"], row["sensitivity_label"],
            row["machine_recommendation"], row["required_check"], row["decision"], row["dedup_confirmation"],
            row["reviewed_title"], row["reviewed_target_library_label"], row["product_boundary_verified"],
            row["original_evidence_verified"], row["original_evidence_path"], row["timeliness_verified"],
            row["sensitive_authorized"], row["formal_display_quality_verified"], row["reviewer"],
            row["reviewed_at"], row["review_notes"],
        ])
    widths = {1: 24, 2: 14, 3: 18, 4: 10, 5: 55, 6: 30, 7: 18, 8: 16, 9: 22, 10: 20,
              11: 16, 12: 22, 13: 55, 14: 20, 15: 22, 16: 30, 17: 18, 18: 18, 19: 18,
              20: 45, 21: 20, 22: 18, 23: 22, 24: 16, 25: 14, 26: 40}
    _style_sheet(candidate_ws, widths)
    candidate_ws.freeze_panes = "A2"
    end_row = max(2, candidate_ws.max_row)
    _add_list_validation(candidate_ws, 14, DECISIONS, 2, end_row)
    _add_list_validation(candidate_ws, 15, DEDUP_CONFIRMATIONS, 2, end_row)
    _add_list_validation(candidate_ws, 17, TARGET_LIBRARY_LABELS, 2, end_row)
    for column in (18, 19, 21, 22, 23):
        _add_list_validation(candidate_ws, column, YES_NO_NA, 2, end_row)
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")
    for row in candidate_ws.iter_rows(min_row=2, min_col=14, max_col=26):
        for cell in row:
            cell.fill = yellow_fill
    candidate_ws.conditional_formatting.add(
        f"N2:N{end_row}",
        FormulaRule(formula=["OR(N2=\"批准为知识资料\",N2=\"批准为正式资产\")"],
                    fill=PatternFill("solid", fgColor="C6EFCE")),
    )
    for cell in candidate_ws["A"][1:]:
        cell.number_format = "@"

    blocked_ws = wb.create_sheet("自动阻断与仅参考")
    blocked_ws.append(BLOCKED_FIELDS)
    for row in package["auto_blocked_or_reference"]:
        blocked_ws.append([
            row["candidate_id"], row["source_document_name"], row["source_page"], row["source_section"],
            row["title"], row["evidence_type_label"], row["dedup_status_label"], row["disposition"], row["reason"],
        ])
    _style_sheet(blocked_ws, {1: 24, 2: 18, 3: 10, 4: 55, 5: 30, 6: 18, 7: 22, 8: 18, 9: 45})
    for cell in blocked_ws["A"][1:]:
        cell.number_format = "@"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def read_workbook_decisions(path: Path) -> dict[str, dict[str, Any]]:
    wb = load_workbook(path, data_only=False, read_only=False)
    sheet_name = "异常候选复核" if "异常候选复核" in wb.sheetnames else "候选审批"
    if sheet_name not in wb.sheetnames:
        raise ValueError("审批表缺少“异常候选复核”工作表")
    ws = wb[sheet_name]
    headers = {_text(cell.value): index for index, cell in enumerate(ws[1], start=1)}
    missing = [field for field in CANDIDATE_FIELDS if field not in headers]
    if missing:
        raise ValueError(f"审批表缺少字段：{'、'.join(missing)}")
    field_map = {
        "审批结论": "decision",
        "去重结论确认": "dedup_confirmation",
        "复核后资料名称": "reviewed_title",
        "复核后目标库": "reviewed_target_library_label",
        "产品边界已核验": "product_boundary_verified",
        "原始证据已核验": "original_evidence_verified",
        "原始证据路径": "original_evidence_path",
        "有效性时效已核验": "timeliness_verified",
        "敏感资料授权": "sensitive_authorized",
        "正式展示质量已确认": "formal_display_quality_verified",
        "复核人": "reviewer",
        "复核日期": "reviewed_at",
        "复核意见": "review_notes",
    }
    decisions: dict[str, dict[str, Any]] = {}
    for row_index in range(2, ws.max_row + 1):
        candidate_id = _text(ws.cell(row_index, headers["候选编号"]).value)
        if not candidate_id:
            continue
        if candidate_id in decisions:
            raise ValueError(f"审批表存在重复候选编号：{candidate_id}")
        row: dict[str, Any] = {}
        for chinese_name, field_name in field_map.items():
            value = ws.cell(row_index, headers[chinese_name]).value
            if isinstance(value, (datetime, date)):
                value = value.isoformat()
            row[field_name] = _text(value)
        decisions[candidate_id] = row
    return decisions


def _write_json(path: Path, metadata: dict[str, Any], records: list[dict[str, Any]], role: str) -> None:
    path.write_text(json.dumps({"metadata": {**metadata, "artifact_role": role, "record_count": len(records)},
                                "records": records}, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_csv(path: Path, records: list[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        for record in records:
            row = dict(record)
            for key, value in row.items():
                if isinstance(value, list):
                    row[key] = json.dumps(value, ensure_ascii=False)
            writer.writerow(row)


def write_artifacts(output_dir: Path, package: dict[str, Any]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    metadata = package["metadata"]
    artifact_specs = (
        ("asset_review_groups", package["groups"], "grouped_manual_review_overview"),
        ("asset_review_candidates", package["review_candidates"], "candidate_level_manual_review"),
        ("asset_policy_auto_accepted", package["policy_auto_accepted"], "policy_auto_accepted_knowledge_only"),
        ("asset_auto_blocked_or_reference", package["auto_blocked_or_reference"], "automatic_block_or_reference"),
        ("asset_approved_decisions", package["approved_decisions"], "validated_policy_or_human_approved_decisions"),
        ("asset_ingestion_candidates", package["ready_for_ingestion"], "ready_for_extraction_and_ingestion_validation"),
    )
    for name, records, role in artifact_specs:
        _write_json(output_dir / f"{name}.json", metadata, records, role)
        fields = sorted({key for row in records for key in row}) or ["candidate_id"]
        _write_csv(output_dir / f"{name}.csv", records, fields)
    (output_dir / "asset_review_validation_errors.json").write_text(
        json.dumps({"metadata": metadata, "records": package["validation_errors"]}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (output_dir / "p0_06_review_manifest.json").write_text(
        json.dumps({"metadata": metadata}, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--matrix", type=Path, default=DEFAULT_MATRIX)
    parser.add_argument("--review-queue", type=Path, default=DEFAULT_REVIEW_QUEUE)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--apply-decisions", type=Path, help="读取已填写的审批表并校验审批结论")
    args = parser.parse_args()

    package = build_review_package(_json(args.matrix), _json(args.review_queue))
    package["metadata"]["input_sha256"] = {
        str(args.matrix.resolve().relative_to(PROJECT_ROOT)): _sha256(args.matrix),
        str(args.review_queue.resolve().relative_to(PROJECT_ROOT)): _sha256(args.review_queue),
    }
    if args.apply_decisions:
        package = apply_review_decisions(package, read_workbook_decisions(args.apply_decisions))
    else:
        write_workbook(args.workbook, package)
    write_artifacts(args.output_dir, package)
    print(json.dumps(package["metadata"], ensure_ascii=False, indent=2))
    if package["validation_errors"]:
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
