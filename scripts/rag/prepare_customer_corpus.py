#!/usr/bin/env python3
"""Prepare customer State Grid tender files before RAG ingestion.

This script is intentionally pre-ingestion only. It reads the Jiangxi/Shanxi
customer corpus, extracts reviewable text/table artifacts, and writes a manifest
plus a quality report. It does not write pgvector or Supabase.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SEED_ROOT = PROJECT_ROOT / "rag_seed" / "power_grid_resources" / "01_tender_documents"
OUTPUT_ROOT = PROJECT_ROOT / "parsed_outputs" / "power_grid_customer_corpus"

JIANGXI_ROOT = SEED_ROOT / "20_国网江西电力2026年第一次配网省网协议库存物资类公开招标采购"
SHANXI_ROOT = SEED_ROOT / "21_国网山西电力2026年第二次物资协议库存公开招标采购"

TEXT_SUFFIXES = {".docx", ".doc"}
TABLE_SUFFIXES = {".xlsx"}
ARCHIVE_SUFFIXES = {".zip", ".rar", ".sign", ".zb"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | TABLE_SUFFIXES | ARCHIVE_SUFFIXES


@dataclass
class ParseRecord:
    source_file: str
    output_file: str | None
    province: str
    batch_no: str
    package_no: str
    package_code: str | None
    material_category: str
    doc_role: str
    qualification_mode: str | None
    source_type: str
    parser: str
    parse_status: str
    sha256: str
    file_size: int
    text_chars: int = 0
    paragraph_count: int = 0
    table_count: int = 0
    sheet_count: int = 0
    row_count: int = 0
    heading_count: int = 0
    warnings: list[str] = field(default_factory=list)
    error: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


def _now_batch_id() -> str:
    return datetime.now().strftime("customer_jx_sx_%Y%m%d_%H%M%S")


def _rel(path: Path) -> str:
    return str(path.relative_to(PROJECT_ROOT))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _safe_name(path: Path) -> str:
    name = re.sub(r"[^\w.\-\u4e00-\u9fff]+", "_", path.stem, flags=re.UNICODE).strip("_")
    digest = hashlib.sha1(str(path).encode("utf-8")).hexdigest()[:10]
    return f"{name[:80]}_{digest}"


def _soffice_bin() -> str | None:
    configured = os.getenv("SOFFICE_BIN")
    if configured:
        return configured
    return shutil.which("soffice") or (
        "/Applications/LibreOffice.app/Contents/MacOS/soffice"
        if Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists()
        else None
    )


def _detect_province(path: Path) -> tuple[str, str]:
    text = str(path)
    if "江西" in text:
        return "江西", "2026-01"
    if "山西" in text:
        return "山西", "2026-02"
    return "unknown", "unknown"


def _detect_package_code(path: Path) -> str | None:
    text = path.name
    match = re.search(r"(?:货物清单_)?([A-Z0-9]{4,8})_", text)
    if match:
        return match.group(1)
    match = re.search(r"\(([A-Z0-9-]{8,})\)", text)
    if match:
        return match.group(1)
    province, _ = _detect_province(path)
    if province == "江西":
        return "1826AA"
    if province == "山西":
        return "0526AB"
    return None


def _detect_material(path: Path) -> str:
    text = str(path)
    if "接地铁" in text:
        return "接地铁"
    if "不锈钢电缆支架" in text:
        return "不锈钢电缆支架"
    if "铁附件" in text:
        return "铁附件/铁构件"
    if "铁构件" in text:
        return "铁构件"
    return "铁构件"


def _detect_doc_role(path: Path) -> str:
    name = path.name
    if path.suffix.lower() in TABLE_SUFFIXES or "货物清单" in name:
        return "goods_list"
    if "招标文件" in name:
        return "main_tender_file"
    if "招标公告" in name:
        return "tender_notice"
    if "投标注意事项" in name:
        return "bid_instructions"
    if "技术规范" in name or "固化ID" in name:
        return "technical_spec"
    if "合同通用条款" in name:
        return "contract_general_terms"
    if "合同" in name or "合同专用条款" in name:
        return "contract_special_terms"
    if path.suffix.lower() in ARCHIVE_SUFFIXES:
        return "archive_only"
    return "unknown"


def _detect_qualification_mode(path: Path) -> str | None:
    text = str(path)
    if "资格预审" in text:
        return "prequalification"
    if "资格后审" in text:
        return "postqualification"
    return None


def _base_metadata(path: Path, batch_id: str) -> dict[str, Any]:
    province, batch_no = _detect_province(path)
    return {
        "seed_corpus": "power_grid_customer_corpus",
        "province": province,
        "batch_no": batch_no,
        "package_no": "包1",
        "package_code": _detect_package_code(path),
        "material_category": _detect_material(path),
        "doc_role": _detect_doc_role(path),
        "qualification_mode": _detect_qualification_mode(path),
        "source_file": _rel(path),
        "source_category": "01_tender_documents",
        "chunker": "parent_child_v2",
        "ingestion_batch_id": batch_id,
        "doc_version": 1,
        "status": "parsed_pending_review",
    }


def _iter_customer_files() -> list[Path]:
    roots = [JIANGXI_ROOT, SHANXI_ROOT]
    files: list[Path] = []
    for root in roots:
        if not root.exists():
            continue
        for path in root.rglob("*"):
            if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
                files.append(path)
    return sorted(files, key=lambda p: str(p))


def _extract_docx(path: Path) -> tuple[str, dict[str, int], list[str]]:
    import mammoth
    from docx import Document

    warnings: list[str] = []
    with path.open("rb") as f:
        result = mammoth.extract_raw_text(f)
    text = result.value or ""
    warnings.extend(str(message) for message in result.messages)

    doc = Document(str(path))
    paragraph_count = len(doc.paragraphs)
    table_count = len(doc.tables)
    heading_count = sum(
        1
        for para in doc.paragraphs
        if para.style is not None and str(para.style.name).lower().startswith("heading")
    )
    stats = {
        "paragraph_count": paragraph_count,
        "table_count": table_count,
        "heading_count": heading_count,
    }
    return text, stats, warnings


def _convert_doc_to_docx(path: Path, converted_dir: Path) -> Path:
    soffice = _soffice_bin()
    if not soffice:
        raise RuntimeError("soffice executable not found; .doc conversion requires LibreOffice")
    converted_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="power-grid-doc-") as profile_dir:
        cmd = [
            soffice,
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            f"-env:UserInstallation=file://{profile_dir}",
            "--convert-to",
            "docx",
            "--outdir",
            str(converted_dir),
            str(path),
        ]
        proc = subprocess.run(cmd, text=True, capture_output=True, timeout=180, check=False)
    converted = converted_dir / f"{path.stem}.docx"
    if proc.returncode != 0 or not converted.exists():
        raise RuntimeError(
            "LibreOffice .doc conversion failed: "
            f"returncode={proc.returncode} stdout={proc.stdout.strip()} stderr={proc.stderr.strip()}"
        )
    return converted


def _write_text_artifact(path: Path, text: str, output_dir: Path, suffix: str = ".md") -> Path:
    target = output_dir / "texts" / f"{_safe_name(path)}{suffix}"
    target.parent.mkdir(parents=True, exist_ok=True)
    title = path.stem
    target.write_text(f"# {title}\n\n{text.strip()}\n", encoding="utf-8")
    return target


def _parse_text_file(path: Path, metadata: dict[str, Any], output_dir: Path) -> ParseRecord:
    parser = "mammoth_docx"
    warnings: list[str] = []
    parse_path = path
    try:
        if path.suffix.lower() == ".doc":
            parser = "libreoffice_doc_to_docx+mammoth"
            parse_path = _convert_doc_to_docx(path, output_dir / "converted_docx")
            warnings.append(f"converted_docx={_rel(parse_path)}")
        text, stats, extractor_warnings = _extract_docx(parse_path)
        warnings.extend(extractor_warnings)
        text_chars = len(text.strip())
        status = "parsed" if text_chars >= 100 else "needs_review"
        if text_chars < 100:
            warnings.append("extracted text is too short")
        out = _write_text_artifact(path, text, output_dir)
        return ParseRecord(
            source_file=_rel(path),
            output_file=_rel(out),
            province=metadata["province"],
            batch_no=metadata["batch_no"],
            package_no=metadata["package_no"],
            package_code=metadata.get("package_code"),
            material_category=metadata["material_category"],
            doc_role=metadata["doc_role"],
            qualification_mode=metadata.get("qualification_mode"),
            source_type=path.suffix.lower().lstrip("."),
            parser=parser,
            parse_status=status,
            sha256=_sha256(path),
            file_size=path.stat().st_size,
            text_chars=text_chars,
            paragraph_count=stats["paragraph_count"],
            table_count=stats["table_count"],
            heading_count=stats["heading_count"],
            warnings=warnings,
            metadata=metadata,
        )
    except Exception as exc:
        return _failed_record(path, metadata, parser, "needs_review", exc, warnings)


def _parse_xlsx(path: Path, metadata: dict[str, Any], output_dir: Path) -> ParseRecord:
    parser = "openpyxl"
    try:
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("openpyxl is not installed; cannot parse .xlsx structure") from exc

        workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
        sheets: list[dict[str, Any]] = []
        row_count = 0
        for sheet in workbook.worksheets:
            rows: list[list[Any]] = []
            for row in sheet.iter_rows(values_only=True):
                values = [value for value in row]
                if any(value not in (None, "") for value in values):
                    rows.append(values)
            row_count += len(rows)
            merged_ranges = [str(rng) for rng in sheet.merged_cells.ranges]
            sheets.append({
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged_ranges": merged_ranges,
                "rows": rows,
                "retrieval_summary": _table_summary(path, sheet.title, rows, metadata),
            })
        payload = {
            "source_file": _rel(path),
            "metadata": metadata,
            "sheets": sheets,
        }
        out = output_dir / "tables" / f"{_safe_name(path)}.json"
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        status = "parsed" if row_count > 0 else "needs_review"
        warnings = [] if row_count > 0 else ["xlsx has no non-empty rows"]
        return ParseRecord(
            source_file=_rel(path),
            output_file=_rel(out),
            province=metadata["province"],
            batch_no=metadata["batch_no"],
            package_no=metadata["package_no"],
            package_code=metadata.get("package_code"),
            material_category=metadata["material_category"],
            doc_role=metadata["doc_role"],
            qualification_mode=metadata.get("qualification_mode"),
            source_type=path.suffix.lower().lstrip("."),
            parser=parser,
            parse_status=status,
            sha256=_sha256(path),
            file_size=path.stat().st_size,
            sheet_count=len(sheets),
            row_count=row_count,
            warnings=warnings,
            metadata=metadata,
        )
    except Exception as exc:
        return _failed_record(path, metadata, parser, "needs_review", exc)


def _table_summary(path: Path, sheet_name: str, rows: list[list[Any]], metadata: dict[str, Any]) -> str:
    preview_cells: list[str] = []
    for row in rows[:5]:
        for cell in row[:8]:
            if cell not in (None, ""):
                preview_cells.append(str(cell))
    preview = "、".join(preview_cells[:12])
    return (
        f"{metadata['province']}{metadata['batch_no']} {metadata['material_category']} "
        f"{metadata.get('package_code') or ''} 货物清单表 {path.name} / {sheet_name}；"
        f"表内关键词：{preview}"
    )


def _archive_record(path: Path, metadata: dict[str, Any]) -> ParseRecord:
    return ParseRecord(
        source_file=_rel(path),
        output_file=None,
        province=metadata["province"],
        batch_no=metadata["batch_no"],
        package_no=metadata["package_no"],
        package_code=metadata.get("package_code"),
        material_category=metadata["material_category"],
        doc_role="archive_only",
        qualification_mode=metadata.get("qualification_mode"),
        source_type=path.suffix.lower().lstrip("."),
        parser="archive_only",
        parse_status="archived",
        sha256=_sha256(path),
        file_size=path.stat().st_size,
        warnings=["archive/signature/platform file is preserved but not parsed for RAG"],
        metadata={**metadata, "doc_role": "archive_only", "status": "archived"},
    )


def _failed_record(
    path: Path,
    metadata: dict[str, Any],
    parser: str,
    status: str,
    exc: Exception,
    warnings: list[str] | None = None,
) -> ParseRecord:
    return ParseRecord(
        source_file=_rel(path),
        output_file=None,
        province=metadata["province"],
        batch_no=metadata["batch_no"],
        package_no=metadata["package_no"],
        package_code=metadata.get("package_code"),
        material_category=metadata["material_category"],
        doc_role=metadata["doc_role"],
        qualification_mode=metadata.get("qualification_mode"),
        source_type=path.suffix.lower().lstrip("."),
        parser=parser,
        parse_status=status,
        sha256=_sha256(path),
        file_size=path.stat().st_size,
        warnings=warnings or [],
        error=str(exc),
        metadata={**metadata, "status": status},
    )


def parse_one(path: Path, batch_id: str, output_dir: Path) -> ParseRecord:
    metadata = _base_metadata(path, batch_id)
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return _parse_text_file(path, metadata, output_dir)
    if suffix in TABLE_SUFFIXES:
        return _parse_xlsx(path, metadata, output_dir)
    return _archive_record(path, metadata)


def _write_manifest(records: list[ParseRecord], output_dir: Path, batch_id: str) -> Path:
    status_counts: dict[str, int] = {}
    role_counts: dict[str, int] = {}
    for rec in records:
        status_counts[rec.parse_status] = status_counts.get(rec.parse_status, 0) + 1
        role_counts[rec.doc_role] = role_counts.get(rec.doc_role, 0) + 1
    payload = {
        "batch_id": batch_id,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source_roots": [_rel(JIANGXI_ROOT), _rel(SHANXI_ROOT)],
        "summary": {
            "total_files": len(records),
            "status_counts": status_counts,
            "doc_role_counts": role_counts,
            "text_chars": sum(r.text_chars for r in records),
            "table_rows": sum(r.row_count for r in records),
        },
        "records": [asdict(rec) for rec in records],
    }
    out = output_dir / "manifest.json"
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


def _write_report(records: list[ParseRecord], output_dir: Path, batch_id: str) -> Path:
    def count(status: str) -> int:
        return sum(1 for rec in records if rec.parse_status == status)

    lines = [
        "# 江西/山西客户资料解析质量报告",
        "",
        f"> 批次：`{batch_id}`",
        f"> 生成时间：{datetime.now().isoformat(timespec='seconds')}",
        "",
        "## 总览",
        "",
        "| 指标 | 数量 |",
        "| --- | ---: |",
        f"| 文件总数 | {len(records)} |",
        f"| 已解析 | {count('parsed')} |",
        f"| 仅归档 | {count('archived')} |",
        f"| 需复核 | {count('needs_review')} |",
        f"| 文本字符数 | {sum(r.text_chars for r in records)} |",
        f"| 表格行数 | {sum(r.row_count for r in records)} |",
        "",
        "## 解析结果",
        "",
        "| 状态 | 省份 | 角色 | 类型 | 文件 | 解析器 | 输出 | 备注 |",
        "| --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for rec in records:
        note_parts = []
        if rec.error:
            note_parts.append(rec.error)
        note_parts.extend(rec.warnings[:2])
        note = "；".join(note_parts).replace("\n", " ")[:160]
        output = rec.output_file or "-"
        lines.append(
            f"| {rec.parse_status} | {rec.province} | `{rec.doc_role}` | {rec.source_type} | "
            f"`{Path(rec.source_file).name}` | {rec.parser} | `{output}` | {note} |"
        )
    lines.extend([
        "",
        "## 入库前结论",
        "",
        "- `parsed` 文件可以进入下一步清洗、父子分块 dry-run。",
        "- `archived` 文件只保留原始归档，不进入 RAG。",
        "- `needs_review` 文件不得静默入库，必须先补依赖、换解析器或人工复核。",
        "- `.xlsx` 必须以 manifest 中的表格 JSON 为基础生成结构化表和检索摘要，不能只按普通长文本入库。",
    ])
    out = output_dir / "parse_quality_report.md"
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--batch-id", default=_now_batch_id())
    parser.add_argument("--output-root", default=str(OUTPUT_ROOT))
    parser.add_argument("--include-archives", action="store_true", help="include zip/rar/sign/zb records in manifest")
    args = parser.parse_args()

    batch_id = args.batch_id
    output_dir = Path(args.output_root) / batch_id
    output_dir.mkdir(parents=True, exist_ok=True)

    files = _iter_customer_files()
    if not args.include_archives:
        files = [path for path in files if path.suffix.lower() not in ARCHIVE_SUFFIXES]

    records: list[ParseRecord] = []
    for index, path in enumerate(files, 1):
        print(f"[{index}/{len(files)}] {path.name}")
        rec = parse_one(path, batch_id, output_dir)
        print(f"  -> {rec.parse_status} role={rec.doc_role} parser={rec.parser}")
        records.append(rec)

    manifest = _write_manifest(records, output_dir, batch_id)
    report = _write_report(records, output_dir, batch_id)
    print(f"\nManifest: {_rel(manifest)}")
    print(f"Report: {_rel(report)}")
    print(f"Done. files={len(records)} parsed={sum(1 for r in records if r.parse_status == 'parsed')} needs_review={sum(1 for r in records if r.parse_status == 'needs_review')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
