#!/usr/bin/env python3
"""Build staging artifacts for the Liaoning/Taichang MVP RAG batch.

This script is intentionally pre-ingestion. It prepares a normal customer
corpus manifest that the existing chunk dry-run and ingest scripts can consume,
and a separate knowledge_assets payload file for Taichang image assets.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))


BATCH_ID = "customer_liaoning_taichang_20260606_p0"
PARSED_ROOT = PROJECT_ROOT / "parsed_outputs" / "power_grid_customer_corpus" / BATCH_ID
SOURCE_MANIFEST = PARSED_ROOT / "manifest.json"
GOODS_ROWS = PARSED_ROOT / "goods_tables" / "goods_rows.json"
PARSE_QUALITY = PARSED_ROOT / "parse_quality_report.json"
ASSET_INDEX = PARSED_ROOT / "asset_index.json"
REFERENCE_TEMPLATES = PARSED_ROOT / "reference_templates" / "haoqian_reference_templates.json"

ARCHIVE_ROLES = {"archive_only", "unknown"}
TENDER_TEXT_ROLES = {
    "technical_spec",
    "contract_general_terms",
    "contract_special_terms",
    "tender_notice",
    "main_tender_file",
}
TEXT_SUFFIXES = {".doc", ".docx"}
DISPLAY_CONTEXTS = ["bid_writing", "knowledge_chat", "asset_search"]


@dataclass
class StagingRecord:
    source_file: str
    output_file: str | None
    province: str
    batch_no: str
    package_no: str
    package_code: str | None
    material_category: str
    doc_role: str
    qualification_mode: str | None
    source_type: str
    parser: str
    parse_status: str
    sha256: str
    file_size: int
    text_chars: int = 0
    paragraph_count: int = 0
    table_count: int = 0
    sheet_count: int = 0
    row_count: int = 0
    heading_count: int = 0
    warnings: list[str] = field(default_factory=list)
    error: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


def _rel(path: Path) -> str:
    return str(path.relative_to(PROJECT_ROOT))


def _load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _safe_name(value: str) -> str:
    clean = re.sub(r"[^\w.\-\u4e00-\u9fff]+", "_", value, flags=re.UNICODE).strip("_")
    digest = hashlib.sha1(value.encode("utf-8")).hexdigest()[:8]
    return f"{clean[:90]}_{digest}"


def _soffice_bin() -> str | None:
    configured = os.getenv("SOFFICE_BIN")
    if configured:
        return configured
    return shutil.which("soffice") or (
        "/Applications/LibreOffice.app/Contents/MacOS/soffice"
        if Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists()
        else None
    )


def _extract_docx(path: Path) -> tuple[str, dict[str, int], list[str]]:
    import mammoth
    from docx import Document

    warnings: list[str] = []
    with path.open("rb") as f:
        result = mammoth.extract_raw_text(f)
    warnings.extend(str(message) for message in result.messages)

    doc = Document(str(path))
    stats = {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "heading_count": sum(
            1
            for para in doc.paragraphs
            if para.style is not None and str(para.style.name).lower().startswith("heading")
        ),
    }
    return result.value or "", stats, warnings


def _convert_doc_to_docx(path: Path, converted_dir: Path) -> Path:
    soffice = _soffice_bin()
    if not soffice:
        raise RuntimeError("soffice executable not found; .doc conversion requires LibreOffice")
    converted_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="taichang-doc-profile-") as profile_dir:
        proc = subprocess.run(
            [
                soffice,
                "--headless",
                "--nologo",
                "--nofirststartwizard",
                f"-env:UserInstallation=file://{profile_dir}",
                "--convert-to",
                "docx",
                "--outdir",
                str(converted_dir),
                str(path),
            ],
            text=True,
            capture_output=True,
            timeout=180,
            check=False,
        )
    converted = converted_dir / f"{path.stem}.docx"
    if proc.returncode != 0 or not converted.exists():
        raise RuntimeError(
            "LibreOffice .doc conversion failed: "
            f"returncode={proc.returncode} stdout={proc.stdout.strip()} stderr={proc.stderr.strip()}"
        )
    return converted


def _write_text(path: Path, title: str, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(f"# {title}\n\n{text.strip()}\n", encoding="utf-8")


def _common_metadata(record: dict[str, Any], *, doc_role: str | None = None) -> dict[str, Any]:
    return {
        "seed_corpus": "power_grid_customer_corpus",
        "source_domain": record.get("source_domain"),
        "province": record.get("province") or "辽宁",
        "batch_no": record.get("batch_no") or "2025-03",
        "package_no": record.get("package_no") or "",
        "package_code": record.get("package_code") or "2225AC",
        "material_category": record.get("material_category") or "",
        "doc_role": doc_role or record.get("doc_role"),
        "qualification_mode": record.get("qualification_mode"),
        "source_file": record.get("source_file"),
        "source_category": (
            "05_enterprise_documents"
            if record.get("source_domain") == "enterprise_fact"
            else "01_tender_documents"
        ),
        "chunker": "parent_child_v2",
        "ingestion_batch_id": BATCH_ID,
        "doc_version": 1,
        "status": "parsed_pending_ingestion",
        "doc_owner": record.get("doc_owner"),
        "enterprise": record.get("enterprise"),
        "reference_only": bool(record.get("reference_only")),
        "fact_source_allowed_for_enterprise": bool(record.get("fact_source_allowed_for_enterprise")),
        "privacy_level": record.get("privacy_level"),
        "target_library": record.get("target_library"),
        "do_not_mix_with": ["河北豪乾参考稿"] if record.get("source_domain") == "enterprise_fact" else [],
    }


def _stage_tender_text(record: dict[str, Any], out_dir: Path) -> StagingRecord:
    source = PROJECT_ROOT / record["source_file"]
    parser = "mammoth_docx"
    warnings: list[str] = []
    parse_path = source
    try:
        if source.suffix.lower() == ".doc":
            parser = "libreoffice_doc_to_docx+mammoth"
            parse_path = _convert_doc_to_docx(source, out_dir / "converted_docx")
            warnings.append(f"converted_docx={_rel(parse_path)}")
        text, stats, extractor_warnings = _extract_docx(parse_path)
        warnings.extend(extractor_warnings)
        text_chars = len(text.strip())
        status = "parsed" if text_chars >= 80 else "needs_review"
        if text_chars < 80:
            warnings.append("extracted text is too short")
        output = out_dir / "texts" / "liaoning_tender" / f"{_safe_name(record['source_file'])}.md"
        _write_text(output, source.stem, text)
        metadata = _common_metadata(record)
        return StagingRecord(
            source_file=record["source_file"],
            output_file=_rel(output),
            province=record.get("province") or "辽宁",
            batch_no=record.get("batch_no") or "2025-03",
            package_no=record.get("package_no") or "",
            package_code=record.get("package_code"),
            material_category=record.get("material_category") or "",
            doc_role=record["doc_role"],
            qualification_mode=None,
            source_type=source.suffix.lower().lstrip("."),
            parser=parser,
            parse_status=status,
            sha256=record.get("sha256") or _sha256_file(source),
            file_size=int(record.get("file_size") or source.stat().st_size),
            text_chars=text_chars,
            paragraph_count=stats["paragraph_count"],
            table_count=stats["table_count"],
            heading_count=stats["heading_count"],
            warnings=warnings,
            metadata=metadata,
        )
    except Exception as exc:
        return StagingRecord(
            source_file=record["source_file"],
            output_file=None,
            province=record.get("province") or "辽宁",
            batch_no=record.get("batch_no") or "2025-03",
            package_no=record.get("package_no") or "",
            package_code=record.get("package_code"),
            material_category=record.get("material_category") or "",
            doc_role=record["doc_role"],
            qualification_mode=None,
            source_type=source.suffix.lower().lstrip("."),
            parser=parser,
            parse_status="needs_review",
            sha256=record.get("sha256") or _sha256_file(source),
            file_size=int(record.get("file_size") or source.stat().st_size),
            warnings=warnings,
            error=str(exc),
            metadata={**_common_metadata(record), "status": "needs_review"},
        )


def _stage_goods_records(out_dir: Path) -> list[StagingRecord]:
    goods_rows = _load_json(GOODS_ROWS)
    staged: list[StagingRecord] = []
    header_keys = [
        "分标编号",
        "包名称",
        "分包编号",
        "项目单位",
        "需求单位",
        "项目名称",
        "工程电压等级",
        "物资名称",
        "物资描述",
        "单位",
        "数量",
        "首批交货日期",
        "最后一批交货日期",
        "交货地点",
        "交货方式",
        "技术规范编码",
        "物料编码",
        "扩展描述",
        "扩展编码",
    ]
    for family, material in [("CPVC", "电缆保护管CPVC"), ("MPP", "电缆保护管MPP")]:
        rows = [row for row in goods_rows if row.get("_material_family") == family]
        if not rows:
            continue
        table_rows = [header_keys]
        for row in rows:
            table_rows.append([row.get(key, "") for key in header_keys])
        summary = (
            f"辽宁 2025-03 2225AC {material} 货物清单，"
            f"共 {len(rows)} 条需求；规格："
            + "、".join(sorted({str(row.get('_spec') or '') for row in rows if row.get('_spec')}))
            + "；技术规范编码："
            + "、".join(sorted({str(row.get('技术规范编码') or '') for row in rows if row.get('技术规范编码')}))
        )
        payload = {
            "source_file": f"{BATCH_ID}/goods_tables/goods_rows.json#{family}",
            "metadata": {
                "seed_corpus": "power_grid_customer_corpus",
                "source_domain": "tender_requirement",
                "province": "辽宁",
                "batch_no": "2025-03",
                "package_no": "包1-包4",
                "package_code": "2225AC",
                "material_category": material,
                "doc_role": "goods_list",
                "source_category": "01_tender_documents",
                "chunker": "table_three_forms",
                "ingestion_batch_id": BATCH_ID,
                "doc_version": 1,
                "status": "parsed_pending_ingestion",
                "doc_owner": "国网辽宁省电力有限公司",
                "reference_only": False,
                "fact_source_allowed_for_enterprise": False,
                "privacy_level": "internal",
                "target_library": "knowledge_library",
            },
            "sheets": [
                {
                    "name": f"{family}_goods_rows",
                    "max_row": len(table_rows),
                    "max_column": len(header_keys),
                    "merged_ranges": [],
                    "rows": table_rows,
                    "retrieval_summary": summary,
                }
            ],
        }
        output = out_dir / "tables" / f"liaoning_2225AC_{family}_goods_rows.json"
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        staged.append(
            StagingRecord(
                source_file=payload["source_file"],
                output_file=_rel(output),
                province="辽宁",
                batch_no="2025-03",
                package_no="包1-包4",
                package_code="2225AC",
                material_category=material,
                doc_role="goods_list",
                qualification_mode=None,
                source_type="json",
                parser="goods_rows_structured_staging",
                parse_status="parsed",
                sha256=_sha256_text(json.dumps(rows, ensure_ascii=False)),
                file_size=output.stat().st_size,
                sheet_count=1,
                row_count=len(rows),
                metadata=payload["metadata"],
            )
        )
    return staged


def _stage_enterprise_records(out_dir: Path) -> list[StagingRecord]:
    report = _load_json(PARSE_QUALITY)
    records: list[StagingRecord] = []
    for item in report.get("records", []):
        markdown = item.get("markdown_path")
        if not markdown:
            continue
        md_path = PROJECT_ROOT / markdown
        if not md_path.exists():
            continue
        text = md_path.read_text(encoding="utf-8")
        source_file = item["source_file"]
        metadata = {
            "seed_corpus": "power_grid_customer_corpus",
            "source_domain": "enterprise_fact",
            "province": "河北",
            "batch_no": "taichang_mvp_20260606",
            "package_no": "",
            "package_code": "2225AC",
            "material_category": "电缆保护管CPVC/MPP",
            "doc_role": "enterprise_evidence",
            "source_file": source_file,
            "source_category": "05_enterprise_documents",
            "chunker": "parent_child_v2",
            "ingestion_batch_id": BATCH_ID,
            "doc_version": 1,
            "status": "parsed_pending_ingestion",
            "source_type": "mineru_pdf",
            "parser": "mineru",
            "doc_owner": "泰昌",
            "enterprise": "泰昌",
            "evidence_type": item.get("evidence_type"),
            "target_library": item.get("target_library"),
            "privacy_level": item.get("privacy_level"),
            "reference_only": False,
            "fact_source_allowed_for_enterprise": True,
            "access_scope": item.get("access_scope") or "taichang_tenant_internal",
            "requires_authorization": bool(item.get("requires_authorization")),
            "display_contexts": item.get("display_contexts") or DISPLAY_CONTEXTS,
            "tenant_visibility": "taichang_only",
            "sensitive_handling": "taichang_internal_use_no_redaction_required",
            "do_not_mix_with": ["河北豪乾参考稿"],
            "quality_flags": item.get("quality_flags") or [],
        }
        records.append(
            StagingRecord(
                source_file=source_file,
                output_file=markdown,
                province="河北",
                batch_no="taichang_mvp_20260606",
                package_no="",
                package_code="2225AC",
                material_category="电缆保护管CPVC/MPP",
                doc_role="enterprise_evidence",
                qualification_mode=None,
                source_type="pdf",
                parser="mineru",
                parse_status="parsed",
                sha256=_sha256_text(text),
                file_size=md_path.stat().st_size,
                text_chars=len(text.strip()),
                warnings=item.get("quality_flags") or [],
                metadata=metadata,
            )
        )
    return records


def _stage_reference_records(out_dir: Path) -> list[StagingRecord]:
    payload = _load_json(REFERENCE_TEMPLATES)
    records: list[StagingRecord] = []
    for item in payload.get("records", []):
        source_file = item["source_file"]
        title = Path(source_file).stem
        toc_lines = item.get("toc_lines") or []
        cover = json.dumps(item.get("cover_fields") or {}, ensure_ascii=False, indent=2)
        allowed = "、".join(item.get("template_extract_allowed_fields") or [])
        forbidden = "、".join(item.get("template_extract_forbidden_fields") or [])
        body = (
            f"# {title}\n\n"
            f"doc_role: `{item.get('doc_role')}`\n\n"
            f"doc_owner: {item.get('doc_owner')}\n\n"
            f"reference_only: true\n\n"
            f"允许抽取字段：{allowed}\n\n"
            f"禁止作为泰昌事实字段：{forbidden}\n\n"
            "## 封面字段\n\n"
            f"```json\n{cover}\n```\n\n"
            "## 目录/章节结构\n\n"
            + "\n".join(f"- {line}" for line in toc_lines[:120])
            + "\n"
        )
        output = out_dir / "texts" / "reference_templates" / f"{_safe_name(source_file)}.md"
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(body, encoding="utf-8")
        metadata = {
            "seed_corpus": "power_grid_customer_corpus",
            "source_domain": "reference_template",
            "province": "辽宁",
            "batch_no": "2025-03",
            "package_no": "",
            "package_code": item.get("package_code") or "2225AC",
            "material_category": item.get("material_category") or "",
            "doc_role": item.get("doc_role"),
            "source_file": source_file,
            "source_category": "01_tender_documents",
            "chunker": "parent_child_v2",
            "ingestion_batch_id": BATCH_ID,
            "doc_version": 1,
            "status": "parsed_pending_ingestion",
            "doc_owner": item.get("doc_owner"),
            "enterprise": None,
            "reference_only": True,
            "fact_source_allowed_for_enterprise": False,
            "privacy_level": "reference_only",
            "target_library": "reference_template_library",
            "template_extract_allowed_fields": item.get("template_extract_allowed_fields") or [],
            "template_extract_forbidden_fields": item.get("template_extract_forbidden_fields") or [],
        }
        records.append(
            StagingRecord(
                source_file=source_file,
                output_file=_rel(output),
                province="辽宁",
                batch_no="2025-03",
                package_no="",
                package_code=item.get("package_code") or "2225AC",
                material_category=item.get("material_category") or "",
                doc_role=item.get("doc_role"),
                qualification_mode=None,
                source_type="pdf_reference_extract",
                parser="reference_template_extractor",
                parse_status="parsed",
                sha256=item.get("sha256") or _sha256_text(body),
                file_size=output.stat().st_size,
                text_chars=len(body),
                metadata=metadata,
            )
        )
    return records


def _asset_payload(asset: dict[str, Any]) -> dict[str, Any]:
    asset_path = PROJECT_ROOT / asset["asset_path"]
    evidence_type = asset.get("evidence_type") or "enterprise_evidence"
    target_library = asset.get("target_library") or "product_library"
    category = "泰昌资信图片" if target_library == "qualification_library" else "泰昌产品能力图片"
    asset_type = "qualification_image" if target_library == "qualification_library" else "product_image"
    tags = ["泰昌", "MVP试点企业", evidence_type]
    tags.extend(asset.get("metadata", {}).get("tags") or [])
    tags = list(dict.fromkeys(str(tag) for tag in tags if tag))
    searchable = asset.get("retrieval_text") or asset.get("ocr_context") or asset.get("caption_candidate") or ""
    title = f"泰昌 {evidence_type} 第{asset.get('page_no')}页图片 {asset.get('asset_id')}"
    metadata = {
        **(asset.get("metadata") or {}),
        "source_batch_id": BATCH_ID,
        "source_domain": "enterprise_fact",
        "doc_owner": "泰昌",
        "enterprise": "泰昌",
        "evidence_type": evidence_type,
        "target_library": target_library,
        "asset_id": asset.get("asset_id"),
        "asset_path": asset.get("asset_path"),
        "asset_sha256": asset.get("asset_sha256"),
        "source_file": asset.get("source_file"),
        "page_no": asset.get("page_no"),
        "bbox": asset.get("bbox"),
        "privacy_level": asset.get("privacy_level"),
        "access_scope": asset.get("access_scope") or "taichang_tenant_internal",
        "tenant_visibility": "taichang_only",
        "reference_only": False,
        "fact_source_allowed_for_enterprise": True,
        "do_not_mix_with": ["河北豪乾参考稿"],
    }
    return {
        "title": title,
        "description": asset.get("caption_candidate") or asset.get("ocr_context") or "",
        "category": category,
        "asset_type": asset_type,
        "file_name": asset_path.name,
        "file_ext": asset_path.suffix.lower(),
        "mime_type": "image/jpeg" if asset_path.suffix.lower() in {".jpg", ".jpeg"} else "image/png",
        "file_size": asset_path.stat().st_size if asset_path.exists() else None,
        "local_path": asset.get("asset_path"),
        "source_type": "taichang_mvp_mineru_asset",
        "license": "customer_internal",
        "attribution": "河北泰昌电力器材科技有限公司",
        "is_synthetic": False,
        "is_sensitive": asset.get("privacy_level") == "taichang_internal_private",
        "anonymized": False,
        "industry": "电力行业",
        "applicable_sections": [],
        "applicable_volumes": asset.get("metadata", {}).get("applicable_volumes") or [],
        "tags": tags,
        "specs": {
            "page_no": asset.get("page_no"),
            "bbox": asset.get("bbox"),
            "evidence_type": evidence_type,
            "target_library": target_library,
        },
        "ocr_text": asset.get("ocr_context") or "",
        "ai_caption": asset.get("caption_candidate") or "",
        "searchable_text": searchable,
        "status": "indexed",
        "metadata": metadata,
    }


def _write_asset_payloads(out_dir: Path) -> Path:
    index = _load_json(ASSET_INDEX)
    payloads = [_asset_payload(asset) for asset in index.get("assets", [])]
    out = out_dir / "asset_staging_payloads.json"
    out.write_text(json.dumps({"batch_id": BATCH_ID, "assets": payloads}, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


def _stage_asset_catalog_records(out_dir: Path) -> list[StagingRecord]:
    """Create searchable text catalogs for image metadata.

    `eval_recall.py` evaluates document_chunks, not knowledge_assets. These
    records make image metadata testable before the formal asset table import.
    """
    index = _load_json(ASSET_INDEX)
    groups: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for asset in index.get("assets", []):
        key = (
            asset.get("evidence_type") or "enterprise_evidence",
            asset.get("target_library") or "product_library",
            asset.get("privacy_level") or "private",
        )
        groups.setdefault(key, []).append(asset)

    records: list[StagingRecord] = []
    for (evidence_type, target_library, privacy_level), assets in sorted(groups.items()):
        lines = [
            f"# 泰昌图片资产目录 - {evidence_type} - {privacy_level}",
            "",
            "本目录用于泰昌 MVP 试点企业内部标书写作、智能问答和图片资产检索。",
            "所有资产均为泰昌企业事实，仅在泰昌租户内使用。",
            "",
        ]
        for asset in assets:
            contexts = asset.get("display_contexts") or asset.get("metadata", {}).get("display_contexts") or []
            lines.extend([
                f"## {asset.get('asset_id')}",
                "",
                f"- asset_path: {asset.get('asset_path')}",
                f"- source_file: {asset.get('source_file')}",
                f"- page_no: {asset.get('page_no')}",
                f"- bbox: {json.dumps(asset.get('bbox'), ensure_ascii=False)}",
                f"- display_contexts: {', '.join(contexts)}",
                f"- requires_authorization=false",
                f"- access_scope: {asset.get('access_scope') or asset.get('metadata', {}).get('access_scope') or 'taichang_tenant_internal'}",
                f"- tenant_visibility: {asset.get('metadata', {}).get('tenant_visibility') or 'taichang_only'}",
                f"- target_library: {target_library}",
                f"- evidence_type: {evidence_type}",
                f"- caption: {asset.get('caption_candidate') or ''}",
                f"- ocr_context: {asset.get('ocr_context') or ''}",
                "",
            ])
        body = "\n".join(lines)
        output = out_dir / "texts" / "asset_catalogs" / f"taichang_{evidence_type}_{privacy_level}.md"
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(body, encoding="utf-8")
        metadata = {
            "seed_corpus": "power_grid_customer_corpus",
            "source_domain": "enterprise_fact",
            "province": "河北",
            "batch_no": "taichang_mvp_20260606",
            "package_no": "",
            "package_code": "2225AC",
            "material_category": "电缆保护管CPVC/MPP",
            "doc_role": "enterprise_evidence",
            "source_file": _rel(output),
            "source_category": "05_enterprise_documents",
            "chunker": "parent_child_v2",
            "ingestion_batch_id": BATCH_ID,
            "doc_version": 1,
            "status": "parsed_pending_ingestion",
            "source_type": "asset_metadata_catalog",
            "parser": "asset_index_catalog",
            "doc_owner": "泰昌",
            "enterprise": "泰昌",
            "evidence_type": evidence_type,
            "target_library": target_library,
            "privacy_level": privacy_level,
            "reference_only": False,
            "fact_source_allowed_for_enterprise": True,
            "access_scope": "taichang_tenant_internal",
            "requires_authorization": False,
            "display_contexts": DISPLAY_CONTEXTS,
            "tenant_visibility": "taichang_only",
            "sensitive_handling": "taichang_internal_use_no_redaction_required",
            "do_not_mix_with": ["河北豪乾参考稿"],
            "is_asset_catalog": True,
        }
        records.append(
            StagingRecord(
                source_file=_rel(output),
                output_file=_rel(output),
                province="河北",
                batch_no="taichang_mvp_20260606",
                package_no="",
                package_code="2225AC",
                material_category="电缆保护管CPVC/MPP",
                doc_role="enterprise_evidence",
                qualification_mode=None,
                source_type="asset_metadata_catalog",
                parser="asset_index_catalog",
                parse_status="parsed",
                sha256=_sha256_text(body),
                file_size=output.stat().st_size,
                text_chars=len(body),
                warnings=[],
                metadata=metadata,
            )
        )
    return records


def _write_manifest(records: list[StagingRecord], out_dir: Path) -> Path:
    status_counts: dict[str, int] = {}
    role_counts: dict[str, int] = {}
    domain_counts: dict[str, int] = {}
    for rec in records:
        status_counts[rec.parse_status] = status_counts.get(rec.parse_status, 0) + 1
        role_counts[rec.doc_role] = role_counts.get(rec.doc_role, 0) + 1
        domain = str(rec.metadata.get("source_domain") or "unknown")
        domain_counts[domain] = domain_counts.get(domain, 0) + 1
    payload = {
        "batch_id": BATCH_ID,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_manifest": _rel(SOURCE_MANIFEST),
        "summary": {
            "total_records": len(records),
            "status_counts": status_counts,
            "doc_role_counts": role_counts,
            "source_domain_counts": domain_counts,
            "text_chars": sum(rec.text_chars for rec in records),
            "table_rows": sum(rec.row_count for rec in records),
        },
        "records": [asdict(rec) for rec in records],
    }
    out = out_dir / "staging_manifest.json"
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


def _write_report(records: list[StagingRecord], assets_path: Path, out_dir: Path) -> Path:
    parsed = [rec for rec in records if rec.parse_status == "parsed"]
    needs_review = [rec for rec in records if rec.parse_status == "needs_review"]
    assets = _load_json(assets_path).get("assets", [])
    internal_assets = [
        asset for asset in assets
        if (asset.get("metadata") or {}).get("privacy_level") == "taichang_internal_private"
    ]
    lines = [
        "# 辽宁 / 泰昌 MVP 入库 Staging 报告",
        "",
        f"> 批次：`{BATCH_ID}`",
        f"> 生成时间：{datetime.now(timezone.utc).isoformat()}",
        "",
        "## 总览",
        "",
        "| 指标 | 数量 |",
        "| --- | ---: |",
        f"| staging records | {len(records)} |",
        f"| parsed | {len(parsed)} |",
        f"| needs_review | {len(needs_review)} |",
        f"| 文本字符数 | {sum(rec.text_chars for rec in records)} |",
        f"| 货物清单行 | {sum(rec.row_count for rec in records)} |",
        f"| 图片 asset payload | {len(assets)} |",
        f"| 泰昌内部私有图片 | {len(internal_assets)} |",
        "",
        "## 边界",
        "",
        "- 辽宁招标要求：`source_domain=tender_requirement`，不得混入泰昌企业事实。",
        "- 泰昌企业事实：`source_domain=enterprise_fact`，`enterprise=泰昌`，不得混入河北豪乾参考稿。",
        "- 河北豪乾参考稿：`source_domain=reference_template`，`reference_only=true`，只用于格式/目录/章节结构参考。",
        "- 泰昌内部图片：泰昌租户内可用于 `knowledge_chat` / `bid_writing` / `asset_search`，禁止跨企业/跨租户使用。",
        "",
        "## 记录明细",
        "",
        "| 状态 | domain | role | material | 文件 | 输出 | 备注 |",
        "| --- | --- | --- | --- | --- | --- | --- |",
    ]
    for rec in records:
        note = (rec.error or "；".join(rec.warnings[:2]) or "").replace("\n", " ")[:140]
        lines.append(
            f"| {rec.parse_status} | `{rec.metadata.get('source_domain')}` | `{rec.doc_role}` | "
            f"`{rec.material_category}` | `{Path(rec.source_file).name}` | `{rec.output_file or '-'}` | {note} |"
        )
    if needs_review:
        lines.extend([
            "",
            "## 需处理后再正式入库",
            "",
        ])
        for rec in needs_review:
            lines.append(f"- `{rec.source_file}`：{rec.error or '解析文本过短'}")
    out = out_dir / "staging_report.md"
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out


def build_staging(out_dir: Path) -> tuple[Path, Path, Path]:
    source = _load_json(SOURCE_MANIFEST)
    records: list[StagingRecord] = []
    for item in source.get("records", []):
        if item.get("source_domain") != "tender_requirement":
            continue
        if item.get("doc_role") not in TENDER_TEXT_ROLES:
            continue
        if str(item.get("suffix") or "").lower() not in TEXT_SUFFIXES:
            continue
        records.append(_stage_tender_text(item, out_dir))
    records.extend(_stage_goods_records(out_dir))
    records.extend(_stage_enterprise_records(out_dir))
    records.extend(_stage_asset_catalog_records(out_dir))
    records.extend(_stage_reference_records(out_dir))
    asset_path = _write_asset_payloads(out_dir)
    manifest_path = _write_manifest(records, out_dir)
    report_path = _write_report(records, asset_path, out_dir)
    return manifest_path, asset_path, report_path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=PARSED_ROOT / "staging")
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path, asset_path, report_path = build_staging(args.output_dir)
    print(json.dumps({
        "staging_manifest": _rel(manifest_path),
        "asset_payloads": _rel(asset_path),
        "report": _rel(report_path),
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
